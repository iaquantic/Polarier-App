"""
Extracción y estandarización de tablas desde archivos mensuales.

Soporta Excel (.xlsx/.xls), CSV, DOCX y PDF.
Produce RawTable: representación uniforme con encabezados, filas y metadatos
de origen independiente del formato fuente.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Any


@dataclass
class RawTable:
    """Tabla extraída de un archivo. Formato estandarizado independiente del origen."""
    source_file: str
    sheet_or_page: str
    headers: List[str]
    rows: List[List[Any]]
    table_index: int = 0
    raw_title: str = ""

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def col_count(self) -> int:
        return len(self.headers)

    def is_empty(self) -> bool:
        return self.row_count == 0 or self.col_count == 0

    def header_text(self) -> str:
        return " | ".join(str(h) for h in self.headers)


def extract_tables_from_file(path: Path) -> List[RawTable]:
    """Extrae todas las tablas de un archivo según su extensión."""
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xls"):
        return _extract_excel(path)
    if ext == ".csv":
        return _extract_csv(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".ods":
        return _extract_ods(path)
    return []


def _extract_excel(path: Path) -> List[RawTable]:
    try:
        import openpyxl
    except ImportError:
        try:
            import xlrd  # type: ignore
            return _extract_excel_xlrd(path)
        except ImportError:
            raise ImportError("Instalar openpyxl: pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True)
    tables: List[RawTable] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            continue

        # Detectar bloques de tabla separados por filas vacías
        blocks = _split_into_blocks(all_rows)
        for idx, block in enumerate(blocks):
            if len(block) < 2:
                continue
            headers = [str(c) if c is not None else "" for c in block[0]]
            rows = [
                [c if c is not None else "" for c in row]
                for row in block[1:]
            ]
            tables.append(RawTable(
                source_file=path.name,
                sheet_or_page=sheet_name,
                headers=headers,
                rows=rows,
                table_index=idx,
            ))

    return tables


def _extract_excel_xlrd(path: Path) -> List[RawTable]:
    import xlrd  # type: ignore
    wb = xlrd.open_workbook(str(path))
    tables: List[RawTable] = []
    for sheet in wb.sheets():
        if sheet.nrows < 2:
            continue
        headers = [str(sheet.cell_value(0, c)) for c in range(sheet.ncols)]
        rows = [
            [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            for r in range(1, sheet.nrows)
        ]
        tables.append(RawTable(
            source_file=path.name,
            sheet_or_page=sheet.name,
            headers=headers,
            rows=rows,
        ))
    return tables


def _extract_csv(path: Path) -> List[RawTable]:
    import csv
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        all_rows = list(reader)
    if len(all_rows) < 2:
        return []
    return [RawTable(
        source_file=path.name,
        sheet_or_page="csv",
        headers=[str(c) for c in all_rows[0]],
        rows=all_rows[1:],
    )]


def _extract_docx(path: Path) -> List[RawTable]:
    try:
        from docx import Document
    except ImportError:
        raise ImportError("Instalar python-docx: pip install python-docx")

    doc = Document(str(path))
    tables: List[RawTable] = []

    for idx, table in enumerate(doc.tables):
        rows_data = []
        for row in table.rows:
            rows_data.append([cell.text.strip() for cell in row.cells])
        if len(rows_data) < 2:
            continue
        tables.append(RawTable(
            source_file=path.name,
            sheet_or_page="docx",
            headers=rows_data[0],
            rows=rows_data[1:],
            table_index=idx,
        ))

    return tables


def _extract_pdf(path: Path) -> List[RawTable]:
    try:
        import fitz  # type: ignore
    except ImportError:
        raise ImportError("Instalar pymupdf: pip install pymupdf")

    doc = fitz.open(str(path))
    tables: List[RawTable] = []
    table_idx = 0

    for page_num, page in enumerate(doc):
        # Usar find_tables si está disponible (pymupdf >= 1.23)
        if hasattr(page, "find_tables"):
            tab_result = page.find_tables()
            for tbl in tab_result.tables:
                data = tbl.extract()
                if not data or len(data) < 2:
                    continue
                headers = [str(c) if c else "" for c in data[0]]
                rows = [[str(c) if c is not None else "" for c in row] for row in data[1:]]
                tables.append(RawTable(
                    source_file=path.name,
                    sheet_or_page=f"pag{page_num + 1}",
                    headers=headers,
                    rows=rows,
                    table_index=table_idx,
                ))
                table_idx += 1

    doc.close()
    return tables


def _extract_ods(path: Path) -> List[RawTable]:
    try:
        import ezodf  # type: ignore
    except ImportError:
        raise ImportError("Instalar ezodf: pip install ezodf")

    doc = ezodf.opendoc(str(path))
    tables: List[RawTable] = []
    for sheet in doc.sheets:
        all_rows = [
            [str(sheet[r, c].value) if sheet[r, c].value is not None else ""
             for c in range(sheet.ncols())]
            for r in range(sheet.nrows())
        ]
        if len(all_rows) < 2:
            continue
        tables.append(RawTable(
            source_file=path.name,
            sheet_or_page=sheet.name,
            headers=all_rows[0],
            rows=all_rows[1:],
        ))
    return tables


def _split_into_blocks(rows: list) -> List[list]:
    """Divide una hoja en bloques separados por filas completamente vacías."""
    blocks: List[list] = []
    current: list = []
    for row in rows:
        if all(c is None or str(c).strip() == "" for c in row):
            if current:
                blocks.append(current)
                current = []
        else:
            current.append(row)
    if current:
        blocks.append(current)
    return blocks


def extract_all_tables(paths: List[Path]) -> List[RawTable]:
    """Extrae tablas de múltiples archivos. Ignora errores por archivo individual."""
    all_tables: List[RawTable] = []
    for path in paths:
        try:
            tables = extract_tables_from_file(path)
            all_tables.extend(tables)
        except Exception as e:
            print(f"[WARN] No se pudieron extraer tablas de {path.name}: {e}")
    return all_tables
