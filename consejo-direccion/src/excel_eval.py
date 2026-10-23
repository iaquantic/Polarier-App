"""
Evaluador de fórmulas para la TABLA CA y el GLOBAL.

Los dos libros se guardan en el VPS con openpyxl (el agente de la TABLA CA y el
Modelo 1), y openpyxl no conserva el resultado de las fórmulas: con
`data_only=True` todo lo calculado sale vacío. Este módulo recalcula lo que haga
falta, celda a celda y bajo demanda, con lo que usan esos libros: aritmética,
comparaciones, referencias (también a otras hojas y a libros externos) y unas
pocas funciones (SUM, AVERAGE…).

Los libros externos se resuelven, por orden:
  1. con un libro vivo que se le pase (p. ej. el GLOBAL-CP26 del VPS), por nombre de fichero;
  2. con la copia en caché que guarda el propio libro (`externalLink`), si la hay.

Uso:
    libro = Libro("TABLA CA 2026.xlsx", externos={"GLOBAL-CP26.xlsx": Libro("GLOBAL-CP26.xlsx")})
    libro.valor("HABANA", 12, 8)        # fila 12, columna H
"""
from __future__ import annotations

import math
import re
import sys
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from urllib.parse import unquote

import openpyxl
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import column_index_from_string, range_boundaries

sys.setrecursionlimit(50000)


class Err:
    """Error de Excel (#DIV/0!, #VALUE!, #REF!…). Se propaga como en Excel."""

    __slots__ = ("codigo",)

    def __init__(self, codigo: str):
        self.codigo = codigo

    def __repr__(self):
        return self.codigo

    def __eq__(self, otro):
        return isinstance(otro, Err) and otro.codigo == self.codigo

    def __hash__(self):
        return hash(self.codigo)


DIV0 = Err("#DIV/0!")
VALUE = Err("#VALUE!")
REF = Err("#REF!")
NA = Err("#N/A")


class Rango:
    """Valores de un rango, para pasarlos a una función."""

    __slots__ = ("valores",)

    def __init__(self, valores: List):
        self.valores = valores


def _compacto(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s).strip().lower()


def _nombre_fichero(target: str) -> str:
    """'file:///C:\\...\\GLOBAL-CP25.xlsx' o '/a/b%20c.xlsx' → 'global-cp25.xlsx'."""
    t = unquote(target or "").replace("\\", "/")
    return _compacto(t.rsplit("/", 1)[-1])


_CELDA = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")


class _CacheExterna:
    """Valores que el libro guarda de un libro externo (parte externalLink)."""

    def __init__(self, link):
        eb = link.externalBook
        self.hojas: List[str] = list(eb.sheetNames.sheetName) if eb.sheetNames else []
        self.datos: Dict[Tuple[str, int, int], object] = {}
        # Hojas que el libro guardó: True si su última actualización falló. Excel
        # da #REF! en una hoja que no guardó, y también en las celdas que faltan
        # de una hoja que falló al actualizarse; en una sana, la que falta vale 0.
        self._hojas_guardadas: Dict[str, bool] = {}
        if eb.sheetDataSet:
            for sd in eb.sheetDataSet.sheetData:
                if sd.sheetId is None or sd.sheetId >= len(self.hojas):
                    continue
                hoja = _compacto(self.hojas[sd.sheetId])
                self._hojas_guardadas[hoja] = bool(getattr(sd, "refreshError", False))
                for fila in sd.row or []:
                    for celda in fila.cell or []:
                        m = _CELDA.match(celda.r or "")
                        if not m:
                            continue
                        v = celda.v
                        if celda.t in (None, "n") and v is not None:
                            try:
                                v = float(v)
                            except ValueError:
                                pass
                        elif celda.t == "e":
                            v = Err(v)
                        self.datos[(hoja, int(m.group(2)), column_index_from_string(m.group(1)))] = v

    def valor(self, hoja: str, fila: int, col: int):
        h = _compacto(hoja)
        if h not in self._hojas_guardadas:
            return REF
        v = self.datos.get((h, fila, col))
        if v is None and self._hojas_guardadas[h]:
            return REF
        return v


class _ConRespaldo:
    """
    Un libro externo vivo con la caché del vínculo como respaldo: si la celda del
    libro vivo está en blanco, vale lo que guardó el vínculo. Es el caso de las
    filas del GLOBAL que se teclean en el PC de quien monta la TABLA CA y no llegan
    al GLOBAL del VPS (turistas-día plan, producción real…).
    """

    def __init__(self, vivo: "Libro", cache: _CacheExterna):
        self.vivo, self.cache = vivo, cache
        self.usos_respaldo = 0

    def valor(self, hoja: str, fila: int, col: int):
        h = self.vivo.hoja(hoja)
        if h is not None and self.vivo.f[h].cell(fila, col).value is not None:
            return self.vivo.valor(h, fila, col)
        v = self.cache.valor(hoja, fila, col)
        if v is not None and not isinstance(v, Err):
            self.usos_respaldo += 1
        return v


class Libro:
    """Un libro de Excel con sus fórmulas, que se evalúan bajo demanda."""

    def __init__(self, ruta, externos: Optional[Dict[str, "Libro"]] = None, usar_cache: bool = False):
        """
        ruta        el .xlsx
        externos    libros vivos para resolver vínculos externos, por nombre de fichero
        usar_cache  si una fórmula trae resultado guardado (el libro lo guardó Excel),
                    usarlo en vez de recalcular. En el VPS no lo hay; en local sirve para
                    comprobar el evaluador comparando con lo que calculó Excel.
        """
        self.ruta = Path(ruta)
        self.f = openpyxl.load_workbook(self.ruta, data_only=False)
        self.v = openpyxl.load_workbook(self.ruta, data_only=True) if usar_cache else None
        self._hojas = {_compacto(n): n for n in self.f.sheetnames}
        self._memo: Dict[Tuple[str, int, int], object] = {}
        self._parsed: Dict[str, object] = {}
        vivos = {_compacto(k): v for k, v in (externos or {}).items()}
        # [n] en una fórmula → el n-ésimo externalLink del libro
        self._ext: Dict[int, object] = {}
        for i, link in enumerate(self.f._external_links, 1):
            nombre = _nombre_fichero(link.file_link.Target if link.file_link else "")
            cache = _CacheExterna(link)
            self._ext[i] = _ConRespaldo(vivos[nombre], cache) if nombre in vivos else cache
        self.nombres_externos = {
            i: _nombre_fichero(l.file_link.Target if l.file_link else "") for i, l in enumerate(self.f._external_links, 1)
        }

    # ── acceso ──────────────────────────────────────────────────

    def hoja(self, nombre: str) -> Optional[str]:
        return self._hojas.get(_compacto(nombre))

    def valor(self, hoja: str, fila: int, col: int):
        """Valor de una celda: número, texto, bool, None (vacía) o Err."""
        h = self.hoja(hoja)
        if h is None:
            return REF
        clave = (h, fila, col)
        if clave in self._memo:
            return self._memo[clave]
        self._memo[clave] = REF  # corta ciclos
        crudo = self.f[h].cell(fila, col).value
        if isinstance(crudo, str) and crudo.startswith("="):
            res = None
            if self.v is not None:
                res = self.v[h].cell(fila, col).value
                if isinstance(res, str) and res.startswith("#"):
                    res = Err(res)
            if res is None:
                try:
                    res = self._evaluar(h, crudo)
                except RecursionError:
                    res = REF
                except ValueError:
                    res = Err("#PARSE!")  # fórmula que este evaluador no sabe leer
            if isinstance(res, Rango):
                res = res.valores[0] if res.valores else None
            if res is None:
                res = 0.0  # una fórmula que apunta a una celda vacía da 0
        elif hasattr(crudo, "text") and not isinstance(crudo, str):  # ArrayFormula
            res = self._evaluar(h, crudo.text) if crudo.text else None
        else:
            res = crudo
        self._memo[clave] = res
        return res

    def numero(self, hoja: str, fila: int, col: int) -> Optional[float]:
        v = self.valor(hoja, fila, col)
        if isinstance(v, bool):
            return float(v)
        if isinstance(v, (int, float)):
            return float(v)
        return None

    # ── evaluación ──────────────────────────────────────────────

    def _evaluar(self, hoja: str, formula: str):
        ast = self._parsed.get(formula)
        if ast is None:
            ast = _Parser(formula).parse()
            self._parsed[formula] = ast
        return self._ev(ast, hoja)

    def _ev(self, nodo, hoja):
        tipo = nodo[0]
        if tipo == "num":
            return nodo[1]
        if tipo == "txt":
            return nodo[1]
        if tipo == "bool":
            return nodo[1]
        if tipo == "err":
            return Err(nodo[1])
        if tipo == "ref":
            return self._ref(nodo[1], hoja)
        if tipo == "neg":
            v = _num(self._escalar(self._ev(nodo[1], hoja)))
            return v if isinstance(v, Err) else -v
        if tipo == "pct":
            v = _num(self._escalar(self._ev(nodo[1], hoja)))
            return v if isinstance(v, Err) else v / 100.0
        if tipo == "bin":
            return _binario(nodo[1], self._escalar(self._ev(nodo[2], hoja)), self._escalar(self._ev(nodo[3], hoja)))
        if tipo == "fn":
            # Una referencia como argumento se trata como rango aunque sea de una
            # celda: SUM y AVERAGE ignoran entonces las vacías y los textos, como Excel.
            args = []
            for a in nodo[2]:
                v = self._ev(a, hoja)
                args.append(Rango([v]) if a[0] == "ref" and not isinstance(v, Rango) else v)
            return self._funcion(nodo[1], args)
        raise ValueError(f"nodo desconocido {tipo}")

    @staticmethod
    def _escalar(v):
        if isinstance(v, Rango):
            return v.valores[0] if len(v.valores) == 1 else VALUE
        return v

    def _ref(self, texto: str, hoja_ctx: str):
        libro_idx, hoja, rango = _partir_ref(texto)
        if libro_idx is not None:
            fuente = self._ext.get(libro_idx)
            if fuente is None:
                return REF
        else:
            fuente = self
        if hoja is None:
            hoja = hoja_ctx
        try:
            c1, f1, c2, f2 = range_boundaries(rango.replace("$", ""))
        except Exception:
            return REF  # nombre definido u otra cosa que no se resuelve
        if c1 is None or f1 is None:
            return REF  # columnas o filas completas: no se usan
        if (c1, f1) == (c2, f2):
            return fuente.valor(hoja, f1, c1)
        return Rango([fuente.valor(hoja, f, c) for f in range(f1, f2 + 1) for c in range(c1, c2 + 1)])

    def _funcion(self, nombre: str, args: List):
        nombre = nombre.upper().replace("_XLFN.", "")
        if nombre in ("SUM", "AVERAGE", "MAX", "MIN", "COUNT", "COUNTA", "PRODUCT"):
            nums, err = [], None
            for a in args:
                if isinstance(a, Rango):
                    for x in a.valores:
                        if isinstance(x, Err):
                            err = err or x
                        elif isinstance(x, bool):
                            continue  # en un rango, los lógicos no cuentan
                        elif isinstance(x, (int, float)):
                            nums.append(float(x))
                        elif nombre == "COUNTA" and x not in (None, ""):
                            nums.append(1.0)
                else:
                    x = _num(a) if a is not None else 0.0
                    if isinstance(x, Err):
                        err = err or x
                    else:
                        nums.append(x)
            if err is not None and nombre != "COUNT":
                return err
            if nombre == "SUM":
                return sum(nums)
            if nombre == "AVERAGE":
                return sum(nums) / len(nums) if nums else DIV0
            if nombre == "MAX":
                return max(nums) if nums else 0.0
            if nombre == "MIN":
                return min(nums) if nums else 0.0
            if nombre in ("COUNT", "COUNTA"):
                return float(len(nums))
            if nombre == "PRODUCT":
                return math.prod(nums) if nums else 0.0
        esc = [self._escalar(a) for a in args]
        if nombre in ("ABS", "IMABS"):
            v = _num(esc[0])
            return v if isinstance(v, Err) else abs(v)
        if nombre == "ROUND":
            v, d = _num(esc[0]), _num(esc[1]) if len(esc) > 1 else 0.0
            if isinstance(v, Err):
                return v
            q = 10 ** int(d)
            return math.floor(abs(v) * q + 0.5) / q * (1 if v >= 0 else -1)
        if nombre == "IF":
            cond = esc[0]
            si = esc[1] if len(esc) > 1 else True
            no = esc[2] if len(esc) > 2 else False
            if isinstance(cond, Err):
                return cond
            return si if _verdad(cond) else no
        if nombre == "IFERROR":
            return esc[1] if isinstance(esc[0], Err) else esc[0]
        if nombre == "ISERROR":
            return isinstance(esc[0], Err)
        return Err("#NAME?")


def _verdad(v) -> bool:
    if isinstance(v, str):
        return v.upper() == "TRUE"
    return bool(v)


def _num(v):
    """Valor como número, a la manera de Excel en aritmética."""
    if isinstance(v, Err):
        return v
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", "").strip())
        except ValueError:
            return VALUE
    return VALUE


def _binario(op: str, a, b):
    if op == "&":
        if isinstance(a, Err):
            return a
        if isinstance(b, Err):
            return b
        return f"{'' if a is None else a}{'' if b is None else b}"
    if op in ("=", "<>", "<", ">", "<=", ">="):
        if isinstance(a, Err):
            return a
        if isinstance(b, Err):
            return b
        a = 0.0 if a is None else a
        b = 0.0 if b is None else b
        if isinstance(a, str) != isinstance(b, str):
            res = {"=": False, "<>": True}.get(op)
            if res is not None:
                return res
            # Excel: los textos son mayores que los números
            a, b = (1, 0) if isinstance(a, str) else (0, 1)
        elif isinstance(a, str):
            a, b = a.lower(), b.lower()
        return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]
    x, y = _num(a), _num(b)
    if isinstance(x, Err):
        return x
    if isinstance(y, Err):
        return y
    if op == "+":
        return x + y
    if op == "-":
        return x - y
    if op == "*":
        return x * y
    if op == "/":
        return DIV0 if y == 0 else x / y
    if op == "^":
        try:
            return x ** y
        except (OverflowError, ZeroDivisionError, ValueError):
            return Err("#NUM!")
    raise ValueError(f"operador desconocido {op}")


_REF_LIBRO = re.compile(r"^\[(\d+)\](.*)$")


def _partir_ref(texto: str) -> Tuple[Optional[int], Optional[str], str]:
    """'[2]LHCP2026!$G$1' → (2, 'LHCP2026', '$G$1');  "'Hoja 1'!A1:B2" → (None, 'Hoja 1', 'A1:B2')."""
    if "!" not in texto:
        return None, None, texto
    hoja, rango = texto.rsplit("!", 1)
    if hoja.startswith("'") and hoja.endswith("'"):
        hoja = hoja[1:-1].replace("''", "'")
    libro = None
    m = _REF_LIBRO.match(hoja)
    if m:
        libro, hoja = int(m.group(1)), m.group(2)
    return libro, hoja, rango


# ── Parser ──────────────────────────────────────────────────────

_PREC = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1, "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}


class _Parser:
    """Convierte una fórmula en un árbol de tuplas, con el tokenizador de openpyxl."""

    def __init__(self, formula: str):
        tok = Tokenizer(formula)
        self.t = [x for x in tok.items if x.type != Token.WSPACE]
        self.i = 0

    def parse(self):
        if not self.t:
            return ("txt", "")
        nodo = self._expr(0)
        return nodo

    def _mirar(self):
        return self.t[self.i] if self.i < len(self.t) else None

    def _tomar(self):
        x = self.t[self.i]
        self.i += 1
        return x

    def _expr(self, prec_min):
        izq = self._unario()
        while True:
            x = self._mirar()
            if x is None or x.type != Token.OP_IN or x.value not in _PREC:
                break
            prec = _PREC[x.value]
            if prec < prec_min:
                break
            self._tomar()
            # ^ es asociativo por la izquierda en Excel, igual que el resto
            der = self._expr(prec + 1)
            izq = ("bin", x.value, izq, der)
        return izq

    def _unario(self):
        x = self._mirar()
        if x is not None and x.type == Token.OP_PRE:
            self._tomar()
            dentro = self._unario()
            return ("neg", dentro) if x.value == "-" else dentro
        nodo = self._primario()
        while self._mirar() is not None and self._mirar().type == Token.OP_POST and self._mirar().value == "%":
            self._tomar()
            nodo = ("pct", nodo)
        return nodo

    def _primario(self):
        x = self._tomar()
        if x.type == Token.OPERAND:
            if x.subtype == Token.NUMBER:
                return ("num", float(x.value))
            if x.subtype == Token.TEXT:
                return ("txt", x.value[1:-1].replace('""', '"'))
            if x.subtype == Token.LOGICAL:
                return ("bool", x.value.upper() == "TRUE")
            if x.subtype == Token.ERROR:
                return ("err", x.value)
            return ("ref", x.value)
        if x.type == Token.FUNC and x.subtype == Token.OPEN:
            nombre = x.value[:-1]
            args = []
            if self._mirar() is not None and self._mirar().type == Token.FUNC and self._mirar().subtype == Token.CLOSE:
                self._tomar()
                return ("fn", nombre, args)
            while True:
                sig = self._mirar()
                if sig is not None and (sig.type == Token.SEP or (sig.type == Token.FUNC and sig.subtype == Token.CLOSE)):
                    args.append(("txt", None))  # argumento vacío: SUM(A1,) o IF(x,,y)
                else:
                    args.append(self._expr(0))
                y = self._tomar()
                if y.type == Token.SEP:
                    continue
                if y.type == Token.FUNC and y.subtype == Token.CLOSE:
                    break
                raise ValueError(f"se esperaba ',' o ')' en la función {nombre}")
            return ("fn", nombre, args)
        if x.type == Token.PAREN and x.subtype == Token.OPEN:
            nodo = self._expr(0)
            y = self._tomar()
            if not (y.type == Token.PAREN and y.subtype == Token.CLOSE):
                raise ValueError("falta ')'")
            return nodo
        raise ValueError(f"token inesperado {x.type}:{x.value}")
