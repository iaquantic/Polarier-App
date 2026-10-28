"""Tablas de lencería desde la lencería del CD (src/lenceria_cd.py). Números inventados."""
import sys
from pathlib import Path
from types import SimpleNamespace

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from lenceria_cd import CONSTRUCTORES_CD, Lenceria  # noqa: E402
from tablas_informe import Periodo, celdas, filas, texto_celda  # noqa: E402

PRENDAS_PROTO = ["Sábana Personal", "Sábana King", "Funda Almohada", "Toalla Alfombra", "Toalla Baño",
                 "Toalla Cara", "Toalla Facial", "Toalla Piscina"]


def proto(titulo, cabecera, rotulos):
    doc = Document()
    t = doc.add_table(rows=2 + len(rotulos), cols=len(cabecera))
    t.cell(0, 0).text = titulo
    for j, c in enumerate(cabecera):
        t.cell(1, j).text = c
    for i, r in enumerate(rotulos):
        t.cell(2 + i, 0).text = r
        for j in range(1, len(cabecera)):
            t.cell(2 + i, j).text = "99"  # cifra del informe de ejemplo, que se sustituye
    return t._tbl


def texto(tbl):
    return [[texto_celda(tc) for tc in celdas(tr)] for tr in filas(tbl)]


def fila(anio, mes, polo, bloque, valores, estado="ok"):
    return {"anio": anio, "mes": mes, "polo": polo, "bloque": bloque, "estado": estado,
            "valores": valores, "totales": {}}


def fu(filas_cd):
    return SimpleNamespace(lenceria=Lenceria(filas_cd))


def cant(**kw):
    return {p: {"cantidad": v} for p, v in kw.items()}


RECHAZO = proto("RECHAZO DE LENCERÍA", ["Lencería", "Habana", "Varadero", "Caibarién", "TOTAL", "TOTAL 2029"],
                PRENDAS_PROTO + ["TOTAL PRENDAS", "TOTAL RECHAZO (kg)", "RECHAZO/PROD. (%)"])


def test_rechazo_vacio_cerrado_y_total():
    datos = [fila(2030, 5, "habana", "rechazo", {}, estado="vacio"),
             fila(2030, 5, "varadero", "rechazo", cant(sabana_personal=1200, funda=30))]
    [r] = CONSTRUCTORES_CD["TABLA_RECHAZO"](fu(datos), Periodo(2030, mes=5), [RECHAZO], 2030)
    t = texto(r.xml)
    assert t[2] == ["Sábana Personal", "—", "1,200", "-", "1,200", ""]
    assert t[3] == ["Sábana King", "—", "-", "-", "-", ""]  # en un bloque rellenado, el blanco es cero
    assert t[10] == ["TOTAL PRENDAS", "—", "1,230", "-", "1,230", ""]
    assert t[11][1:] == ["", "", "", "", ""]  # kg: falta el peso de cada prenda
    assert any("Habana" in a and "sin rellenar" in a for a in r.avisos)
    assert any("peso de cada prenda" in a for a in r.avisos)
    assert any("año anterior va en blanco" in a for a in r.avisos)


def test_trimestre_suma_y_anio_anterior():
    datos = [fila(2030, m, "varadero", "mal_uso", cant(toalla_bano=10 * m)) for m in (4, 5, 6)]
    datos += [fila(2029, 5, "caibarien", "mal_uso", cant(toalla_bano=7))]
    tb = proto("MAL USO", ["Lencería", "Varadero", "TOTAL", "TOTAL 2029"], PRENDAS_PROTO + ["TOTAL PRENDAS"])
    [r] = CONSTRUCTORES_CD["TABLA_MAL_USO"](fu(datos), Periodo(2030, trimestre=2), [tb], 2030)
    t = texto(r.xml)
    bano = next(f for f in t if f[0] == "Toalla Baño")
    assert bano == ["Toalla Baño", "150", "150", "7"]
    assert not any("año anterior" in a for a in r.avisos)


def test_montas_promedio_de_los_polos_del_prototipo():
    idx = lambda v: {"sabana_personal": {"indice": v}}  # noqa: E731
    datos = [fila(2030, 5, "habana", "montas_habitacion", idx(1.0)),
             fila(2030, 5, "varadero", "montas_habitacion", idx(2.0)),
             fila(2030, 5, "holguin", "montas_habitacion", idx(9.0))]  # no está en el prototipo
    tb = proto("ESTADO DE LAS MONTAS POR HABITACIÓN", ["Lencería", "Habana", "Varadero", "Promedio"], PRENDAS_PROTO)
    [r] = CONSTRUCTORES_CD["TABLA_MONTAS_HABITACION"](fu(datos), Periodo(2030, mes=5), [tb], 2030)
    assert texto(r.xml)[2] == ["Sábana Personal", "1.00", "2.00", "1.50"]


def test_faltantes_sin_perdida_economica():
    datos = [fila(2030, 5, "habana", "faltantes", {"sabana_king": {"faltantes": 12, "sobrantes": 3}})]
    tb = proto("TOTAL DE FALTANTES DETECTADOS EN INVENTARIOS MAYO",
               ["PRENDAS", "Habana", "Cayo Coco", "Total", "Pérdida económica (CUP)", "FALTANTES 2029"],
               ["Sábanas Personales", "Sábanas King", "Total"])
    [r] = CONSTRUCTORES_CD["TABLA_FALTANTES"](fu(datos), Periodo(2030, mes=5), [tb], 2030)
    t = texto(r.xml)
    assert t[3] == ["Sábanas King", "12", "-", "12", "", ""]
    assert any("Economía" in a for a in r.avisos)
    assert not any("perdida economica" in a for a in r.avisos)


def test_desmanche_porcentaje_de_las_sumas():
    doc = Document()
    t = doc.add_table(rows=5, cols=5)
    t.cell(0, 0).text = "TRATAMIENTOS DE DESMANCHE MAYO"
    t.cell(1, 0).text = "Lencería"
    a = t.cell(1, 1).merge(t.cell(1, 2))
    a.text = "Varadero"
    b = t.cell(1, 3).merge(t.cell(1, 4))
    b.text = "TOTAL"
    for j, s in enumerate(["", "Trat.", "%Rec.", "Trat.", "%Rec."]):
        t.cell(2, j).text = s
    t.cell(3, 0).text = "Toalla Baño"
    t.cell(4, 0).text = "TOTAL"
    datos = [fila(2030, m, "varadero", "desmanche", {"toalla_bano": {"tratadas": 100, "recuperadas": 80 + m}})
             for m in (4, 5, 6)]
    [r] = CONSTRUCTORES_CD["TABLA_DESMANCHE"](fu(datos), Periodo(2030, trimestre=2), [t._tbl], 2030)
    out = texto(r.xml)
    assert out[3] == ["Toalla Baño", "300", "85%", "300", "85%"]
