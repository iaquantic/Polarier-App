"""Tests para excel_eval.py, formato_excel.py y el relleno de tablas (sin datos reales)."""

import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT / "src"))

from excel_eval import Err, Libro  # noqa: E402
from formato_excel import formatear  # noqa: E402
from tablas_informe import Periodo, clave_polo, correr_anios, poner_periodo  # noqa: E402


@pytest.fixture()
def libro(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HABANA"
    ws["A1"], ws["A2"], ws["A3"] = 10, 20, None
    ws["B1"] = "=SUM(A1:A3)"
    ws["B2"] = "=AVERAGE(A1,A3,A2)"        # la vacía no cuenta
    ws["B3"] = "=A1/A3"                     # entre vacía → #DIV/0!
    ws["B4"] = "=A3"                        # apunta a vacía → 0
    ws["B5"] = "=-(A1+A2)*2%"
    ws["B6"] = "=CONSOLIDADO!A1+1"
    ws["B7"] = "=B3+1"                      # el error se propaga
    ws["B8"] = "=SUM(A1,)"                  # argumento vacío al final
    ws["B9"] = "=AVERAGE(A3)"               # sin números → #DIV/0!
    otra = wb.create_sheet("CONSOLIDADO")
    otra["A1"] = 5
    ruta = tmp_path / "prueba.xlsx"
    wb.save(ruta)
    return Libro(ruta)


def test_aritmetica_y_funciones(libro):
    assert libro.valor("HABANA", 1, 2) == 30
    assert libro.valor("HABANA", 2, 2) == 15
    assert libro.valor("HABANA", 5, 2) == pytest.approx(-0.6)
    assert libro.valor("HABANA", 6, 2) == 6
    assert libro.valor("HABANA", 8, 2) == 10


def test_errores_y_vacias_como_excel(libro):
    assert libro.valor("HABANA", 3, 2) == Err("#DIV/0!")
    assert libro.valor("HABANA", 4, 2) == 0
    assert libro.valor("HABANA", 7, 2) == Err("#DIV/0!")
    assert libro.valor("HABANA", 9, 2) == Err("#DIV/0!")


def test_hoja_por_nombre_sin_tildes(libro):
    assert libro.hoja("habana") == "HABANA"
    assert libro.hoja("NO EXISTE") is None


@pytest.mark.parametrize("valor,fmt,esperado", [
    (12345678.905, "#,##0.00", "12,345,678.91"),
    (42.164, "0.00", "42.16"),
    (0.5, "0%", "50%"),
    (0.12345, "0.00%", "12.35%"),
    (51234.4, "#,##0", "51,234"),
    (0, '_-* #,##0.00\\ _€_-;\\-* #,##0.00\\ _€_-;_-* "-"??\\ _€_-;_-@_-', "-"),
    (-1234.5, '_-* #,##0.00\\ _€_-;\\-* #,##0.00\\ _€_-;_-* "-"??\\ _€_-;_-@_-', "-1,234.50"),
    (1234.5, '_-* #,##0.00\\ _€_-;\\-* #,##0.00\\ _€_-;_-* "-"??\\ _€_-;_-@_-', "1,234.50"),
    (5, "General", "5"),
    (-0.0001, "0.00", "0.00"),
    (Err("#DIV/0!"), "0.00", "-"),
    (None, "0.00", ""),
])
def test_formato_de_excel(valor, fmt, esperado):
    assert formatear(valor, fmt) == esperado


def test_periodo_y_anios_en_rotulos():
    junio = Periodo(2026, mes=6)
    t2 = Periodo(2026, trimestre=2)
    assert poner_periodo("Costo y gasto por Peso de Ingreso MAYO", junio) == "Costo y gasto por Peso de Ingreso JUNIO"
    assert poner_periodo("RETIRO GENERADO JUNIO", t2) == "RETIRO GENERADO II TRIMESTRE"
    assert correr_anios("REAL 26 - REAL 25", 1) == "REAL 27 - REAL 26"
    assert correr_anios("Real 2025", 1) == "Real 2026"
    assert t2.meses == [4, 5, 6]


@pytest.mark.parametrize("texto,clave", [
    ("La Habana", "HA"), ("077 LAVANDERÍA HABANA", "HA"), ("082 Centro Lavandería", "CA"),
    ("Caibarién", "CA"), ("LAV. CAYO CRUZ 0083", "CX"), ("Holguín", "HO"), ("SUMA", "TOTAL"),
    ("Consolidado", "CONSOLIDADO"), ("Promedio", "PROMEDIO"), ("Plan", None),
])
def test_clave_de_polo(texto, clave):
    assert clave_polo(texto) == clave
