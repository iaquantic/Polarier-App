"""
Monta el informe del Consejo: plantilla maestra + DOCX de tablas + textos de análisis.

Es el mismo montaje que scripts/generate_report.py, con rutas explícitas para que
lo llame n8n por el pythonrunner:

    python scripts/montar_consejo.py --tablas trabajo/2026-06/tablas.docx \
        --periodo JUNIO --anio 2026 --nombre consejo_direccion_JUNIO_2026.docx \
        (--analisis analisis.json | --analisis-b64 <trozo> [<trozo>…])

--analisis-b64: el JSON de análisis en base64 (UTF-8), partido en trozos para no
pasar del límite de longitud de un argumento. Se admite {"sections": {...}} o el
diccionario de secciones directamente.

Salida por stdout: {"ok": true, "nombre", "base64", "tablas_insertadas", ...} o
{"ok": false, "error"}. El DOCX queda también en trabajo/<período>/.
"""
from __future__ import annotations

import argparse
import base64
import json
import os
import re
import shutil
import sys
import traceback
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
os.chdir(RAIZ)
sys.stdout.reconfigure(encoding="utf-8")  # el JSON va a n8n en UTF-8 también desde Windows
sys.path.insert(0, str(RAIZ / "src"))

from docx import Document  # noqa: E402

from docx_content_replacer import replace_metadata  # noqa: E402
from report_builder import build_report  # noqa: E402
from table_markers import classify_tables, describe_tables, load_rules, template_table_markers  # noqa: E402

MASTER = RAIZ / "templates" / "master" / "consejo_direccion_master.docx"


def salir(obj) -> None:
    print(json.dumps(obj, ensure_ascii=False))
    sys.stdout.flush()


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--tablas", required=True)
    ap.add_argument("--periodo", required=True, help="JUNIO, II TRIMESTRE…")
    ap.add_argument("--anio", required=True)
    ap.add_argument("--nombre", required=True)
    ap.add_argument("--analisis")
    ap.add_argument("--analisis-b64", nargs="+")
    a = ap.parse_args()

    if a.analisis:
        datos = json.loads(Path(a.analisis).read_text(encoding="utf-8"))
    elif a.analisis_b64:
        datos = json.loads(base64.b64decode("".join(a.analisis_b64)).decode("utf-8"))
    else:
        raise ValueError("falta el análisis (--analisis o --analisis-b64)")
    secciones = datos.get("sections", datos) if isinstance(datos, dict) else {}
    if not isinstance(secciones, dict) or not secciones:
        raise ValueError("el análisis no trae secciones")

    ruta_tablas = Path(a.tablas)
    if not ruta_tablas.is_absolute():
        ruta_tablas = RAIZ / ruta_tablas
    tablas = describe_tables(Document(str(ruta_tablas)))
    clas = classify_tables(tablas, load_rules())
    if not clas.ok:
        raise ValueError("tablas sin reconocer: " + "; ".join(clas.errors))
    if not MASTER.exists():
        raise FileNotFoundError(f"no está la plantilla maestra {MASTER}")
    en_plantilla = template_table_markers(Document(str(MASTER)))
    faltan = sorted({x.marker for x in clas.assignments} - en_plantilla)
    if faltan:
        raise ValueError("marcadores que no están en la plantilla: " + ", ".join(faltan))
    mapa_xml = clas.marker_to_xml(tablas)

    nombre = Path(a.nombre).name
    salida = ruta_tablas.parent / nombre
    shutil.copy2(MASTER, salida)
    replace_metadata(salida, mes=a.periodo, año=str(a.anio))

    analisis = {k: v for k, v in secciones.items()
                if k.startswith("ANALISIS_") and isinstance(v, str) and v.strip()
                and not re.match(r"^\[dato no disponible\]", v.strip(), re.I)}
    otros = {k: v for k, v in secciones.items() if k.startswith(("SECCION_", "INTRO_")) and isinstance(v, str)}
    sin_texto = sorted(m.replace("TABLA_", "ANALISIS_") for m in mapa_xml
                       if m.replace("TABLA_", "ANALISIS_") not in analisis)
    res = build_report(output_path=salida, source_xml_map=mapa_xml, analysis_texts=analisis, section_texts=otros)

    pendientes = []
    for p in Document(str(salida)).paragraphs:
        pendientes += re.findall(r"\{[A-Z_ÁÉÍÓÚÑ]+\}", p.text)
    salir({
        "ok": True,
        "nombre": nombre,
        "archivo": str(salida),
        "tablas_insertadas": res.get("tables_inserted"),
        "tablas_eliminadas": res.get("tables_removed"),
        "analisis_insertados": res.get("analyses_inserted"),
        "analisis_que_faltan": sin_texto,
        "marcadores_pendientes": pendientes,
        "base64": base64.b64encode(salida.read_bytes()).decode("ascii"),
    })


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        salir({"ok": False, "error": f"{type(e).__name__}: {e}", "traza": traceback.format_exc()[-2000:]})
