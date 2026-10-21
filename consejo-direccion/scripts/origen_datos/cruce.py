"""Cruce por valores: busca cada cifra de las tablas de un informe en las celdas del GLOBAL y la TABLA CA.

Uso: py -3 -P cruce.py <informe.json> <salida.json> <libro1.xlsx> [<libro2.xlsx> ...]

Los GLOBAL se guardaron con openpyxl y no traen valores de fórmulas en caché:
se evalúan aquí (SUM, aritmética y referencias entre hojas, que es todo lo que usan).
"""
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, range_boundaries

REF = re.compile(r"(?:(?:\[\d+\])?'?([^'!=+\-*/(),]+?)'?!)?\$?([A-Z]{1,3})\$?(\d+)")
SUM = re.compile(r"SUM\(([^()]*)\)", re.I)
ARIT = re.compile(r"[\d\.\s+\-*/()eE]*")


class Libro:
    def __init__(self, ruta):
        self.ruta = ruta
        self.nombre = Path(ruta).name
        self.f = openpyxl.load_workbook(ruta, data_only=False)
        self.v = openpyxl.load_workbook(ruta, data_only=True)
        self.memo = {}

    def hoja(self, nombre):
        if nombre in self.f.sheetnames:
            return nombre
        for n in self.f.sheetnames:  # 3er. Trimestre lleva un espacio final, etc.
            if n.strip() == nombre.strip():
                return n
        return None

    def valor(self, hoja, fila, col):
        clave = (hoja, fila, col)
        if clave in self.memo:
            return self.memo[clave]
        self.memo[clave] = None  # corta ciclos
        raw = self.f[hoja].cell(fila, col).value
        if isinstance(raw, str) and raw.startswith("="):
            cache = self.v[hoja].cell(fila, col).value
            res = float(cache) if isinstance(cache, (int, float)) and not isinstance(cache, bool) else self._formula(hoja, raw)
        elif isinstance(raw, (int, float)) and not isinstance(raw, bool):
            res = float(raw)
        else:
            res = None
        self.memo[clave] = res
        return res

    def _celda(self, hoja_ctx, hoja_ref, col, fila):
        h = self.hoja(hoja_ref) if hoja_ref else hoja_ctx
        if h is None:
            return None
        return self.valor(h, int(fila), column_index_from_string(col))

    def _formula(self, hoja, formula):
        expr = formula[1:].replace("$", "")
        if "#REF!" in expr:
            expr = expr.replace("#REF!", "0")

        def suma(m):
            total = 0.0
            for parte in m.group(1).split(","):
                parte = parte.strip()
                hoja_ref = None
                if "!" in parte:
                    hoja_ref, parte = parte.rsplit("!", 1)
                    hoja_ref = re.sub(r"^\[\d+\]", "", hoja_ref).strip("'")
                h = self.hoja(hoja_ref) if hoja_ref else hoja
                if h is None:
                    continue
                if ":" in parte:
                    c1, f1, c2, f2 = range_boundaries(parte)
                    for f in range(f1, f2 + 1):
                        for c in range(c1, c2 + 1):
                            x = self.valor(h, f, c)
                            total += x or 0.0
                else:
                    m2 = REF.fullmatch(parte)
                    if m2:
                        total += self._celda(h, None, m2.group(2), m2.group(3)) or 0.0
            return repr(total)

        expr = SUM.sub(suma, expr)
        expr = REF.sub(lambda m: repr(self._celda(hoja, m.group(1), m.group(2), m.group(3)) or 0.0), expr)
        if not ARIT.fullmatch(expr):
            return None
        try:
            return float(eval(expr, {"__builtins__": {}}, {}))
        except Exception:
            return None


def texto(v):
    return v.strip() if isinstance(v, str) and v.strip() and not v.strip().startswith("=") else None


def etiquetas(libro, hoja, fila, col):
    """Etiqueta de fila (texto más cercano a la izquierda) y de columna (texto más cercano arriba)."""
    ws = libro.f[hoja]
    fila_txt = None
    for c in range(col - 1, 0, -1):
        t = texto(ws.cell(fila, c).value)
        if t:
            fila_txt = t
            break
    col_txt = None
    for f in range(fila - 1, max(0, fila - 80), -1):
        t = texto(ws.cell(f, col).value)
        if t:
            col_txt = t
            # GLOBAL: bajo PLAN/REAL/Qp, el mes está en la fila de arriba, en la columna del PLAN
            if t in ("PLAN", "REAL", "Qp") and f > 1:
                for c in range(col, max(0, col - 3), -1):
                    t2 = texto(ws.cell(f - 1, c).value)
                    if t2:
                        col_txt = f"{t2} · {t}"
                        break
            break
    return fila_txt, col_txt


def indexar(libros):
    idx = defaultdict(list)  # valor redondeado a 2 → [(libro, hoja, celda, valor, fila_txt, col_txt)]
    for lb in libros:
        for hoja in lb.f.sheetnames:
            ws = lb.f[hoja]
            for row in ws.iter_rows():
                for c in row:
                    x = lb.valor(hoja, c.row, c.column)
                    if x is None or x == 0:
                        continue
                    ft, ct = etiquetas(lb, hoja, c.row, c.column)
                    reg = (lb.nombre, hoja, c.coordinate, x, ft, ct)
                    idx[round(x, 2)].append(reg)
                    if abs(x) < 10:  # fracciones que el informe muestra en %
                        idx[round(x * 100, 2)].append(reg + ("x100",))
    return idx


NUM = re.compile(r"^\(?(-?[\d,]*\.?\d+)\)?\s*(%?)$")


def parsear(celda):
    s = celda.replace("−", "-").replace(" ", "")
    m = NUM.match(s)
    if not m:
        return None
    cuerpo = m.group(1)
    dec = len(cuerpo.split(".")[1]) if "." in cuerpo else 0
    try:
        return float(cuerpo.replace(",", "")), dec, bool(m.group(2))
    except ValueError:
        return None


def buscar(idx, v, dec):
    cands = []
    for clave in (round(v, 2), round(v, 2) + 0.01, round(v, 2) - 0.01) if dec < 2 else (round(v, 2),):
        for reg in idx.get(round(clave, 2), []):
            if round(reg[3] * (100 if len(reg) > 6 else 1), dec) == round(v, dec):
                cands.append(reg)
    # sin duplicados
    vistos, out = set(), []
    for r in cands:
        k = (r[0], r[1], r[2], len(r))
        if k not in vistos:
            vistos.add(k)
            out.append(r)
    return out


def main():
    informe, salida, *rutas = sys.argv[1:]
    libros = [Libro(r) for r in rutas]
    idx = indexar(libros)
    tablas = json.load(open(informe, encoding="utf-8"))
    res = []
    for t in tablas:
        filas_out = []
        tot = ok = 0
        for i, fila in enumerate(t["filas"]):
            etiqueta = fila[0] if fila else ""
            celdas = []
            for j, celda in enumerate(fila[1:], 1):
                p = parsear(celda)
                if not p or p[0] == 0:
                    continue
                v, dec, pct = p
                # Enteros pequeños y fracciones con pocas cifras casan con cualquier cosa
                trivial = (dec == 0 and abs(v) < 100) or (abs(v) < 1 and dec <= 2)
                cands = buscar(idx, v, dec)
                tot += 1
                especifico = 0 < len(cands) <= 12 and not trivial
                if especifico:
                    ok += 1
                celdas.append({
                    "col": j, "texto": celda, "trivial": trivial, "n": len(cands),
                    "origen": [{"libro": c[0], "hoja": c[1], "celda": c[2], "fila": c[4], "col": c[5], "x100": len(c) > 6} for c in cands[:6]] if especifico else [],
                })
            filas_out.append({"etiqueta": etiqueta, "celdas": celdas})
        res.append({"pos": t["pos"], "titulo": t["titulo"], "encabezado": t["encabezado"], "seccion": t["seccion"],
                    "numericas": tot, "casadas": ok, "filas": filas_out})
    json.dump(res, open(salida, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    for r in res:
        hojas = Counter()
        for f in r["filas"]:
            for c in f["celdas"]:
                for o in c["origen"]:
                    hojas[f"{o['libro'][:12]}:{o['hoja']}"] += 1
        print(f"[{r['pos']:>2}] {r['titulo'][:46]:<46} {r['casadas']:>3}/{r['numericas']:<3} {dict(hojas.most_common(3))}")


if __name__ == "__main__":
    main()
