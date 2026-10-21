"""Cruce por etiquetas entre las tablas del informe y las hojas del CD de un polo.

Uso: py -3 -P cruce_cd2.py <tablas.json> <polo> <posiciones> <CD.xlsx>

- Tablas «por prenda» (filas = prendas, columnas = polos): toma la columna del polo
  y busca en el CD una columna cuyas filas, emparejadas por nombre de prenda, den los
  mismos valores.
- Tablas «por polo» (filas = polos o centros, p. ej. vehículos o vacantes): toma la
  fila del polo y busca en el CD una fila con la misma secuencia de valores.
"""
import json
import re
import sys
import unicodedata

from cruce import Libro, parsear


def norm(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    t = re.sub(r"[^a-z ]", " ", t)
    t = re.sub(r"\btoallas\b", "toalla", t)
    t = re.sub(r"\bsabanas\b", "sabana", t)
    t = re.sub(r"\bfundas?\b.*", "funda", t)
    t = re.sub(r"\bpersonales\b", "personal", t)
    t = re.sub(r"\bfaciales\b", "facial", t)
    t = re.sub(r"\bde\b", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def columna_del_polo(filas, polo):
    for i, fila in enumerate(filas[:3]):
        for j, celda in enumerate(fila):
            if polo.lower() in celda.lower():
                out = {}
                for f in filas[i + 1:]:
                    if j < len(f) and f[0].strip():
                        p = parsear(f[j])
                        if p:
                            out[norm(f[0])] = p[:2]
                return out
    return {}


def fila_del_polo(filas, polo):
    for f in filas:
        if f and polo.lower() in f[0].lower():
            return [parsear(c)[:2] for c in f[1:] if parsear(c)]
    return []


def igual(a, b, dec):
    a = a or 0.0
    return round(a, dec) == round(b, dec) or round(a * 100, dec) == round(b, dec)


def main():
    tablas = json.load(open(sys.argv[1], encoding="utf-8"))
    polo = sys.argv[2]
    pos = {int(x) for x in sys.argv[3].split(",")}
    lb = Libro(sys.argv[4])

    # Índice del CD: para cada hoja, cada columna → {etiqueta de fila normalizada: valor}
    # La etiqueta de una fila es el último texto a su izquierda.
    columnas = []
    filas_cd = []
    for hoja in lb.f.sheetnames:
        ws = lb.f[hoja]
        etiqueta_fila = {}
        for r in range(1, ws.max_row + 1):
            etq = None
            vals = []
            for c in range(1, ws.max_column + 1):
                raw = ws.cell(r, c).value
                if isinstance(raw, str) and raw.strip() and not raw.startswith("="):
                    etq = (c, norm(raw))
                else:
                    x = lb.valor(hoja, r, c)
                    if x is not None and etq:
                        etiqueta_fila[(r, c)] = etq[1]
                    vals.append((c, x))
            filas_cd.append((hoja, r, vals))
        por_col = {}
        for (r, c), etq in etiqueta_fila.items():
            por_col.setdefault(c, []).append((r, etq, lb.valor(hoja, r, c)))
        for c, celdas in por_col.items():
            columnas.append((hoja, c, celdas))

    for t in tablas:
        if t["pos"] not in pos:
            continue
        col = columna_del_polo(t["filas"], polo)
        print(f"\n##### [{t['pos']}] {t['titulo'][:55]}")
        if col and any(v for v, _ in col.values()):
            mejor = (0, None)
            for hoja, c, celdas in columnas:
                # bloques contiguos de filas: se prueba cada arranque
                for i in range(len(celdas)):
                    bloque = {}
                    for r, etq, x in celdas[i:i + 14]:
                        bloque.setdefault(etq, x)
                    comunes = [k for k in col if k in bloque]
                    if len(comunes) < 4:
                        continue
                    ok = sum(igual(bloque[k], *col[k][::-1][::-1]) if False else igual(bloque[k], col[k][0], col[k][1]) for k in comunes)
                    nz = sum(1 for k in comunes if col[k][0] != 0)
                    score = (ok, -len(comunes))
                    if ok == len(comunes) and nz >= 2 and ok > mejor[0]:
                        from openpyxl.utils import get_column_letter as L
                        mejor = (ok, f"{hoja} · columna {L(c)} desde la fila {celdas[i][0]} ({ok}/{len(comunes)} prendas, {len(col)} en el informe)")
            print(f"   columna «{polo}»: {', '.join(f'{k[:12]}={v:g}' for k, (v, _) in list(col.items())[:10])}")
            print("   " + ("✔ " + mejor[1] if mejor[1] else "✘ ninguna columna del CD da esos valores por prenda"))
            continue
        fila = fila_del_polo(t["filas"], polo)
        if fila and any(v for v, _ in fila):
            print(f"   fila «{polo}»: {[v for v, _ in fila]}")
            hit = None
            for recorte in (len(fila), len(fila) - 1):
                v = fila[:recorte]
                if len(v) < 2:
                    continue
                for hoja, r, vals in filas_cd:
                    xs = [x for _, x in vals]
                    for k in range(len(xs) - len(v) + 1):
                        if all(igual(xs[k + j], b, d) for j, (b, d) in enumerate(v)):
                            hit = f"{hoja} · fila {r} (primeras {recorte} columnas)"
                            break
                    if hit:
                        break
                if hit:
                    break
            print("   " + ("✔ " + hit if hit else "✘ no está como fila"))
            continue
        print("   (sin datos del polo, o todo a cero)")


if __name__ == "__main__":
    main()
