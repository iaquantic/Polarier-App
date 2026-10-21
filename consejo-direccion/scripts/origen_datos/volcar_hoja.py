"""Vuelca una hoja de un xlsx como rejilla compacta (valores con caché).
Uso: py -3 -P volcar_hoja.py <libro.xlsx> <hoja (subcadena)> [filas_max] [cols_max]"""
import sys

import openpyxl
from openpyxl.utils import get_column_letter as L

ruta, sub = sys.argv[1], sys.argv[2].lower()
maxr = int(sys.argv[3]) if len(sys.argv) > 3 else 80
maxc = int(sys.argv[4]) if len(sys.argv) > 4 else 40
wb = openpyxl.load_workbook(ruta, data_only=True)
ws = next(w for w in wb.worksheets if sub in w.title.lower())
print(f"## {ws.title} ({ws.max_row}x{ws.max_column})")
for r in range(1, min(ws.max_row, maxr) + 1):
    celdas = []
    for c in range(1, min(ws.max_column, maxc) + 1):
        v = ws.cell(r, c).value
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if isinstance(v, float):
            v = f"{v:.2f}".rstrip("0").rstrip(".")
        celdas.append(f"{L(c)}:{str(v).strip()[:22]}")
    if celdas:
        print(r, " | ".join(celdas))
