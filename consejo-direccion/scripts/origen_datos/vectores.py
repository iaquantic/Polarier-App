"""Busca filas completas de una tabla del informe (como secuencia de valores) en los libros.
Para tablas de enteros pequeños o fracciones, donde una celda suelta casa con cualquier cosa.
Uso: py -3 -P vectores.py <tablas.json> <pos1,pos2,...> <libro1.xlsx> [...]"""
import json
import sys

import openpyxl

from cruce import Libro, parsear


def secuencias(lb):
    """Para cada hoja: filas y columnas como listas de (coordenada, valor)."""
    out = []
    for hoja in lb.f.sheetnames:
        ws = lb.f[hoja]
        maxr, maxc = ws.max_row, ws.max_column
        grid = {}
        for r in range(1, maxr + 1):
            for c in range(1, maxc + 1):
                v = lb.valor(hoja, r, c)
                if v is not None:
                    grid[(r, c)] = v
        for r in range(1, maxr + 1):
            out.append((hoja, f"fila {r}", [(c, grid.get((r, c))) for c in range(1, maxc + 1)]))
        for c in range(1, maxc + 1):
            out.append((hoja, f"col {openpyxl.utils.get_column_letter(c)}", [(r, grid.get((r, c))) for r in range(1, maxr + 1)]))
    return out


def casa(seq, vec):
    n = len(vec)
    vals = [v for _, v in seq]
    for i in range(len(vals) - n + 1):
        ok = True
        for j, (x, dec) in enumerate(vec):
            y = vals[i + j]
            if y is None:
                y = 0.0
            if round(y, dec) != round(x, dec) and round(y * 100, dec) != round(x, dec):
                ok = False
                break
        if ok:
            return seq[i][0]
    return None


def main():
    tablas = json.load(open(sys.argv[1], encoding="utf-8"))
    posiciones = {int(p) for p in sys.argv[2].split(",")}
    libros = [Libro(r) for r in sys.argv[3:]]
    seqs = [(lb.nombre, *s) for lb in libros for s in secuencias(lb)]
    for t in tablas:
        if t["pos"] not in posiciones:
            continue
        filas = t["filas"]
        ok = tot = 0
        print(f"\n##### [{t['pos']}] {t['titulo'][:60]}")
        for fila in filas:
            vec = [parsear(c) for c in fila[1:]]
            vec = [(p[0], p[1]) for p in vec if p]
            if len(vec) < 3 or all(v == 0 for v, _ in vec):
                continue
            tot += 1
            hit = None
            # El informe quita los polos cerrados: el total/promedio deja de ir pegado
            for intento, v in (("", vec), (" (sin la última columna)", vec[:-1]), (" (solo los polos)", vec[:-2])):
                if len(v) < 3 or hit:
                    continue
                for libro, hoja, eje, seq in seqs:
                    inicio = casa(seq, v)
                    if inicio is not None:
                        hit = f"{libro[:13]}·{hoja}·{eje} desde {inicio}{intento}"
                        break
            ok += bool(hit)
            print(f"  {fila[0][:28]:<28} {'✔ ' + hit if hit else '✘'}")
        print(f"  → {ok}/{tot} filas encontradas enteras")


if __name__ == "__main__":
    main()
