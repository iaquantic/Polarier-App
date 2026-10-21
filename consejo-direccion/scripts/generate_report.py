"""
Genera el informe mensual o trimestral del Consejo de Dirección de Polarier.

Uso:
    py -3 scripts/generate_report.py

El script detecta automáticamente:
  - El DOCX de entrada en data/input/mes_actual/  (debe haber exactamente uno)
  - El JSON de análisis en data/processed/         (busca analisis_{mes}_{año}.json)
  - El período a partir del nombre del DOCX:
      * Mensual:    "Informe MES tablas.docx"         → mes en español
      * Trimestral: "Informe I Trimestre YYYY.docx"   → "I TRIMESTRE", "II TRIMESTRE", etc.
  - El año del sistema si no se encuentra en el nombre del archivo
  - A qué marcador {TABLA_*} va cada tabla, por su título (config/table_markers.yaml)

Si alguna tabla no se reconoce, se detiene antes de crear el informe y dice cuál.

Ejecutar siempre desde la raíz del proyecto.
"""
import sys
import json
import re
from datetime import datetime
from pathlib import Path

sys.path.insert(0, 'src')

from docx import Document
from template_cloner import clone_master, MASTER_PATH
from docx_content_replacer import replace_metadata
from report_builder import build_report
from table_markers import (
    describe_tables,
    load_rules,
    classify_tables,
    template_table_markers,
)


MESES_ES = {
    "enero": "ENERO", "febrero": "FEBRERO", "marzo": "MARZO",
    "abril": "ABRIL", "mayo": "MAYO", "junio": "JUNIO",
    "julio": "JULIO", "agosto": "AGOSTO", "septiembre": "SEPTIEMBRE",
    "octubre": "OCTUBRE", "noviembre": "NOVIEMBRE", "diciembre": "DICIEMBRE",
}

TRIMESTRES_ES = {
    "i trimestre": "I TRIMESTRE",
    "ii trimestre": "II TRIMESTRE",
    "iii trimestre": "III TRIMESTRE",
    "iv trimestre": "IV TRIMESTRE",
    "1er trimestre": "I TRIMESTRE",
    "2do trimestre": "II TRIMESTRE",
    "3er trimestre": "III TRIMESTRE",
    "4to trimestre": "IV TRIMESTRE",
    "primer trimestre": "I TRIMESTRE",
    "segundo trimestre": "II TRIMESTRE",
    "tercer trimestre": "III TRIMESTRE",
    "cuarto trimestre": "IV TRIMESTRE",
}

INPUT_DIR = Path("data/input/mes_actual")
PROCESSED_DIR = Path("data/processed")


def _find_source_docx() -> Path:
    """Localiza el único DOCX en data/input/mes_actual/."""
    docx_files = list(INPUT_DIR.glob("*.docx"))
    if not docx_files:
        raise FileNotFoundError(
            f"No se encontró ningún DOCX en {INPUT_DIR}\n"
            "Coloque el archivo 'Informe MES tablas.docx' del mes en esa carpeta."
        )
    if len(docx_files) > 1:
        raise ValueError(
            f"Se encontraron {len(docx_files)} archivos DOCX en {INPUT_DIR}:\n"
            + "\n".join(f"  {f.name}" for f in docx_files)
            + "\nDeje solo el archivo del mes actual."
        )
    return docx_files[0]


def _extract_mes_año(docx_path: Path) -> tuple[str, str]:
    """
    Infiere período y año desde el nombre del archivo.
    Patrones soportados:
      - Mensual:    'Informe MES tablas.docx'
      - Trimestral: 'Informe I Trimestre YYYY -tablas.docx'
    Si no se puede extraer el año, usa el año del sistema.
    Retorna (período_upper, año) donde período puede ser un mes o un trimestre.
    """
    stem = docx_path.stem.lower()
    año = str(datetime.now().year)

    # Buscar año de 4 dígitos en el nombre
    year_match = re.search(r'\b(20\d{2})\b', stem)
    if year_match:
        año = year_match.group(1)

    # Buscar trimestre (antes que mes para evitar falsos positivos).
    # Se ordenan las claves por longitud descendente y se exige límite de palabra
    # para evitar que "i trimestre" haga falso match dentro de "ii trimestre"
    # (substring "i trimestre" está contenido literalmente en "ii trimestre").
    for trim_lower in sorted(TRIMESTRES_ES, key=len, reverse=True):
        trim_upper = TRIMESTRES_ES[trim_lower]
        if re.search(r'(?<![a-záéíóúñ])' + re.escape(trim_lower) + r'(?![a-záéíóúñ])', stem):
            return trim_upper, año

    # Buscar nombre de mes en español
    for mes_lower, mes_upper in MESES_ES.items():
        if mes_lower in stem:
            return mes_upper, año

    raise ValueError(
        f"No se pudo determinar el período desde el nombre del archivo: '{docx_path.name}'\n"
        f"El nombre debe contener el mes en español o el trimestre, "
        f"p.ej.: 'Informe JUNIO tablas.docx' o 'Informe I Trimestre 2026 -tablas.docx'"
    )


def _find_analisis_json(mes: str, año: str) -> Path:
    """Busca el JSON de análisis para el período y año dados."""
    # Normalizar: reemplazar espacios por guiones bajos para el nombre de archivo
    mes_slug = mes.lower().replace(" ", "_")
    expected = PROCESSED_DIR / f"analisis_{mes_slug}_{año}.json"
    if expected.exists():
        return expected
    # Búsqueda flexible: cualquier JSON que contenga el período en el nombre
    candidates = list(PROCESSED_DIR.glob(f"analisis_{mes_slug}*.json"))
    if candidates:
        return candidates[0]
    raise FileNotFoundError(
        f"No se encontró el JSON de análisis para {mes} {año}.\n"
        f"Archivo esperado: {expected}\n"
        f"Créelo con los textos de análisis del período antes de ejecutar este script."
    )


def _stop(message: str) -> None:
    print(f"\n[ERROR] {message}", file=sys.stderr)
    sys.exit(1)


# ─── Ejecución ───────────────────────────────────────────────────

print("=" * 60)
print("  Generador de Informe — Consejo de Dirección Polarier")
print("=" * 60)

print("\n1. Detectando archivos del mes...")
source_docx = _find_source_docx()
mes, año = _extract_mes_año(source_docx)
analisis_json = _find_analisis_json(mes, año)
print(f"   DOCX entrada : {source_docx}")
print(f"   Período      : {mes} {año}")
print(f"   JSON análisis: {analisis_json}")

# Las tablas se reconocen antes de tocar nada: si una falla, no queda
# un informe a medias en data/output/.
print("\n2. Reconociendo las tablas del DOCX de entrada...")
tables = describe_tables(Document(str(source_docx)))
classification = classify_tables(tables, load_rules())
print(f"   -> {len(tables)} tablas encontradas")
for a in classification.assignments:
    print(f"      {a.position:>2}  {a.marker:<32} [{a.how}] {a.caption[:50]}")

if not classification.ok:
    _stop(
        "No se pudo asignar alguna tabla a su sitio en el informe:\n"
        + "\n".join(f"   - {e}" for e in classification.errors)
        + "\nRevise el título de esa tabla en el DOCX o añada su patrón en config/table_markers.yaml."
    )

if not MASTER_PATH.exists():
    _stop(f"Plantilla maestra no encontrada: {MASTER_PATH}")
in_template = template_table_markers(Document(str(MASTER_PATH)))
missing = sorted({a.marker for a in classification.assignments} - in_template)
if missing:
    _stop(
        "Estos marcadores no existen en la plantilla maestra y sus tablas se perderían: "
        + ", ".join(missing)
        + f"\nAñádalos a {MASTER_PATH} o corrija config/table_markers.yaml."
    )

source_xml_map = classification.marker_to_xml(tables)
print(f"   -> {len(source_xml_map)} marcadores con tabla")

print("\n3. Clonando plantilla maestra...")
out_path = clone_master(mes=mes, año=año)
print(f"   -> {out_path}")

print("\n4. Reemplazando metadatos ({MES}, {AÑO})...")
replace_metadata(out_path, mes=mes, año=año)

print("\n5. Cargando textos de análisis...")
with open(analisis_json, encoding="utf-8") as f:
    data = json.load(f)
sections = data["sections"]

analysis_texts = {
    k: v for k, v in sections.items()
    if k.startswith("ANALISIS_") and v.strip()
    and not re.match(r'^\[dato no disponible\]', v.strip(), re.IGNORECASE)
}

for marker in source_xml_map:
    analisis_key = marker.replace("TABLA_", "ANALISIS_")
    if analisis_key not in analysis_texts:
        print(f"   [AVISO] Sin análisis para {analisis_key} — se omitirá la subsección")

print(f"   -> {len(analysis_texts)} análisis cargados")

print("\n6. Construyendo informe...")
section_texts = {
    k: v for k, v in sections.items()
    if k.startswith("SECCION_") or k.startswith("INTRO_")
}
result = build_report(
    output_path=out_path,
    source_xml_map=source_xml_map,
    analysis_texts=analysis_texts,
    section_texts=section_texts,
)
print(f"   -> tablas insertadas : {result['tables_inserted']}")
print(f"   -> tablas eliminadas : {result['tables_removed']}")
print(f"   -> análisis insertados: {result['analyses_inserted']}")

print("\n7. Verificando marcadores pendientes...")
doc_out = Document(str(out_path))
remaining = [p.text.strip() for p in doc_out.paragraphs if "{" in p.text and "}" in p.text]
if remaining:
    print(f"   [AVISO] {len(remaining)} marcador(es) sin reemplazar:")
    for m in remaining:
        print(f"      {m}")
else:
    print("   -> OK, ningún marcador pendiente")

print(f"\n{'=' * 60}")
print(f"  LISTO -> {out_path}")
print(f"{'=' * 60}\n")
