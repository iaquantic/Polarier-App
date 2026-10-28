"""
Prepara las tablas de un informe del Consejo a partir de la TABLA CA y el GLOBAL.

Sustituye al DOCX de tablas que se montaba a mano: escribe uno igual en
trabajo/<período>/tablas.docx, rellenando cada tabla sobre la misma tabla del
último informe de ejemplo (ejemplos/entrada/), y saca por stdout un JSON con las
tablas en texto para quien redacta los análisis (la routine de Claude, vía n8n).

Uso (desde cualquier carpeta; el script trabaja en la raíz del proyecto):
    python scripts/preparar_consejo.py --anio 2026 --mes 6
    python scripts/preparar_consejo.py --anio 2026 --trimestre 2
Opciones:
    --tabla RUTA     TABLA CA <año>.xlsx  (por defecto, la del VPS: ../CD/TABLA CA <año>.xlsx)
    --global RUTA    GLOBAL-CP<aa>.xlsx   (por defecto, ../GLOBAL-CP<aa>.xlsx)
    --sin-global     no abrir el GLOBAL: los vínculos de la TABLA CA salen de su propia caché
                     (para probar en local con una TABLA CA guardada por Excel)
    --ejemplo RUTA   DOCX de tablas que sirve de prototipo (por defecto, el más reciente
                     del mismo tipo en ejemplos/entrada/)
    --lenceria RUTA  JSON con las filas de cd_lenceria del período y del mismo período del
                     año anterior (lo que devuelve lenceria_informe_consejo en Supabase).
                     Con él, las ocho tablas de lencería salen del CD; sin él, de las hojas
                     tecleadas de la TABLA CA
    --lenceria-b64 TROZOS…  lo mismo en base64, partido en trozos (desde n8n)

Salida: {"ok": true, ...} o {"ok": false, "error": "..."}; siempre con código 0 para
que n8n pueda leer el motivo.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import threading
import traceback
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
os.chdir(RAIZ)
sys.stdout.reconfigure(encoding="utf-8")  # el JSON va a n8n en UTF-8 también desde Windows
sys.path.insert(0, str(RAIZ / "src"))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from constructores_tablas import CONSTRUCTORES, SinFuente  # noqa: E402
from excel_eval import Libro  # noqa: E402
from lenceria_cd import Lenceria  # noqa: E402
from table_markers import classify_tables, describe_tables, load_rules, template_table_markers  # noqa: E402
from tablas_informe import Fuentes, Periodo, celdas, filas, texto_celda  # noqa: E402

EJEMPLOS = RAIZ / "ejemplos" / "entrada"
ANALISIS_EJEMPLO = RAIZ / "ejemplos" / "analisis"
MESES_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
            "septiembre", "octubre", "noviembre", "diciembre"]
ROMANOS = ["i", "ii", "iii", "iv"]
TRABAJO = RAIZ / "trabajo"
MASTER = RAIZ / "templates" / "master" / "consejo_direccion_master.docx"
ESTILOS_TITULO = re.compile(r"^(heading|t[ií]tulo)", re.I)


def salir(obj) -> None:
    print(json.dumps(obj, ensure_ascii=False))
    sys.stdout.flush()


def periodo_de_nombre(nombre: str):
    """'tablas_2026_06.docx' → (2026, 6, None); 'tablas_2026_T2.docx' → (2026, None, 2)."""
    m = re.search(r"(20\d\d)[_-](?:(\d{1,2})|T([1-4]))\b", nombre, re.I)
    if not m:
        return None
    return int(m.group(1)), (int(m.group(2)) if m.group(2) else None), (int(m.group(3)) if m.group(3) else None)


def elegir_ejemplo(per: Periodo, forzado: str | None) -> Path:
    if forzado:
        return Path(forzado)
    candidatos = []
    for f in EJEMPLOS.glob("*.docx"):
        p = periodo_de_nombre(f.name)
        if not p:
            continue
        anio, mes, tri = p
        if (tri is not None) != per.es_trimestre:
            continue
        orden = (anio, tri or 0, mes or 0)
        candidatos.append((orden, f))
    if not candidatos:
        raise FileNotFoundError(f"no hay DOCX de tablas de ejemplo {'trimestral' if per.es_trimestre else 'mensual'} "
                                f"en {EJEMPLOS} (nombre esperado: tablas_AAAA_MM.docx o tablas_AAAA_Tn.docx)")
    return max(candidatos)[1]


def analisis_de_ejemplo(per: Periodo):
    """
    El último análisis aprobado del mismo tipo (mensual o trimestral), anterior al
    período, para que quien redacta vea el tono y la extensión. Nombres:
    analisis_<mes>_<año>.json o analisis_<romano>_trimestre_<año>.json.
    """
    candidatos = []
    for f in ANALISIS_EJEMPLO.glob("analisis_*.json"):
        n = f.stem.lower()
        m = re.fullmatch(r"analisis_(i|ii|iii|iv)_trimestre_(20\d\d)", n)
        if m and per.es_trimestre:
            orden = (int(m.group(2)), ROMANOS.index(m.group(1)) + 1)
        else:
            m = re.fullmatch(r"analisis_([a-z]+)_(20\d\d)", n)
            if not m or per.es_trimestre or m.group(1) not in MESES_ES:
                continue
            orden = (int(m.group(2)), MESES_ES.index(m.group(1)) + 1)
        actual = (per.anio, per.trimestre if per.es_trimestre else per.mes)
        if orden < actual:
            candidatos.append((orden, f))
    if not candidatos:
        return None
    f = max(candidatos)[1]
    datos = json.loads(f.read_text(encoding="utf-8"))
    return {"archivo": f.name, "sections": datos.get("sections", datos)}


def texto_tabla(tbl) -> list:
    return [[texto_celda(tc) for tc in celdas(tr)] for tr in filas(tbl)]


def sha(ruta: Path) -> str:
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


def fuente_info(ruta: Path) -> dict:
    st = ruta.stat()
    return {"archivo": ruta.name, "modificado": datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat(),
            "sha256": sha(ruta)}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--anio", type=int, required=True)
    ap.add_argument("--mes", type=int)
    ap.add_argument("--trimestre", type=int)
    ap.add_argument("--tabla")
    ap.add_argument("--global", dest="global_")
    ap.add_argument("--sin-global", action="store_true")
    ap.add_argument("--ejemplo")
    ap.add_argument("--lenceria")
    ap.add_argument("--lenceria-b64", nargs="+")
    a = ap.parse_args()
    if bool(a.mes) == bool(a.trimestre):
        raise ValueError("indica --mes o --trimestre (uno de los dos)")
    per = Periodo(a.anio, mes=a.mes, trimestre=a.trimestre)

    leslie = RAIZ.parent
    ruta_tabla = Path(a.tabla) if a.tabla else leslie / "CD" / f"TABLA CA {a.anio}.xlsx"
    ruta_global = Path(a.global_) if a.global_ else leslie / f"GLOBAL-CP{str(a.anio)[2:]}.xlsx"
    if not ruta_tabla.exists():
        raise FileNotFoundError(f"no existe {ruta_tabla}")
    G = None
    if not a.sin_global:
        if not ruta_global.exists():
            raise FileNotFoundError(f"no existe {ruta_global}")
        G = Libro(ruta_global)
    externos = {ruta_global.name: G} if G is not None else {}
    L = Libro(ruta_tabla, externos=externos)
    fu = Fuentes(L, G)
    if a.lenceria:
        fu.lenceria = Lenceria(json.loads(Path(a.lenceria).read_text(encoding="utf-8")))
    elif a.lenceria_b64:
        import base64
        fu.lenceria = Lenceria(json.loads(base64.b64decode("".join(a.lenceria_b64)).decode("utf-8")))

    ejemplo = elegir_ejemplo(per, a.ejemplo)
    p_ej = periodo_de_nombre(ejemplo.name)
    anio_proto = p_ej[0] if p_ej else per.anio
    doc = Document(str(ejemplo))
    tablas = describe_tables(doc)
    clas = classify_tables(tablas, load_rules())
    if not clas.ok:
        raise ValueError("el ejemplo tiene tablas sin reconocer: " + "; ".join(clas.errors))
    por_pos = {t.position: t for t in tablas}
    protos = {}
    for asg in clas.assignments:
        protos.setdefault(asg.marker, []).append(por_pos[asg.position].xml)

    rellenas, avisos, sin_fuente, salida_tablas = {}, {}, [], []
    for marcador, lista in protos.items():
        fn = CONSTRUCTORES.get(marcador)
        if fn is None:
            sin_fuente.append({"marcador": marcador, "motivo": "no hay constructor para esta tabla"})
            continue
        try:
            res = fn(fu, per, lista, anio_proto)
        except SinFuente as e:
            sin_fuente.append({"marcador": marcador, "motivo": str(e)})
            continue
        rellenas[marcador] = [r.xml for r in res]
        av = sorted({x for r in res for x in r.avisos})
        if av:
            avisos[marcador] = av
        salida_tablas.append({
            "marcador": marcador,
            "tecleado": any(x.startswith("tecleado sin mes") for x in av),
            "tablas": [texto_tabla(r.xml) for r in res],
            "avisos": av,
        })

    # DOCX de tablas: el del ejemplo, con sus tablas cambiadas por las rellenas, sin las
    # tablas que no tienen fuente y sin sus notas (eran de otro período; se quedan los títulos)
    cuerpo = doc.element.body
    usados = {}
    for asg in clas.assignments:
        tbl = por_pos[asg.position].xml
        lista = rellenas.get(asg.marker)
        i = usados.get(asg.marker, 0)
        usados[asg.marker] = i + 1
        if lista is None or i >= len(lista):
            cuerpo.remove(tbl)
        else:
            tbl.addprevious(lista[i])
            cuerpo.remove(tbl)
    for p in list(cuerpo.iterchildren(qn("w:p"))):
        estilo = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        nombre_estilo = estilo.get(qn("w:val")) if estilo is not None else ""
        texto = "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()
        if texto and not ESTILOS_TITULO.match(nombre_estilo or ""):
            cuerpo.remove(p)
    ident = f"{per.anio}-T{per.trimestre}" if per.es_trimestre else f"{per.anio}-{per.mes:02d}"
    carpeta = TRABAJO / ident
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta_docx = carpeta / "tablas.docx"
    doc.save(str(ruta_docx))

    # Comprobación: el DOCX resultante tiene que reconocerse igual que el del ejemplo
    comprob = classify_tables(describe_tables(Document(str(ruta_docx))), load_rules())
    if not comprob.ok:
        raise ValueError("el DOCX de tablas generado no se reconoce: " + "; ".join(comprob.errors))

    # Textos que tendrá que redactar la routine: los de la plantilla que tienen tabla, más
    # las introducciones y las secciones finales
    marcadores_plantilla = []
    import re as _re
    for par in Document(str(MASTER)).paragraphs:
        marcadores_plantilla += _re.findall(r"\{([A-Z_]+)\}", par.text)
    con_tabla = {m for m in rellenas}
    textos = []
    for m in marcadores_plantilla:
        if m.startswith("TABLA_"):
            continue
        if m.startswith("ANALISIS_"):
            if "TABLA_" + m[len("ANALISIS_"):] in con_tabla:
                textos.append(m)
        elif m.startswith(("INTRO_", "SECCION_")):
            textos.append(m)

    usos = sum(getattr(x, "usos_respaldo", 0) for x in L._ext.values())
    salir({
        "ok": True,
        "periodo": {"anio": per.anio, "mes": per.mes, "trimestre": per.trimestre, "nombre": per.nombre,
                    "id": ident},
        "ejemplo": ejemplo.name,
        "archivo_tablas": str(ruta_docx.relative_to(RAIZ.parents[1])) if RAIZ.parents[1] in ruta_docx.parents else str(ruta_docx),
        "fuentes": {"tabla_ca": fuente_info(ruta_tabla), "global": fuente_info(ruta_global) if G is not None else None,
                    "celdas_del_global_desde_la_cache_de_la_tabla_ca": usos,
                    "lenceria": ({"origen": "cd_lenceria", "filas": fu.lenceria.n} if fu.lenceria is not None
                                 else {"origen": "hojas tecleadas de la TABLA CA"})},
        "tablas": salida_tablas,
        "sin_fuente": sin_fuente,
        "textos": textos,
        "ejemplo_analisis": analisis_de_ejemplo(per),
    })


def _correr():
    try:
        main()
    except Exception as e:  # el motivo tiene que llegar a n8n
        salir({"ok": False, "error": f"{type(e).__name__}: {e}", "traza": traceback.format_exc()[-2000:]})


if __name__ == "__main__":
    threading.stack_size(200 * 1024 * 1024)  # las cadenas de fórmulas son profundas
    hilo = threading.Thread(target=_correr)
    hilo.start()
    hilo.join()
