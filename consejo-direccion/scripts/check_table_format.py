import sys
sys.path.insert(0, 'src')
from docx import Document
from docx.oxml.ns import qn

WTCPR = qn('w:tcPr')
WSHD  = qn('w:shd')
WFILL = qn('w:fill')
WCOLOR = qn('w:color')

def table_info(doc_path, table_idx, label):
    doc = Document(doc_path)
    t = doc.tables[table_idx]
    print(f'=== {label} - Table {table_idx} ===')
    print(f'  Style: {t.style.name if t.style else None}')
    print(f'  Rows: {len(t.rows)}, Cols: {len(t.columns)}')
    for ri in range(min(2, len(t.rows))):
        for ci in range(min(3, len(t.columns))):
            cell = t.rows[ri].cells[ci]
            para = cell.paragraphs[0] if cell.paragraphs else None
            font_name = None
            font_size = None
            bold = None
            color = None
            if para and para.runs:
                run = para.runs[0]
                font_name = run.font.name
                font_size = run.font.size
                bold = run.font.bold
                color = run.font.color.rgb if run.font.color and run.font.color.type else None
            tcPr = cell._tc.find(WTCPR)
            fill = None
            if tcPr is not None:
                shd = tcPr.find(WSHD)
                if shd is not None:
                    fill = shd.get(WFILL)
            print(f'  [{ri},{ci}] text={repr(cell.text[:30]):35s} font={font_name} sz={font_size} bold={bold} fill={fill} color={color}')
    print()

table_info('data/input/mes_actual/Informe MARZO tablas.docx', 0, 'INPUT')
table_info('data/input/mes_actual/Informe MARZO tablas.docx', 1, 'INPUT')
table_info('data/output/consejo_direccion_MARZO_2026.docx',   0, 'OUTPUT')
table_info('data/output/consejo_direccion_MARZO_2026.docx',   1, 'OUTPUT')
