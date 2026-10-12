"""
Normaliza el GLOBAL (GLOBAL-CP<aa>.xlsx) para cargarlo en Supabase (global_ai).

No escribe nada: imprime por stdout un JSON con una fila por hoja de polo ×
partida × mes. n8n lo recoge y llama a la función sync_global_ai() de Supabase,
que reemplaza las hojas recibidas en una sola transacción (migración 010).

Reglas (ver vault/Generacion-Documentos.md, D6):
  - Fichero: el GLOBAL-CP*.xlsx modificado más recientemente en la carpeta del
    script (el que Modelo1.py acaba de escribir), o el que se pase como argumento.
    El año sale del nombre de la hoja (<COD>CP<AAAA>), no va fijo.
  - Hojas: solo las de polo, <COD>CP<AAAA> con COD en LH, VAR, CA, CC, CCZ, HOL.
  - Meses: se localizan por su nombre en la fila 4; PLAN y REAL en la fila 5.
  - Fórmulas: Modelo1.py guarda con openpyxl y eso borra los valores cacheados,
    así que las fórmulas (sumas, restas, SUM de rangos de la misma hoja) se
    evalúan aquí. Para decidir si un mes está cargado, una fórmula cuyas celdas
    están todas vacías cuenta como vacía (si no, un mes sin cargar parecería un
    mes con ventas cero). Dentro de un mes cargado se evalúan como Excel: un
    SUM sobre celdas vacías es 0, igual que hacía el importador original.
  - Mes cargado: un mes cuenta para una hoja solo si "Ventas Totales" (código 1)
    tiene PLAN o REAL. Los meses futuros de la plantilla traen algún 0 escrito a
    mano que no es un dato; sin esta regla aparecerían como meses con ceros.
  - Fila: dentro de un mes cargado, se emite si la celda de PLAN o la de REAL no
    está en blanco (valor o fórmula). La que esté en blanco va como null.

Uso:
  python sync_global_ai.py [ruta/GLOBAL-CP26.xlsx]
"""

import glob
import hashlib
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string

MESES = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9, "OCTUBRE": 10,
    "NOVIEMBRE": 11, "DICIEMBRE": 12,
}
NOMBRE_MES = {
    1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril", 5: "Mayo", 6: "Junio",
    7: "Julio", 8: "Agosto", 9: "Septiembre", 10: "Octubre", 11: "Noviembre",
    12: "Diciembre",
}
# Código de hoja -> (polo, polo_key). Mismos valores que ya usa global_ai.
POLOS = {
    "LH": ("Habana", "habana"),
    "VAR": ("Varadero", "varadero"),
    "CA": ("Caibarien", "caibarien"),
    "CC": ("Cayo Coco", "cayo_coco"),
    "CCZ": ("Cayo Cruz", "cayo_cruz"),
    "HOL": ("Holguin", "holguin"),
}
CODIGO_MES_CARGADO = 1  # Ventas Totales
HOJA_POLO = re.compile(r"^(LH|VAR|CCZ|CC|CA|HOL)CP(\d{4})$")
FILA_MESES = 4
FILA_PLAN_REAL = 5
REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
RANGO_SUM = re.compile(r"SUM\(\s*(\$?[A-Z]{1,3}\$?\d+)\s*:\s*(\$?[A-Z]{1,3}\$?\d+)\s*\)", re.I)
SOLO_ARITMETICA = re.compile(r"[0-9eE.+\-*/() ]*")


def elegir_fichero() -> str:
    if len(sys.argv) > 1:
        return sys.argv[1]
    carpeta = os.path.dirname(os.path.abspath(__file__))
    candidatos = glob.glob(os.path.join(carpeta, "GLOBAL-CP*.xlsx"))
    if not candidatos:
        raise FileNotFoundError(f"No hay ningún GLOBAL-CP*.xlsx en {carpeta}")
    return max(candidatos, key=os.path.getmtime)


class Evaluador:
    """Evalúa las celdas de una hoja resolviendo sus fórmulas aritméticas.

    Devuelve None para una celda vacía y para una fórmula cuyas referencias
    están todas vacías.
    """

    def __init__(self, hoja):
        self.hoja = hoja
        self.memo = {}

    def valor(self, fila: int, col: int):
        clave = (fila, col)
        if clave in self.memo:
            return self.memo[clave]
        self.memo[clave] = None  # corta ciclos
        v = self.hoja.cell(fila, col).value
        if isinstance(v, str) and v.startswith("="):
            res = self._formula(v)
        elif isinstance(v, bool):
            res = None
        elif isinstance(v, (int, float)):
            res = float(v)
        else:
            res = None
        self.memo[clave] = res
        return res

    def excel(self, fila: int, col: int):
        """Como Excel: celda en blanco -> None; fórmula sin datos -> 0."""
        v = self.hoja.cell(fila, col).value
        if v is None or (isinstance(v, str) and not v.startswith("=")):
            return None
        x = self.valor(fila, col)
        return 0.0 if x is None else x

    def _ref(self, texto: str):
        m = REF.fullmatch(texto.replace(" ", ""))
        return int(m.group(2)), column_index_from_string(m.group(1))

    def _formula(self, formula: str):
        expr = formula[1:]
        hay_dato = False

        def suma_rango(m):
            nonlocal hay_dato
            f1, c1 = self._ref(m.group(1))
            f2, c2 = self._ref(m.group(2))
            total = 0.0
            for f in range(min(f1, f2), max(f1, f2) + 1):
                for c in range(min(c1, c2), max(c1, c2) + 1):
                    x = self.valor(f, c)
                    if x is not None:
                        hay_dato = True
                        total += x
            return repr(total)

        def celda(m):
            nonlocal hay_dato
            x = self.valor(int(m.group(2)), column_index_from_string(m.group(1)))
            if x is None:
                return "0.0"
            hay_dato = True
            return repr(x)

        expr = RANGO_SUM.sub(suma_rango, expr)
        expr = REF.sub(celda, expr)
        if not SOLO_ARITMETICA.fullmatch(expr):
            raise ValueError(f"Fórmula no soportada: {formula}")
        if not hay_dato:
            return None
        try:
            return float(eval(expr, {"__builtins__": {}}, {}))  # solo aritmética, validada arriba
        except ZeroDivisionError:
            return None


def columnas_meses(hoja):
    """{mes: (col_plan, col_real, col_qp)} a partir de las filas 4 y 5."""
    meses = {}
    for col in range(1, hoja.max_column + 1):
        nombre = hoja.cell(FILA_MESES, col).value
        if not isinstance(nombre, str):
            continue
        mes = MESES.get(nombre.strip().upper())
        if mes is None:
            continue
        cols = {}
        for c in range(col, col + 3):
            # Algunas hojas rotulan "PLAN (Qp)" en vez de "PLAN".
            etiqueta = str(hoja.cell(FILA_PLAN_REAL, c).value or "").strip().upper()
            for clave in ("PLAN", "REAL", "QP"):
                if etiqueta.startswith(clave) and clave not in cols:
                    cols[clave] = c
                    break
        if "PLAN" in cols and "REAL" in cols:
            meses[mes] = (cols["PLAN"], cols["REAL"], cols.get("QP"))
    return meses


def letra(col: int) -> str:
    return openpyxl.utils.get_column_letter(col)


def main():
    ruta = elegir_fichero()
    source_file = os.path.splitext(os.path.basename(ruta))[0]
    libro = openpyxl.load_workbook(ruta, data_only=False)

    filas = []
    resumen = {}
    anios = set()
    for nombre in libro.sheetnames:
        m = HOJA_POLO.match(nombre)
        if not m:
            continue
        cod, anio = m.group(1), int(m.group(2))
        anios.add(anio)
        polo, polo_key = POLOS[cod]
        hoja = libro[nombre]
        ev = Evaluador(hoja)
        meses = columnas_meses(hoja)
        if len(meses) != 12:
            raise ValueError(f"{nombre}: se esperaban 12 meses en la fila {FILA_MESES}, hay {sorted(meses)}")
        filas_codigo = []
        for fila in range(FILA_PLAN_REAL + 1, hoja.max_row + 1):
            codigo = hoja.cell(fila, 1).value
            if isinstance(codigo, (int, float)) and not isinstance(codigo, bool) and hoja.cell(fila, 2).value is not None:
                filas_codigo.append((fila, int(codigo)))
        fila_control = next((f for f, c in filas_codigo if c == CODIGO_MES_CARGADO), None)
        if fila_control is None:
            raise ValueError(f"{nombre}: no encuentro la fila con código {CODIGO_MES_CARGADO} (Ventas Totales)")
        cargados = {
            mes: cols for mes, cols in meses.items()
            if ev.valor(fila_control, cols[0]) is not None or ev.valor(fila_control, cols[1]) is not None
        }
        n_hoja = 0
        meses_con_real = set()
        for fila, codigo in filas_codigo:
            partida_raw = hoja.cell(fila, 2).value
            for mes, (c_plan, c_real, c_qp) in sorted(cargados.items()):
                plan = ev.excel(fila, c_plan)
                real = ev.excel(fila, c_real)
                if plan is None and real is None:
                    continue
                if real is not None:
                    meses_con_real.add(mes)
                filas.append({
                    "hoja": nombre,
                    "anio": anio,
                    "polo": polo,
                    "polo_key": polo_key,
                    "mes": mes,
                    "periodo": NOMBRE_MES[mes],
                    "period_key": f"{anio}-{mes:02d}",
                    "codigo_fila": codigo,
                    "partida_raw": str(partida_raw),
                    "plan": plan,
                    "real": real,
                    "cell_plan": f"{nombre}!{letra(c_plan)}{fila}",
                    "cell_real": f"{nombre}!{letra(c_real)}{fila}",
                    "cell_qp": f"{nombre}!{letra(c_qp)}{fila}" if c_qp else None,
                    "fila_excel": fila,
                })
                n_hoja += 1
        resumen[nombre] = {
            "filas": n_hoja,
            "meses_cargados": sorted(cargados),
            "meses_con_real": sorted(meses_con_real),
        }

    if len(anios) != 1:
        raise ValueError(f"Las hojas de polo mezclan años: {sorted(anios)}")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    json.dump(
        {
            "source_file": source_file,
            "anio": anios.pop(),
            "script_sha256": hashlib.sha256(open(__file__, "rb").read()).hexdigest(),
            "resumen": resumen,
            "filas": filas,
        },
        sys.stdout,
        ensure_ascii=False,
        separators=(",", ":"),
    )


if __name__ == "__main__":
    main()
