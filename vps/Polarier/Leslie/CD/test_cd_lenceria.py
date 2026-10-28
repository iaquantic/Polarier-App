"""Tests de cd_lenceria.py con CD inventados (los dos modelos que se han visto).

    cd vps/Polarier/Leslie/CD && python -m pytest test_cd_lenceria.py -q

Todos los números son inventados: el repo es público.
"""

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import cd_lenceria as cl  # noqa: E402

NOMBRES = ["Sabana Personal", "Sabana King", "Funda Almohada", "Toalla Alfombra", "Toalla Baño",
           "Toalla Cara", "Toalla Facial", "Toalla Bidet", "Toalla Piscina"]


def poner(ws, fila, col, valores):
    for i, v in enumerate(valores):
        ws.cell(row=fila, column=col + i, value=v)


def bloque_hoteles(ws, fila, titulo, datos, suma=None):
    """Título, cabecera «Lencería | H1 | H2 | SUMA» y una fila por prenda. datos: {prenda: (h1, h2)}."""
    poner(ws, fila, 1, [titulo])
    poner(ws, fila + 1, 1, ["Lencería", "HOTEL UNO", "HOTEL DOS", "SUMA"])
    for i, nombre in enumerate(NOMBRES):
        h1, h2 = datos.get(nombre, (None, None))
        s = suma if suma is not None else (h1 or 0) + (h2 or 0)
        poner(ws, fila + 2 + i, 1, [nombre, h1, h2, s])
    poner(ws, fila + 2 + len(NOMBRES), 1, ["TOTAL"])


def cd_modelo_general(ruta):
    """Como los CD de Varadero y Caibarién: «5.Lenceria 1» y «Lencería 2»."""
    wb = Workbook()
    l1 = wb.active
    l1.title = "5.Lenceria 1"
    poner(l1, 1, 1, ["ESTADO DE LAS MONTAS POR HABITACIÓN"])
    poner(l1, 2, 1, ["Lencería", "H1", "H2", "TOTAL"])
    for i, nombre in enumerate(NOMBRES):
        poner(l1, 3 + i, 1, [nombre, 1.0, 2.0, 1.5 + i / 10])
    poner(l1, 19, 1, ["ESTADO DE LAS MONTAS POR CLIENTE"])
    poner(l1, 20, 1, ["Lencería", "H1", "H2", "TOTAL"])
    poner(l1, 21, 1, ["TD TOTALES", 0, 900, 900])
    poner(l1, 22, 1, ["TD PROMEDIO", 0, 30, 30])
    for i, nombre in enumerate(NOMBRES):  # el total da #DIV/0! y solo H2 tiene datos
        poner(l1, 23 + i, 1, [nombre, "#DIV/0!", 5.0 + i, "#DIV/0!"])

    l2 = wb.create_sheet("Lencería 2")
    bloque_hoteles(l2, 1, "RESUMEN DE LOS RECHAZOS POR TIPO DE PRENDA",
                   {"Sabana Personal": (10, 5), "Funda Almohada": (0, 7)})
    bloque_hoteles(l2, 14, "RESUMEN DEL MAL USO DETECTADO POR HOTELES", {"Toalla Baño": (3, 4)})
    bloque_hoteles(l2, 27, "RESUMEN DE LAS INYECCIONES POR TIPO DE PRENDA",
                   {n: (0, 0) for n in NOMBRES})
    # Desmanche con los códigos del modelo «cuerpo» (TC = baño, TM = cara, TB = bidet)
    poner(l2, 1, 8, ["RECUPERACIONES DE LENCERÍA"])
    poner(l2, 2, 8, ["Prenda", "Producción", "Recuperado", "% Recuperación"])
    for i, (cod, trat, rec) in enumerate([("SP", 100, 90), ("SK", 50, 40), ("F", 0, 0),
                                          ("ALF", 0, 0), ("TC", 200, 150), ("TM", 80, 60),
                                          ("TF", 0, 0), ("TB", 0, 0), ("TP", 10, 10)]):
        poner(l2, 3 + i, 8, [cod, trat, rec, None])
    poner(l2, 12, 8, ["Total"])
    # Descartes: «en el mes» con negativos y el mes en la columna rotulada 2025
    poner(l2, 1, 14, ["DESCARTES - INVENTARIOS DE DESCARTE"])
    poner(l2, 2, 14, ["Prenda", "Inventario Inicial", "Descarte generado en el mes",
                      "Descarte generado   2025", "Ventas"])
    for i, (cod, mes, otra) in enumerate([("SP", 0, 12), ("SK", -30, 4), ("F", 0, 0), ("TA", -5, 0),
                                          ("TB", 0, 7), ("TC", 0, 0), ("TF", 0, 0), ("TP", 0, 3)]):
        poner(l2, 3 + i, 14, [cod, 100, mes, otra, 0])
    poner(l2, 11, 14, ["Total"])
    # Inventarios: solo la diferencia tecleada (Inicial y Real en blanco)
    poner(l2, 16, 20, ["Resultados de los inventarios realizados"])
    poner(l2, 17, 20, ["MAYO", None, "SP", "SK", "F", "ALF", "TC", "TM", "TF", "TB", "TP"])
    grupos = [("HOTEL UNO", [-7, 3, 0, 0, 0, 0, 0, 0, -2]), (None, [0] * 9)]
    fila = 18
    for nombre, difs in grupos:
        poner(l2, fila, 20, [nombre, "Inicial"])
        poner(l2, fila + 1, 21, ["Real"])
        poner(l2, fila + 2, 21, ["Diferencia"] + difs)
        poner(l2, fila + 3, 20, ["Importe", 12.5])
        fila += 4
    wb.save(ruta)


def cd_modelo_habana(ruta):
    """Como el CD de La Habana: hojas con sufijo, otros títulos y bloques sin rellenar."""
    wb = Workbook()
    l1 = wb.active
    l1.title = "5.Lenceria 1 (2)"
    poner(l1, 1, 1, ["ESTADO DE LAS MONTAS POR HABITACIÓN"])
    poner(l1, 2, 1, ["Lencería", "H1", "TOTAL"])
    for i, nombre in enumerate(NOMBRES):
        poner(l1, 3 + i, 1, [nombre, 1.1, 1.1])
    # Inventarios en la hoja 1, con nombre de hotel y Real
    poner(l1, 20, 10, ["MAYO", None, "SP", "SK", "F", "TA", "TB", "TC", "TF", "TP"])
    poner(l1, 21, 10, ["HOTEL TRES", "Inicial", 10, 10, 10, 10, 10, 10, 10, 10])
    poner(l1, 22, 11, ["Real", 8, 12, 10, 9, 10, 10, 10, 10])
    poner(l1, 23, 11, ["Diferencia", -2, 2, 0, -1, 0, 0, 0, 0])
    poner(l1, 24, 10, ["Importe", 40])

    l2 = wb.create_sheet("Lencería 2 MAYO")
    bloque_hoteles(l2, 1, "RESUMEN DE LOS RECHAZOS DE LENCERIA MENSUAL", {}, suma=0)
    bloque_hoteles(l2, 14, "RESUMEN DE LOS RECHAZOS POR TIPO DE PRENDA", {}, suma=0)
    bloque_hoteles(l2, 27, "RESUMEN DEL MAL USO DETECTADO POR HOTELES", {}, suma=0)
    bloque_hoteles(l2, 40, "RESUMEN DEL DESCARTE  POR HOTELES", {"Sabana King": (2, 1)})
    bloque_hoteles(l2, 53, "RESUMEN INYECCIONES  POR HOTELES", {"Funda Almohada": (100, 50)})
    # Desmanche partido en «sin remojo» y «con remojo» (códigos del modelo «baño»), y
    # «mal uso hotel», que no es desmanche
    for col, titulo, vals in [(8, "RECUPERACIONES DE LENCERÍA sin remojo", (20, 18)),
                              (13, "RECUPERACIONES DE LENCERÍA con remojo", (10, 6)),
                              (18, "RECUPERACIONES DE LENCERÍA MAL USO HOTEL", (999, 999))]:
        poner(l2, 1, col, [titulo])
        poner(l2, 2, col, ["Prenda", "Producción", "Recuperado", "% Recuperación"])
        for i, cod in enumerate(["SP", "SK", "F", "TA", "TB", "TC", "TF", "TP"]):
            poner(l2, 3 + i, col, [cod] + (list(vals) if cod == "TB" else [None, None]) + ["#DIV/0!"])
        poner(l2, 11, col, ["Total"])
    wb.save(ruta)


@pytest.fixture
def general(tmp_path):
    ruta = tmp_path / "CACD203001.xlsx"
    cd_modelo_general(ruta)
    return cl.leer_cd(ruta, "MAYO")


@pytest.fixture
def habana(tmp_path):
    ruta = tmp_path / "HACA203001.xlsx"
    cd_modelo_habana(ruta)
    return cl.leer_cd(ruta, "MAYO")


def test_salen_los_ocho_bloques(general):
    assert list(general) == cl.BLOQUES


def test_montas_por_habitacion_columna_total(general):
    b = general["montas_habitacion"]
    assert b["estado"] == "ok"
    assert b["valores"]["sabana_personal"]["indice"] == 1.5
    assert b["valores"]["toalla_piscina"]["indice"] == pytest.approx(2.3)


def test_montas_por_cliente_con_total_roto_usa_el_unico_hotel(general):
    b = general["montas_cliente"]
    assert b["estado"] == "ok"
    assert b["valores"]["sabana_personal"]["indice"] == 5
    assert b["totales"]["td_totales"] == 900
    assert any("solo un hotel" in a for a in b["avisos"])


def test_bloques_por_hoteles_leen_la_suma(general):
    assert general["rechazo"]["valores"]["sabana_personal"]["cantidad"] == 15
    assert general["rechazo"]["valores"]["funda"]["cantidad"] == 7
    assert general["mal_uso"]["valores"]["toalla_bano"]["cantidad"] == 7


def test_ceros_escritos_no_son_vacio(general):
    b = general["inyecciones"]
    assert b["estado"] == "ok"
    assert all(v["cantidad"] == 0 for v in b["valores"].values())


def test_desmanche_con_codigos_del_modelo_cuerpo(general):
    v = general["desmanche"]["valores"]
    assert v["toalla_bano"] == {"tratadas": 200, "recuperadas": 150, "pct": 0.75}
    assert v["toalla_cara"]["tratadas"] == 80
    assert v["toalla_bidet"]["tratadas"] == 0
    assert v["funda"]["pct"] is None


def test_retiro_usa_la_columna_sin_negativos(general):
    b = general["retiro"]
    assert b["estado"] == "ok"
    assert b["valores"]["sabana_personal"]["cantidad"] == 12
    assert b["valores"]["toalla_bano"]["cantidad"] == 7  # TB = baño en este bloque (sin TM ni ALF)
    assert any("negativos" in a for a in b["avisos"])


def test_faltantes_con_solo_la_diferencia(general):
    b = general["faltantes"]
    assert b["estado"] == "ok"
    assert b["valores"]["sabana_personal"] == {"faltantes": 7, "sobrantes": 0}
    assert b["valores"]["sabana_king"] == {"faltantes": 0, "sobrantes": 3}
    assert b["valores"]["toalla_piscina"]["faltantes"] == 2
    assert b["totales"] == {"hoteles": ["HOTEL UNO"], "importe": 12.5}


def test_habana_bloques_sin_rellenar_son_vacio(habana):
    for bloque in ("rechazo", "mal_uso"):
        assert habana[bloque]["estado"] == "vacio", bloque


def test_habana_descarte_e_inyecciones_por_hoteles(habana):
    assert habana["retiro"]["estado"] == "ok"
    assert habana["retiro"]["valores"]["sabana_king"]["cantidad"] == 3
    assert habana["inyecciones"]["valores"]["funda"]["cantidad"] == 150


def test_habana_desmanche_suma_sin_y_con_remojo(habana):
    b = habana["desmanche"]
    assert b["estado"] == "ok"
    assert b["valores"]["toalla_bano"] == {"tratadas": 30, "recuperadas": 24, "pct": 0.8}
    assert any("se suman" in a.lower() for a in b["avisos"])


def test_habana_inventarios_en_la_hoja_1(habana):
    b = habana["faltantes"]
    assert b["hoja"] == "5.Lenceria 1 (2)"
    assert b["valores"]["sabana_personal"]["faltantes"] == 2
    assert b["valores"]["toalla_alfombra"]["faltantes"] == 1
    assert b["totales"]["hoteles"] == ["HOTEL TRES"]


def test_bloque_que_no_existe(tmp_path):
    ruta = tmp_path / "VACD203001.xlsx"
    wb = Workbook()
    wb.active.title = "Lencería 2"
    bloque_hoteles(wb.active, 1, "RESUMEN DE LOS RECHAZOS POR TIPO DE PRENDA", {"Sabana King": (1, 1)})
    wb.save(ruta)
    r = cl.leer_cd(ruta)
    assert r["rechazo"]["estado"] == "ok"
    assert r["montas_habitacion"]["estado"] == "no_encontrado"
    assert r["faltantes"]["estado"] == "no_encontrado"


@pytest.mark.parametrize("texto,clave", [
    ("Sabana Personal", "sabana_personal"), ("Sábanas Personales", "sabana_personal"),
    ("Fundas", "funda"), ("Toallas Alfombra", "toalla_alfombra"), ("Toalla Baño", "toalla_bano"),
    ("Toallas Faciales", "toalla_facial"), ("Toallas Piscina", "toalla_piscina"),
    ("Mancha Amarilla", None), ("TOTAL", None),
])
def test_prenda_por_nombre(texto, clave):
    assert cl.prenda_por_nombre(texto) == clave


def test_codigos_segun_el_bloque():
    assert cl.tabla_codigos(["SP", "ALF", "TC", "TM"])["TC"] == "toalla_bano"
    assert cl.tabla_codigos(["SP", "TA", "TB", "TC"])["TC"] == "toalla_cara"
