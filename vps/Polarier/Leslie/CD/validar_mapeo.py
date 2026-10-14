#!/usr/bin/env python3
"""
Comprueba el mapa de agente_ca.py contra una TABLA CA ya rellenada a mano.

No escribe nada. Para cada CA extrae los indicadores con la misma lógica que el
agente y los compara con la columna del mes en la hoja del polo de la tabla.

    python validar_mapeo.py --tabla "TABLA CA 2026.xlsx" --bloque anterior CACA202504.xlsx ...
    python validar_mapeo.py --tabla "TABLA CA 2025.xlsx" --bloque actual --quimicos CACA202504.xlsx

--bloque anterior   compara con el bloque «AÑO ANTERIOR» (los CA son del año
                    previo al de la tabla; es como se validó el 2026-10-14)
--bloque actual     compara con el bloque del año de la tabla
--quimicos          compara también el bloque de productos químicos (solo existe
                    para el año actual de cada tabla)

Las seis cifras económicas salen del GLOBAL en el agente; aquí se comparan las
del propio CA (contraste), porque el GLOBAL de ese año puede no estar a mano.
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import agente_ca as A  # noqa: E402


def bloque_tabla(ws, mes, bloque):
    """{clave normalizada: valor} de la columna del mes en el bloque pedido."""
    cabeceras = [r for r in range(1, ws.max_row + 1)
                 if A.clave_norm(ws.cell(r, 1).value or "").startswith(("ano actual", "ano anterior"))]
    if not cabeceras:
        raise SystemExit(f"{ws.title}: no encuentro las cabeceras de año")
    r0 = cabeceras[0] if bloque == "actual" else next(
        r for r in cabeceras if A.clave_norm(ws.cell(r, 1).value).startswith("ano anterior"))
    col = A.columnas_meses_tabla(ws, r0)[mes]
    out, vistos = {}, set()
    for r in range(r0 + 1, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if not k or not str(k).strip():
            if out:
                break
            continue
        k = A.clave_norm(k).replace("alquiler", "alqujler")
        if k in vistos:
            break
        vistos.add(k)
        out[k] = ws.cell(r, col).value
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tabla", required=True)
    ap.add_argument("--bloque", choices=["anterior", "actual"], default="anterior")
    ap.add_argument("--quimicos", action="store_true")
    ap.add_argument("ficheros", nargs="+")
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    tabla = load_workbook(args.tabla, data_only=True)
    total = fallos = 0
    for f in args.ficheros:
        ruta = Path(f)
        ctx = A.contexto(ruta.name)
        ctx["ruta_cd"] = ruta
        valores, errores, detalle, contraste = A.extraccion_cd(ctx)
        for k, (v, o) in contraste.items():
            if k in A.GLOBAL_FILAS:
                valores.setdefault(k, {"valor": A.limpio(v), "origen": f"{o} (CA)"})
        ws = A.hoja_tabla(tabla, ctx)
        esperado = bloque_tabla(ws, ctx["mes"], args.bloque)
        print(f"\n## {ruta.name} → {ws.title} {ctx['mes']} (transporte: {ctx['criterio_transporte']})")
        for clave in A.CLAVES:
            e = esperado.get(A.clave_norm(clave))
            v = valores.get(clave, {}).get("valor")
            if e is None and not v:
                ok = True
            elif e is None or v is None:
                ok = False
            else:
                ok = abs(v - e) <= max(0.011, abs(e) * 1e-6) or (
                    clave in ("CDT EQUIPOS", "CDO EQUIPOS") and round(v) == round(e))
            total += 1
            fallos += not ok
            if not ok:
                print(f"  ✗ {clave:32} extraído={v!s:>16}  tabla={e!s:>16}  "
                      f"{valores.get(clave, {}).get('origen') or errores.get(clave, '')}")
        if args.quimicos:
            _, cols, productos = A.mapa_quimicos(ws)
            q = {A.compacto(k): v for k, v in (detalle.get("quimicos") or {}).items()}
            if not q:
                print(f"  ✗ químicos: {detalle.get('quimicos_error')}")
            for pk, (r, etq) in productos.items():
                e = ws.cell(r, cols[ctx["mes"]]).value
                v = q.get(pk)
                ok = (not e and not v) or (e is not None and v is not None and abs(v - e) < 0.06)
                total += 1
                fallos += not ok
                if not ok:
                    print(f"  ✗ químico {etq:24} extraído={v!s:>10}  tabla={e!s:>10}")
    print(f"\n{total - fallos}/{total} coinciden")


if __name__ == "__main__":
    main()
