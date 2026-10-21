"""
Resume la TABLA CA <año> por mes y polo para cargarla en Supabase (tabla_ca_meses).

No escribe nada: imprime por stdout un JSON con una fila por mes × polo ×
indicador que dice si el dato está o no (sin el valor). n8n lo recoge y llama a
la función sync_tabla_ca_meses() de Supabase, que reemplaza el año en una sola
transacción (migración 013). La app lo usa para decir, mes a mes, qué falta
para generar el informe del Consejo (vault/Informe-Consejo-Origen-Datos.md).

Reglas:
  - Fichero: el «TABLA CA <año>*.xlsx» modificado más recientemente en la
    carpeta del script, o el que se pase como argumento. El año sale del nombre.
  - Hojas de polo: HABANA, VARADERO, CAIBARIEN, CAYOCOCO, CAYOCRUZ, HOLGUÍN
    (se comparan sin tildes ni espacios: la de Cayo Coco llevó un espacio final).
  - Bloques: «AÑO ACTUAL» (con la fila de meses ENERO…DICIEMBRE), «AÑO ANTERIOR»
    y «CONSUMO DE PRODUCTOS QUÍMICOS»; los indicadores se buscan por su etiqueta
    en la columna A dentro de su bloque, nunca por número de fila.
  - Presente: una celda con número escrito (0 incluido: un 0 tecleado es un dato).
    Las fórmulas no se fían de su caché:
      * vínculo a una celda del GLOBAL (=[n]VARCP2026!$G$2): presente si esa
        celda de GLOBAL-CP<aa>.xlsx (carpeta de arriba) tiene un número, o si la
        caché de la TABLA CA trae un número distinto de 0. Un 0 en caché no
        cuenta: es lo que deja Excel en los meses aún sin cargar;
      * cualquier otra fórmula: presente solo si su caché es un número distinto
        de 0 (una suma de celdas vacías da 0).
  - Ingresos del CD: presentes si lo está alquiler, lavado u otros (la fila
    «REAL INGRESOS» es su suma). Químicos: si lo está algún producto.

Uso:
  python sync_tabla_ca.py ["ruta/TABLA CA 2026.xlsx"] [--global "ruta/GLOBAL-CP26.xlsx"]
"""

import glob
import hashlib
import json
import os
import re
import sys
import unicodedata

import openpyxl
from openpyxl.utils import column_index_from_string

POLOS = {
    "HABANA": "habana",
    "VARADERO": "varadero",
    "CAIBARIEN": "caibarien",
    "CAYOCOCO": "cayo_coco",
    "CAYOCRUZ": "cayo_cruz",
    "HOLGUIN": "holguin",
}

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

# indicador → (bloque, etiquetas en la columna A). Varias etiquetas = basta una.
INDICADORES = {
    "td_plan":      ("actual",   ["PLAN TURISTAS DIAS"]),
    "td_real":      ("actual",   ["REAL TURISTAS DIAS"]),
    "prod_plan":    ("actual",   ["PLAN PRODUCCION"]),
    "prod_real":    ("actual",   ["REAL PRODUCCION"]),
    "prendas":      ("actual",   ["PRENDAS PROCESADAS"]),
    "ingresos_cd":  ("actual",   ["INGRESOS ALQUILER DE LENCERIA", "INGRESOS ALQUJLER DE LENCERIA",
                                  "INGRESOS LAVADO", "OTROS INGRESOS"]),
    "trabajadores": ("actual",   ["PROMEDIO DE TRABAJADORES"]),
    "cdt":          ("actual",   ["CDT EQUIPOS"]),
    "cdo":          ("actual",   ["CDO EQUIPOS"]),
    "diesel":       ("actual",   ["CONSUMO DIESEL PRODUCCION"]),
    "electricidad": ("actual",   ["CONSUMO ELECTRICIDAD"]),
    "agua":         ("actual",   ["CONSUMO AGUA"]),
    "calidad":      ("actual",   ["EVALUACION CALIDAD HOTELES"]),
    "td_real_aa":   ("anterior", ["REAL TURISTAS DIAS"]),
    "prod_real_aa": ("anterior", ["REAL PRODUCCION"]),
    "ingresos_aa":  ("anterior", ["REAL INGRESOS"]),
}

VINCULO_GLOBAL = re.compile(r"^=\s*(?:\[\d+\])?'?([A-Z]{2,3}CP\d{4})'?!\$?([A-Z]{1,3})\$?(\d+)\s*$")


def norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().upper()


def elegir_fichero(args):
    rutas = [a for a in args if not a.startswith("--")]
    if rutas:
        return rutas[0]
    carpeta = os.path.dirname(os.path.abspath(__file__))
    candidatos = glob.glob(os.path.join(carpeta, "TABLA CA *.xlsx"))
    if not candidatos:
        raise FileNotFoundError(f"No hay ninguna «TABLA CA <año>.xlsx» en {carpeta}")
    return max(candidatos, key=os.path.getmtime)


def ruta_global(args, ruta_tabla, anio):
    if "--global" in args:
        return args[args.index("--global") + 1]
    arriba = os.path.dirname(os.path.dirname(os.path.abspath(ruta_tabla)))
    return os.path.join(arriba, f"GLOBAL-CP{anio % 100:02d}.xlsx")


class Lector:
    """Decide si una celda de la TABLA CA tiene un dato (ver las reglas arriba)."""

    def __init__(self, ruta_tabla, ruta_global):
        self.f = openpyxl.load_workbook(ruta_tabla, data_only=False)
        self.v = openpyxl.load_workbook(ruta_tabla, data_only=True)
        self.globl = None
        if ruta_global and os.path.exists(ruta_global):
            self.globl = openpyxl.load_workbook(ruta_global, data_only=False)
        self.sin_global = 0

    def presente(self, hoja, fila, col) -> bool:
        raw = self.f[hoja].cell(fila, col).value
        if isinstance(raw, bool):
            return False
        if isinstance(raw, (int, float)):
            return True
        if isinstance(raw, str) and raw.startswith("="):
            m = VINCULO_GLOBAL.match(raw.replace(" ", ""))
            if m and self.globl is not None and m.group(1) in self.globl.sheetnames:
                destino = self.globl[m.group(1)].cell(int(m.group(3)), column_index_from_string(m.group(2))).value
                if isinstance(destino, (int, float)) and not isinstance(destino, bool):
                    return True
                # Vacía o con fórmula en el GLOBAL: decide la caché (abajo). Un 0 en
                # caché no cuenta, que es lo que deja Excel en los meses sin cargar.
            elif m:
                self.sin_global += 1
            cache = self.v[hoja].cell(fila, col).value
            return isinstance(cache, (int, float)) and not isinstance(cache, bool) and cache != 0
        return False


def bloques(ws):
    """Filas de cabecera de cada bloque y columnas de los meses."""
    actual = anterior = quimicos = None
    for fila in range(1, ws.max_row + 1):
        etq = norm(ws.cell(fila, 1).value)
        if etq.startswith("ANO ACTUAL") and actual is None:
            actual = fila
        elif etq.startswith("ANO ANTERIOR") and anterior is None:
            anterior = fila
        elif etq.startswith("CONSUMO DE PRODUCTOS QUIMICOS") and quimicos is None:
            quimicos = fila
    if actual is None or anterior is None:
        raise ValueError(f"{ws.title}: no encuentro los bloques «AÑO ACTUAL» y «AÑO ANTERIOR»")
    cols = {}
    for col in range(1, ws.max_column + 1):
        mes = norm(ws.cell(actual, col).value)
        if mes in MESES and mes not in cols:
            cols[mes] = col
    if len(cols) != 12:
        raise ValueError(f"{ws.title}: la fila {actual} debería tener los 12 meses, tiene {sorted(cols)}")
    return actual, anterior, quimicos, cols


def filas_por_etiqueta(ws, desde, hasta):
    out = {}
    for fila in range(desde + 1, hasta):
        etq = norm(ws.cell(fila, 1).value)
        if etq and etq not in out:
            out[etq] = fila
    return out


def main(argv):
    ruta = elegir_fichero(argv)
    m = re.search(r"TABLA CA (\d{4})", os.path.basename(ruta))
    if not m:
        raise ValueError(f"No saco el año del nombre «{os.path.basename(ruta)}»")
    anio = int(m.group(1))
    lector = Lector(ruta, ruta_global(argv, ruta, anio))

    filas, resumen = [], {}
    for hoja in lector.f.sheetnames:
        polo = POLOS.get(norm(hoja).replace(" ", ""))
        if not polo:
            continue
        ws = lector.f[hoja]
        actual, anterior, quimicos, cols = bloques(ws)
        fin_anterior = quimicos or ws.max_row + 1
        etq_actual = filas_por_etiqueta(ws, actual, anterior)
        etq_anterior = filas_por_etiqueta(ws, anterior, fin_anterior)
        productos = []
        if quimicos:
            for fila in range(quimicos + 1, ws.max_row + 1):
                etq = norm(ws.cell(fila, 1).value)
                if etq.startswith("TOTAL") or etq.startswith("INDICE"):
                    break
                if etq:
                    productos.append(fila)

        meses_con_cd = []
        for n_mes, nombre_mes in enumerate(MESES, start=1):
            col = cols[nombre_mes]
            hay_cd = False
            for indicador, (bloque, etiquetas) in INDICADORES.items():
                mapa = etq_actual if bloque == "actual" else etq_anterior
                filas_ind = [mapa[e] for e in etiquetas if e in mapa]
                if not filas_ind:
                    raise ValueError(f"{hoja}: no encuentro «{etiquetas[0]}» en el bloque {bloque}")
                presente = any(lector.presente(hoja, f, col) for f in filas_ind)
                if indicador == "td_real" and presente:
                    hay_cd = True
                filas.append({
                    "mes": n_mes, "polo": polo, "indicador": indicador, "presente": presente,
                    "fuente": f"{hoja.strip()}!{openpyxl.utils.get_column_letter(col)}{filas_ind[0]}",
                })
            filas.append({
                "mes": n_mes, "polo": polo, "indicador": "quimicos",
                "presente": any(lector.presente(hoja, f, col) for f in productos),
                "fuente": f"{hoja.strip()}!{openpyxl.utils.get_column_letter(col)}{productos[0]}" if productos else None,
            })
            if hay_cd:
                meses_con_cd.append(n_mes)
        resumen[polo] = {"meses_con_td_real": meses_con_cd}

    if len(resumen) != len(POLOS):
        faltan = sorted(set(POLOS.values()) - set(resumen))
        raise ValueError(f"Faltan hojas de polo en la TABLA CA: {faltan}")

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    json.dump({
        "archivo": os.path.basename(ruta),
        "anio": anio,
        "script_sha256": hashlib.sha256(open(__file__, "rb").read()).hexdigest(),
        "global_encontrado": lector.globl is not None,
        "vinculos_sin_global": lector.sin_global,
        "resumen": resumen,
        "filas": filas,
    }, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main(sys.argv[1:])
