#!/usr/bin/env python3
"""
Herramientas del agente "Rellenar TABLA CA" (n8n + Claude).

Sustituye a rellenar_ca.py. En vez de un proceso ciego de principio a fin,
expone operaciones pequeñas que el agente de n8n llama por el pythonrunner:

    POST http://polarierauto_pythonrunner:8000/run
    {"script": "Leslie/CD/agente_ca.py", "args": ["<comando>", "<json>"]}

Comandos (todos reciben un JSON con al menos {"cd": "<nombre del fichero CD>"}):

    preparar        Contexto (polo, año, mes, rutas), extracción de referencia
                    (la misma lógica que rellenar_ca.py, indicador a indicador,
                    con su origen o su error) y el estado actual de la tabla.
    ver_hoja        Celdas no vacías de una hoja del CD, con coordenadas Excel.
    ver_global      Filas de la hoja del polo en el GLOBAL (Plan/Real del mes).
    estado_tabla    Valores y fórmulas de la columna del mes en la TABLA CA.
    escribir        Escribe {clave: valor} en la columna del mes. Nunca pisa
                    una celda con fórmula. Hace copia de seguridad antes.
    limpiar         Borra el CD ya procesado.

Rutas (relativas a este script, así funciona igual en local y en el VPS):
    <CD>                         CD subido desde la app (lo escribe n8n)
    TABLA CA <año>.xlsx          tabla que se rellena
    ../GLOBAL-CP<aa>.xlsx        GLOBAL vivo de la rama Leslie (o ./GLOBAL-CP<aa>.xlsx)
    respaldos/                   copias de la tabla antes de cada escritura

Toda salida es JSON por stdout. Los errores salen como {"ok": false, "error": ...}
con código 0, para que el agente pueda leerlos y razonar en vez de cortar el flujo.
"""

import hashlib
import json
import re
import shutil
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
ENTRADA = BASE_DIR  # n8n deja el CD junto a la tabla, como hacía rellenar_ca.py
RESPALDOS = BASE_DIR / "respaldos"
MAX_RESPALDOS = 15

# prefijo del fichero CD -> (prefijo de hoja en GLOBAL, hoja en TABLA CA, nombre)
POLOS = {
    "LH": ("LHCP", "HABANA", "La Habana"),
    "VA": ("VARCP", "VARADERO", "Varadero"),
    "CA": ("CACP", "CAIBARIEN", "Caibarién"),
    "CC": ("CCCP", "CAYOCOCO", "Cayo Coco"),
    "CCX": ("CCZCP", "CAYOCRUZ", "Cayo Cruz"),
    "HOL": ("HOLCP", "HOLGUIN", "Holguín"),
}
ALIAS_POLO = {"HA": "LH", "CX": "CCX", "CCZ": "CCX", "HO": "HOL"}

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

# Claves de la TABLA CA que rellena este proceso (columna A de la hoja del polo).
# "ALQUJLER" va con la errata tal como está escrito en la tabla.
CLAVES = [
    "REAL TURISTAS DÍAS", "REAL PRODUCCIÓN", "PRENDAS PROCESADAS", "REAL INGRESOS",
    "INGRESOS ALQUJLER DE LENCERÍA", "INGRESOS LAVADO", "OTROS INGRESOS",
    "REAL COSTO DE LAVANDERÍA", "REAL GASTOS TOTALES", "UTILIDAD EN OPERACIONES",
    "UTILIDAD NETA", "VALOR AGREGADO BRUTO", "GASTO DE SALARIO", "DÍAS TRABAJADOS",
    "HORAS TRABAJADAS", "PROMEDIO DE TRABAJADORES", "PROM TRAB PRODUCTIVOS DIARIOS",
    "VEHÍCULOS EN USO TOTALES", "DE ELLOS PROPIOS", "KILÓMETROS RECORRIDOS",
    "CONSUMO DIESEL TRANSPORTE", "CDT EQUIPOS", "CDO EQUIPOS",
    "CONSUMO DIESEL PRODUCCIÓN", "CONSUMO ELECTRICIDAD", "CONSUMO AGUA",
    "CAPACIDAD INSTALADA POR HORA", "EVALUACIÓN CALIDAD HOTELES",
]


# ----------------------------------------------------------------- utilidades

def sin_tildes(s):
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c))


def clave_norm(s):
    """Normaliza una clave de la tabla: sin tildes, minúsculas, espacios simples."""
    return re.sub(r"\s+", " ", sin_tildes(s).strip().lower())


def compacto(s):
    """Normalización agresiva para nombres de hoja: sin tildes ni separadores."""
    return re.sub(r"[\s.\-_/]+", "", sin_tildes(s).strip().lower())


def num(v):
    """Convierte a float si es numérico; None en cualquier otro caso."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)  # descarta NaN
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def limpio(v):
    """Valor listo para JSON: enteros sin .0, floats redondeados, texto tal cual."""
    if isinstance(v, float):
        if v != v:
            return None
        return int(v) if v.is_integer() else round(v, 6)
    if isinstance(v, datetime):
        return v.isoformat()
    return v


def coord(r, c):
    """(fila, columna) 0-based -> 'C7'."""
    letras = ""
    c += 1
    while c:
        c, resto = divmod(c - 1, 26)
        letras = chr(65 + resto) + letras
    return f"{letras}{r + 1}"


def salir(obj):
    sys.stdout.write(json.dumps(obj, ensure_ascii=False, default=str))
    sys.stdout.write("\n")
    sys.exit(0)


class ErrorAgente(Exception):
    pass


# ------------------------------------------------------------------ contexto

def contexto(nombre_cd):
    nombre_cd = str(nombre_cd or "").strip()
    if not nombre_cd or "/" in nombre_cd or "\\" in nombre_cd or ".." in nombre_cd:
        raise ErrorAgente(f"Nombre de CD no válido: {nombre_cd!r}")

    m = re.match(r"^([A-Za-z]+?)CD[\s_\-]*(\d{4})[\s_\-]*(\d{2})", nombre_cd, re.IGNORECASE)
    if not m:
        raise ErrorAgente(
            f"El nombre '{nombre_cd}' no sigue el patrón <POLO>CD<AAAA><MM>.xlsx "
            "(p. ej. CACD202609.xlsx)."
        )
    prefijo = m.group(1).upper()
    prefijo = ALIAS_POLO.get(prefijo, prefijo)
    if prefijo not in POLOS:
        raise ErrorAgente(
            f"Prefijo de polo '{m.group(1)}' desconocido. Válidos: {', '.join(POLOS)}."
        )
    anio, mes_num = int(m.group(2)), int(m.group(3))
    if not 1 <= mes_num <= 12:
        raise ErrorAgente(f"Mes '{m.group(3)}' fuera de rango en '{nombre_cd}'.")

    global_cp, hoja_tabla, nombre_polo = POLOS[prefijo]
    aa = str(anio)[-2:]
    candidatos_global = [BASE_DIR.parent / f"GLOBAL-CP{aa}.xlsx", BASE_DIR / f"GLOBAL-CP{aa}.xlsx"]
    ruta_global = next((p for p in candidatos_global if p.exists()), candidatos_global[0])

    return {
        "cd": nombre_cd,
        "ruta_cd": ENTRADA / nombre_cd,
        "polo": nombre_polo,
        "prefijo": prefijo,
        "anio": anio,
        "mes_num": mes_num,
        "mes": MESES[mes_num - 1],
        "hoja_global": f"{global_cp}{anio}",
        "hoja_tabla": hoja_tabla,
        "ruta_global": ruta_global,
        "ruta_tabla": BASE_DIR / f"TABLA CA {anio}.xlsx",
    }


def contexto_publico(ctx):
    out = {k: v for k, v in ctx.items() if not k.startswith("ruta_")}
    for k in ("ruta_cd", "ruta_global", "ruta_tabla"):
        out[k.replace("ruta_", "existe_")] = ctx[k].exists()
        out[k.replace("ruta_", "fichero_")] = ctx[k].name
    return out


# ----------------------------------------------------------------- lectura CD

def abrir_cd(ctx):
    if not ctx["ruta_cd"].exists():
        raise ErrorAgente(f"No encuentro el CD '{ctx['cd']}' en la carpeta del script.")
    return load_workbook(ctx["ruta_cd"], data_only=True, read_only=True)


def hoja_fuzzy(wb, clave):
    k = compacto(clave)
    coincidencias = [s for s in wb.sheetnames if k in compacto(s)]
    if not coincidencias:
        raise ErrorAgente(f"No hay hoja para '{clave}'. Hojas del CD: {wb.sheetnames}")
    return sorted(coincidencias, key=len)[0]


def rejilla(ws):
    return [list(fila) for fila in ws.iter_rows(values_only=True)]


def celda(g, r, c):
    if r < 0 or r >= len(g) or c < 0 or c >= len(g[r]):
        return None
    return g[r][c]


def buscar(g, texto, quitar_tildes=False):
    """Coordenadas de las celdas cuyo texto normalizado es exactamente `texto`.
    Recorre por filas, como np.where en el script original, y salta la fila 1
    (pandas la usaba de cabecera y nunca buscaba en ella)."""
    hallados = []
    for r in range(1, len(g)):
        for c, v in enumerate(g[r]):
            if v is None:
                continue
            s = str(v).strip().lower()
            if quitar_tildes:
                s = sin_tildes(s)
            s = re.sub(r"\s+", " ", s)
            if s == texto:
                hallados.append((r, c))
    if not hallados:
        raise ErrorAgente(f"No encuentro la etiqueta '{texto}'")
    return hallados


def columna(g, desde, c):
    return [celda(g, r, c) for r in range(desde, len(g))]


# ------------------------------------------------- extracción de referencia

def extraccion_cd(ctx):
    """Misma lógica que rellenar_ca.py, pero cada indicador por separado:
    si uno falla, se anota el error y se sigue con los demás."""
    valores, errores = {}, {}

    def poner(clave, valor, origen):
        valores[clave] = {"valor": limpio(valor), "origen": origen}

    def intentar(claves, fn):
        try:
            fn()
        except Exception as e:  # el agente decide qué hacer con el hueco
            for k in claves:
                errores[k] = f"{type(e).__name__}: {e}"

    wb = abrir_cd(ctx)
    hojas = {}

    def H(clave):
        if clave not in hojas:
            nombre = hoja_fuzzy(wb, clave)
            hojas[clave] = (nombre, rejilla(wb[nombre]))
        return hojas[clave]

    # --- Economía
    def turistas():
        nombre, g = H("economia")
        hits = buscar(g, "turistas dias", quitar_tildes=True)
        if len(hits) < 2:
            raise ErrorAgente("'turistas dias' aparece una sola vez (el script usaba la 2ª)")
        r, c = hits[1]
        poner("REAL TURISTAS DÍAS", num(celda(g, r, c + 2)), f"{nombre}!{coord(r, c + 2)} (2ª 'Turistas días', +2 col)")
    intentar(["REAL TURISTAS DÍAS"], turistas)

    def arrendamiento():
        nombre, g = H("economia")
        r, c = buscar(g, "arrendamiento")[0]
        poner("INGRESOS ALQUJLER DE LENCERÍA", num(celda(g, r, c + 2)), f"{nombre}!{coord(r, c + 2)} ('Arrendamiento', +2 col)")
    intentar(["INGRESOS ALQUJLER DE LENCERÍA"], arrendamiento)

    def lavado_y_otros():
        nombre, g = H("economia")
        r, c = buscar(g, "lavado")[0]
        poner("INGRESOS LAVADO", num(celda(g, r, c + 2)), f"{nombre}!{coord(r, c + 2)} ('Lavado', +2 col)")
        col = [num(v) or 0.0 for v in columna(g, r + 1, c + 2)]
        no_cero = [v for v in col if v != 0]
        otros = round(sum(col) - (no_cero[-1] if no_cero else 0), 2)
        poner("OTROS INGRESOS", otros,
              f"{nombre}!{coord(r + 1, c + 2)}:{coord(len(g) - 1, c + 2)} (suma bajo 'Lavado' menos el último no nulo, que se asume total)")
    intentar(["INGRESOS LAVADO", "OTROS INGRESOS"], lavado_y_otros)

    def real_ingresos():
        partes = ["INGRESOS ALQUJLER DE LENCERÍA", "INGRESOS LAVADO", "OTROS INGRESOS"]
        if any(p not in valores for p in partes):
            raise ErrorAgente("falta alguna de sus tres partes")
        total = sum((valores[p]["valor"] or 0) for p in partes)
        poner("REAL INGRESOS", total, "suma de alquiler + lavado + otros ingresos")
    intentar(["REAL INGRESOS"], real_ingresos)

    # --- Producción
    def produccion():
        nombre, g = H("produccion")
        r, c = buscar(g, "produccion (kg)", quitar_tildes=True)[0]
        poner("REAL PRODUCCIÓN", num(celda(g, r, c + 2)), f"{nombre}!{coord(r, c + 2)} ('Producción (kg)', +2 col)")
    intentar(["REAL PRODUCCIÓN"], produccion)

    def prendas():
        nombre, g = H("produccion")
        r, c = buscar(g, "procesado")[0]
        total = sum(num(celda(g, i, c)) or 0 for i in range(r + 1, r + 32))
        poner("PRENDAS PROCESADAS", round(total, 2), f"{nombre}!{coord(r + 1, c)}:{coord(r + 31, c)} (31 filas bajo 'Procesado')")
    intentar(["PRENDAS PROCESADAS"], prendas)

    def capacidad():
        nombre, g = H("produccion")
        r, c = buscar(g, "capacidad tunel lavado (x hora)", quitar_tildes=True)[0]
        a, b = num(celda(g, r, c + 2)), num(celda(g, r + 2, c + 2))
        x, y = num(celda(g, r + 1, c + 2)), num(celda(g, r + 3, c + 2))
        if None in (a, b, x, y):
            raise ErrorAgente(f"celdas vacías en {coord(r, c + 2)}:{coord(r + 3, c + 2)}")
        poner("CAPACIDAD INSTALADA POR HORA", a * b + x * y,
              f"{nombre}!{coord(r, c + 2)}*{coord(r + 2, c + 2)} + {coord(r + 1, c + 2)}*{coord(r + 3, c + 2)}")
    intentar(["CAPACIDAD INSTALADA POR HORA"], capacidad)

    # --- RRHH
    def rrhh(clave, etiqueta, tildes=False, alternativa=None):
        def fn():
            nombre, g = H("rrhh")
            try:
                r, c = buscar(g, etiqueta, quitar_tildes=tildes)[0]
                usada = etiqueta
            except ErrorAgente:
                if not alternativa:
                    raise
                r, c = buscar(g, alternativa)[0]
                usada = alternativa
            poner(clave, num(celda(g, r, c + 2)), f"{nombre}!{coord(r, c + 2)} ('{usada}', +2 col)")
        intentar([clave], fn)

    rrhh("DÍAS TRABAJADOS", "dias trabajados", tildes=True)
    rrhh("HORAS TRABAJADAS", "horas promedio trabajadas")
    rrhh("PROMEDIO DE TRABAJADORES", "promedio de trabajadores", alternativa="total trabajadores")
    rrhh("PROM TRAB PRODUCTIVOS DIARIOS", "trabajadores productivos")

    # --- Transporte
    def contar_no_cero(g, r, c):
        col = [num(v) or 0.0 for v in columna(g, r + 1, c)]
        return max(sum(1 for v in col if v != 0) - 1, 0)

    def vehiculos():
        nombre, g = H("transporte")
        r, c = buscar(g, "km")[0]
        propios = contar_no_cero(g, r, c)
        poner("DE ELLOS PROPIOS", propios, f"{nombre}!{coord(r + 1, c)}↓ (nº de valores no nulos bajo 'KM' menos el total)")
        r2, c2 = buscar(g, "importe pagado")[0]
        alquilados = contar_no_cero(g, r2, c2)
        poner("VEHÍCULOS EN USO TOTALES", propios + alquilados,
              f"propios ({propios}) + alquilados ({alquilados}, no nulos bajo 'Importe pagado' en {nombre}!{coord(r2, c2)} menos el total)")
    intentar(["DE ELLOS PROPIOS", "VEHÍCULOS EN USO TOTALES"], vehiculos)

    def kilometros():
        nombre, g = H("transporte")
        r, c = buscar(g, "km")[0]
        vals = [(i, v) for i, v in enumerate(columna(g, r + 2, c), start=r + 2) if v is not None]
        if not vals:
            raise ErrorAgente("columna 'KM' vacía")
        i, v = vals[-1]
        poner("KILÓMETROS RECORRIDOS", num(v), f"{nombre}!{coord(i, c)} (último valor bajo 'KM')")
    intentar(["KILÓMETROS RECORRIDOS"], kilometros)

    def diesel_transporte():
        nombre, g = H("transporte")
        try:
            r, c = buscar(g, "comb")[0]
            desde, col_c, etiqueta = r + 2, c, "Comb"
        except ErrorAgente:
            r, c = buscar(g, "consumo")[0]
            desde, col_c, etiqueta = r + 1, c + 1, "Consumo (+1 col)"
        vals = [(i, v) for i, v in enumerate(columna(g, desde, col_c), start=desde) if v is not None]
        if not vals:
            raise ErrorAgente(f"columna bajo '{etiqueta}' vacía")
        i, v = vals[-1]
        poner("CONSUMO DIESEL TRANSPORTE", num(v), f"{nombre}!{coord(i, col_c)} (último valor bajo '{etiqueta}')")
    intentar(["CONSUMO DIESEL TRANSPORTE"], diesel_transporte)

    # --- SSTT
    def cdt_cdo():
        nombre, g = H("sstt")
        r, c = buscar(g, "cdt")[0]
        cdt, cdo = num(celda(g, r, c + 1)), num(celda(g, r - 1, c + 1))
        if cdt is None or cdo is None:
            raise ErrorAgente(f"celdas vacías en {coord(r - 1, c + 1)}:{coord(r, c + 1)}")
        poner("CDT EQUIPOS", round(cdt * 100), f"{nombre}!{coord(r, c + 1)} ×100 ('CDT', +1 col)")
        poner("CDO EQUIPOS", round(cdo * 100), f"{nombre}!{coord(r - 1, c + 1)} ×100 (fila sobre 'CDT', +1 col)")
    intentar(["CDT EQUIPOS", "CDO EQUIPOS"], cdt_cdo)

    def consumos():
        nombre, g = H("sstt")
        r, c = buscar(g, "consumos")[0]
        diesel, elec, agua = (num(celda(g, r, c + k)) for k in (1, 2, 3))
        poner("CONSUMO DIESEL PRODUCCIÓN", diesel, f"{nombre}!{coord(r, c + 1)} ('Consumos', +1 col)")
        poner("CONSUMO ELECTRICIDAD", None if elec is None else int(elec), f"{nombre}!{coord(r, c + 2)} ('Consumos', +2 col)")
        poner("CONSUMO AGUA", None if agua is None else int(agua), f"{nombre}!{coord(r, c + 3)} ('Consumos', +3 col)")
    intentar(["CONSUMO DIESEL PRODUCCIÓN", "CONSUMO ELECTRICIDAD", "CONSUMO AGUA"], consumos)

    # --- Calidad
    def calidad():
        nombre, g = H("calidad")
        r, c = buscar(g, "promedio total")[0]
        v, col = num(celda(g, r, c + 1)), c + 1
        if v is None:
            v, col = num(celda(g, r, c + 2)), c + 2
        poner("EVALUACIÓN CALIDAD HOTELES", None if v is None else round(v, 2), f"{nombre}!{coord(r, col)} ('Promedio total')")
    intentar(["EVALUACIÓN CALIDAD HOTELES"], calidad)

    wb.close()
    return valores, errores


# ---------------------------------------------------------------- GLOBAL

def _evaluador_global():
    """Evaluador de fórmulas de sync_global_ai.py (carpeta padre, ya en el VPS).
    Modelo1.py guarda el GLOBAL con openpyxl y eso borra los valores calculados,
    así que leer con data_only daría las fórmulas vacías."""
    try:
        if str(BASE_DIR.parent) not in sys.path:
            sys.path.insert(0, str(BASE_DIR.parent))
        from sync_global_ai import Evaluador, columnas_meses
        return Evaluador, columnas_meses
    except Exception:
        return None, None


def filas_global(ctx):
    """Filas de la hoja del polo en el GLOBAL con Plan y Real del mes.
    Devuelve (filas, letra de la columna Real, si las fórmulas se evaluaron)."""
    if not ctx["ruta_global"].exists():
        raise ErrorAgente(f"No existe {ctx['ruta_global'].name} en el servidor.")
    Evaluador, columnas_meses = _evaluador_global()
    wb = load_workbook(ctx["ruta_global"], data_only=Evaluador is None)
    if ctx["hoja_global"] not in wb.sheetnames:
        raise ErrorAgente(f"El GLOBAL no tiene la hoja {ctx['hoja_global']}. Hojas: {wb.sheetnames}")
    ws = wb[ctx["hoja_global"]]

    cab = next((r for r in range(1, 21) if str(ws.cell(r, 1).value or "").strip() == "Filas"), None)
    if cab is None:
        raise ErrorAgente("No encuentro la fila de cabecera 'Filas' en el GLOBAL.")

    if Evaluador:
        cols = columnas_meses(ws).get(ctx["mes_num"])
        if not cols:
            raise ErrorAgente(f"No encuentro las columnas PLAN/REAL de {ctx['mes']} en el GLOBAL.")
        c_plan, c_real = cols[0], cols[1]
        ev = Evaluador(ws)

        def leer(r, c):
            try:
                return ev.excel(r, c)
            except ValueError:  # fórmula que el evaluador no soporta
                return None
    else:
        c_plan = next((c for c in range(1, ws.max_column + 1)
                       if str(ws.cell(cab, c).value or "").strip().upper() == ctx["mes"]), None)
        if c_plan is None:
            raise ErrorAgente(f"No encuentro la columna {ctx['mes']} en la cabecera del GLOBAL.")
        c_real = c_plan + 1

        def leer(r, c):
            return num(ws.cell(r, c).value)

    filas = []
    for r in range(cab + 2, ws.max_row + 1):
        codigo, titulo = ws.cell(r, 1).value, ws.cell(r, 2).value
        if codigo is None and titulo is None:
            continue
        filas.append({
            "excel_fila": r,
            "fila": limpio(num(codigo)) if num(codigo) is not None else codigo,
            "titulo": titulo,
            "plan": limpio(leer(r, c_plan)),
            "real": limpio(leer(r, c_real)),
        })
    wb.close()
    return filas, coord(0, c_real - 1)[:-1], Evaluador is not None


VAB_SUMA = [30, 144]
VAB_RESTA = [69, 77, 93, 129, 112, 115, 121, 136, 176, 177, 178, 180, 182, 198,
             113, 114, 116, 117, 123, 181, 195, 200]


def extraccion_global(ctx):
    valores, errores = {}, {}
    try:
        filas, col_real, evaluado = filas_global(ctx)
    except Exception as e:
        msg = f"{type(e).__name__}: {e}"
        for k in ["REAL COSTO DE LAVANDERÍA", "REAL GASTOS TOTALES", "UTILIDAD EN OPERACIONES",
                  "UTILIDAD NETA", "GASTO DE SALARIO", "VALOR AGREGADO BRUTO"]:
            errores[k] = msg
        return valores, errores, None

    hoja = ctx["hoja_global"]

    def por_titulo(t):
        return [f for f in filas if f["titulo"] is not None and clave_norm(f["titulo"]) == t]

    def poner(clave, valor, origen):
        valores[clave] = {"valor": limpio(valor), "origen": origen}

    def real(f):
        return f["real"] or 0

    try:
        hits = por_titulo("lavanderia")
        if len(hits) < 2:
            raise ErrorAgente("'Lavandería' aparece menos de dos veces (el script usaba la 2ª)")
        poner("REAL COSTO DE LAVANDERÍA", hits[1]["real"], f"{hoja}!{col_real}{hits[1]['excel_fila']} (2ª fila 'Lavandería', Real)")
    except Exception as e:
        errores["REAL COSTO DE LAVANDERÍA"] = str(e)

    try:
        hits = por_titulo("gastos fijos") + por_titulo("gastos indirectos")
        if not hits:
            raise ErrorAgente("no hay filas 'Gastos Fijos' ni 'Gastos Indirectos'")
        poner("REAL GASTOS TOTALES", sum(real(f) for f in hits),
              " + ".join(f"{hoja}!{col_real}{f['excel_fila']} ({str(f['titulo']).strip()})" for f in hits))
    except Exception as e:
        errores["REAL GASTOS TOTALES"] = str(e)

    for clave, titulo in [("UTILIDAD EN OPERACIONES", "utilidades antes de cargo fijo"),
                          ("UTILIDAD NETA", "utilidades antes de impuestos"),
                          ("GASTO DE SALARIO", "salario")]:
        hits = por_titulo(titulo)
        if hits:
            poner(clave, hits[0]["real"], f"{hoja}!{col_real}{hits[0]['excel_fila']} ('{hits[0]['titulo'].strip()}', Real)")
        else:
            errores[clave] = f"no hay fila '{titulo}'"

    try:
        por_num = {}
        for f in filas:
            if isinstance(f["fila"], (int, float)) and f["fila"] not in por_num:
                por_num[f["fila"]] = f
        faltan = [n for n in VAB_SUMA + VAB_RESTA if n not in por_num]
        if faltan:
            raise ErrorAgente(f"faltan las filas {faltan} del GLOBAL")
        vab = sum(real(por_num[n]) for n in VAB_SUMA) - sum(real(por_num[n]) for n in VAB_RESTA)
        poner("VALOR AGREGADO BRUTO", round(vab, 2),
              "Real filas 30 + 144 − (69, 77, 93, 112–117, 121, 123, 129, 136, 176–178, 180–182, 195, 198, 200) del GLOBAL")
    except Exception as e:
        errores["VALOR AGREGADO BRUTO"] = str(e)

    vacias = sum(1 for f in filas if not f["real"])
    avisos = []
    if filas and vacias / len(filas) > 0.9:
        avisos.append(f"La columna Real de {ctx['mes']} en {hoja} está vacía o a cero en {vacias}/{len(filas)} "
                      "filas: el mes no parece cargado en el GLOBAL (falta subir el PDF CP de ese mes).")
    if not evaluado:
        avisos.append("No se pudo cargar el evaluador de sync_global_ai.py: el GLOBAL se leyó con los valores "
                      "guardados, que faltan en las celdas con fórmula si Modelo1.py lo guardó sin recalcular.")
    return valores, errores, " ".join(avisos) or None


# ---------------------------------------------------------------- TABLA CA

def abrir_tabla(ctx, solo_valores):
    if not ctx["ruta_tabla"].exists():
        raise ErrorAgente(f"No existe '{ctx['ruta_tabla'].name}' en el servidor.")
    return load_workbook(ctx["ruta_tabla"], data_only=solo_valores)


def hoja_tabla(wb, ctx):
    objetivo = compacto(ctx["hoja_tabla"])
    for s in wb.sheetnames:
        if compacto(s) == objetivo:
            return wb[s]
    raise ErrorAgente(f"La tabla no tiene hoja para {ctx['polo']} ({ctx['hoja_tabla']}). Hojas: {wb.sheetnames}")


def mapa_tabla(ws, ctx):
    """Fila de cabecera, columna de cada mes y filas de las claves del año actual
    (de la cabecera hacia abajo hasta la primera clave repetida, que marca el
    comienzo del bloque del año anterior)."""
    cab = None
    for r in range(1, min(ws.max_row, 20) + 1):
        if any(v is not None and clave_norm(v) == "enero" for v in
               (ws.cell(r, c).value for c in range(1, min(ws.max_column, 30) + 1))):
            cab = r
            break
    if cab is None:
        raise ErrorAgente("No encuentro la fila de cabecera con los meses en la tabla.")

    cols_mes = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(cab, c).value
        if v is None:
            continue
        n = clave_norm(v).upper()
        n = "SEPTIEMBRE" if n == "SETIEMBRE" else n
        if n in MESES and n not in cols_mes:
            cols_mes[n] = c
    if ctx["mes"] not in cols_mes:
        raise ErrorAgente(f"La cabecera de la tabla no tiene la columna {ctx['mes']}.")

    filas, vistos = {}, set()
    for r in range(cab + 1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None or not str(v).strip():
            continue
        k = clave_norm(v)
        if k in vistos:
            break
        vistos.add(k)
        filas[k] = (r, str(v).strip())
    return cab, cols_mes, filas


def estado_tabla(ctx, meses_previos=3):
    wb_f = abrir_tabla(ctx, solo_valores=False)
    wb_v = abrir_tabla(ctx, solo_valores=True)
    ws_f, ws_v = hoja_tabla(wb_f, ctx), hoja_tabla(wb_v, ctx)
    cab, cols_mes, filas = mapa_tabla(ws_f, ctx)
    c_mes = cols_mes[ctx["mes"]]
    previos = [m for m in MESES[:ctx["mes_num"] - 1] if m in cols_mes][-meses_previos:]

    objetivo = {clave_norm(k) for k in CLAVES}
    estado = []
    for k, (r, etiqueta) in filas.items():
        if k not in objetivo:
            continue
        bruto = ws_f.cell(r, c_mes).value
        formula = isinstance(bruto, str) and bruto.startswith("=")
        estado.append({
            "clave": etiqueta,
            "fila": r,
            "celda": coord(r - 1, c_mes - 1),
            "unidad": ws_f.cell(r, 2).value,
            "es_formula": formula,
            "formula": bruto if formula else None,
            "valor_actual": limpio(ws_v.cell(r, c_mes).value),
            "meses_previos": {m: limpio(ws_v.cell(r, cols_mes[m]).value) for m in previos},
        })
    no_encontradas = [k for k in CLAVES if clave_norm(k) not in filas]
    return {
        "hoja": ws_f.title,
        "fila_cabecera": cab,
        "columna_mes": coord(cab - 1, c_mes - 1)[:-len(str(cab))],
        "filas": estado,
        "claves_sin_fila": no_encontradas,
        "nota": "valor_actual de una fórmula es el último valor que guardó Excel; "
                "puede estar vacío o desfasado si el fichero se guardó sin recalcular.",
    }


def respaldar(ctx):
    RESPALDOS.mkdir(exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d-%H%M%S")
    destino = RESPALDOS / f"{ctx['ruta_tabla'].stem} {sello}.xlsx"
    shutil.copy2(ctx["ruta_tabla"], destino)
    copias = sorted(RESPALDOS.glob(f"{ctx['ruta_tabla'].stem} *.xlsx"))
    for vieja in copias[:-MAX_RESPALDOS]:
        vieja.unlink(missing_ok=True)
    return destino.name


def escribir(ctx, valores):
    if not isinstance(valores, dict) or not valores:
        raise ErrorAgente("'valores' debe ser un objeto {clave: número} no vacío.")
    permitidas = {clave_norm(k) for k in CLAVES}

    wb = abrir_tabla(ctx, solo_valores=False)
    ws = hoja_tabla(wb, ctx)
    _, cols_mes, filas = mapa_tabla(ws, ctx)
    c_mes = cols_mes[ctx["mes"]]

    escritas, omitidas = [], []
    for clave, valor in valores.items():
        k = clave_norm(clave)
        if k not in permitidas:
            omitidas.append({"clave": clave, "motivo": "no es una de las claves que rellena este proceso"})
            continue
        if k not in filas:
            omitidas.append({"clave": clave, "motivo": "no existe esa fila en la hoja"})
            continue
        v = num(valor)
        if v is None:
            omitidas.append({"clave": clave, "motivo": f"valor no numérico: {valor!r}"})
            continue
        r, etiqueta = filas[k]
        cell = ws.cell(r, c_mes)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            omitidas.append({"clave": etiqueta, "motivo": f"la celda {cell.coordinate} tiene fórmula ({cell.value}); no se pisa"})
            continue
        anterior = cell.value
        cell.value = int(v) if v.is_integer() else v
        escritas.append({"clave": etiqueta, "celda": cell.coordinate, "anterior": limpio(anterior), "nuevo": limpio(v)})

    respaldo = None
    if escritas:
        respaldo = respaldar(ctx)
        wb.save(ctx["ruta_tabla"])
    return {"escritas": escritas, "omitidas": omitidas, "respaldo": respaldo,
            "hoja": ws.title, "mes": ctx["mes"]}


# ---------------------------------------------------------------- volcados

def ver_hoja(ctx, hoja, desde=1, hasta=None, max_celdas=1500):
    wb = abrir_cd(ctx)
    nombre = hoja_fuzzy(wb, hoja)
    g = rejilla(wb[nombre])
    wb.close()
    hasta = min(hasta or len(g), len(g))
    lineas, total = [], 0
    for r in range(max(desde, 1) - 1, hasta):
        partes = []
        for c, v in enumerate(g[r]):
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            partes.append(f"{coord(r, c)}={limpio(v)}")
        if partes:
            lineas.append(" | ".join(partes))
            total += len(partes)
        if total >= max_celdas:
            lineas.append(f"... (cortado en la fila {r + 1}; pide 'desde': {r + 2} para seguir)")
            break
    return {"hoja": nombre, "filas_totales": len(g), "celdas": "\n".join(lineas)}


def ver_global(ctx, buscar_txt=None):
    filas, col_real, _ = filas_global(ctx)
    if buscar_txt:
        b = clave_norm(buscar_txt)
        filas = [f for f in filas if f["titulo"] is not None and b in clave_norm(f["titulo"])]
    return {"hoja": ctx["hoja_global"], "mes": ctx["mes"], "columna_real": col_real,
            "filas": [f for f in filas if f["titulo"] is not None][:400]}


# ---------------------------------------------------------------- comandos

def cmd_preparar(ctx, _):
    pub = contexto_publico(ctx)
    if not ctx["ruta_cd"].exists():
        raise ErrorAgente(f"No encuentro el CD '{ctx['cd']}' en la carpeta del script.")
    wb = abrir_cd(ctx)
    pub["hojas_cd"] = wb.sheetnames
    wb.close()

    v_cd, e_cd = extraccion_cd(ctx)
    v_gl, e_gl, aviso_global = extraccion_global(ctx)
    try:
        tabla = estado_tabla(ctx)
    except Exception as e:
        tabla = {"error": f"{type(e).__name__}: {e}"}

    return {
        "script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "contexto": pub,
        "extraccion": {"valores": {**v_cd, **v_gl}, "errores": {**e_cd, **e_gl}},
        "aviso_global": aviso_global,
        "tabla": tabla,
    }


COMANDOS = {
    "preparar": cmd_preparar,
    "ver_hoja": lambda ctx, a: ver_hoja(ctx, a.get("hoja", ""), int(a.get("desde") or 1),
                                        int(a["hasta"]) if a.get("hasta") else None),
    "ver_global": lambda ctx, a: ver_global(ctx, a.get("buscar")),
    "estado_tabla": lambda ctx, a: estado_tabla(ctx, int(a.get("meses_previos") or 3)),
    "escribir": lambda ctx, a: escribir(ctx, a.get("valores")),
    "limpiar": lambda ctx, a: {"borrado": (ctx["ruta_cd"].unlink(missing_ok=True) or True)},
}


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2 or sys.argv[1] not in COMANDOS:
        salir({"ok": False, "error": f"Uso: agente_ca.py <{'|'.join(COMANDOS)}> '<json>'"})
    try:
        args = json.loads(sys.argv[2]) if len(sys.argv) > 2 and sys.argv[2].strip() else {}
        if isinstance(args, str):  # el agente a veces manda el JSON doblemente serializado
            args = json.loads(args)
        if isinstance(args.get("valores"), str):
            args["valores"] = json.loads(args["valores"])
        ctx = contexto(args.get("cd"))
        salir({"ok": True, **COMANDOS[sys.argv[1]](ctx, args)})
    except ErrorAgente as e:
        salir({"ok": False, "error": str(e)})
    except Exception as e:
        salir({"ok": False, "error": f"{type(e).__name__}: {e}"})


if __name__ == "__main__":
    main()
