#!/usr/bin/env python3
"""
Lencería del CD para el informe del Consejo de Dirección.

Lee del CD de un polo las tablas de lencería que el informe necesita y las
devuelve en JSON, para guardarlas por mes en Supabase (`cd_lenceria`). No
escribe en ningún fichero: la TABLA CA no se toca (decisión de Leo,
2026-10-22). Diseño en vault/Informe-Consejo-Lenceria-CD.md.

    POST http://polarierauto_pythonrunner:8000/run
    {"script": "Leslie/CD/cd_lenceria.py", "args": ["leer", "<CD>"]}

    python cd_lenceria.py leer <CD> [--ruta <fichero>] [--polo LH --anio 2026 --mes 5]

<CD> es el nombre del fichero subido (<POLO>CA<AAAA><MM>.xlsx), que n8n deja
junto a este script, como para agente_ca.py. Para recargar meses pasados con
CD de nombre libre («HACA2026 mayo.xlsx») se pasan --ruta, --polo, --anio y
--mes.

Bloques (respuestas del formulario del 2026-10-28, P2–P10, P24–P25):

    montas_habitacion  «ESTADO DE LAS MONTAS POR HABITACIÓN», columna TOTAL
    montas_cliente     «ESTADO DE LAS MONTAS POR CLIENTE», columna TOTAL (+ TD)
    rechazo            «RESUMEN DE LOS RECHAZOS …» por prenda, columna SUMA
    mal_uso            «RESUMEN DEL MAL USO …», columna SUMA, sin tocar (P2)
    inyecciones        «RESUMEN (DE LAS) INYECCIONES …», columna SUMA
    desmanche          «RECUPERACIONES DE LENCERÍA»: Producción (= tratadas) y
                       Recuperado; las variantes «sin/con remojo» se suman
    retiro             «DESCARTES …» (Descarte generado en el mes) o
                       «RESUMEN DEL DESCARTE POR HOTELES» (SUMA): P6
    faltantes          resultados de los inventarios del mes: suma de las
                       diferencias negativas de los hoteles inventariados (P10)

Cada bloque sale con `estado`:
    ok              encontrado y rellenado (los ceros escritos cuentan)
    vacio           encontrado pero sin rellenar: celdas de hotel en blanco
                    (la SUMA da 0 por fórmula, pero no es un cero real)
    no_encontrado   el CD no tiene ese bloque

Los modelos de CD no son iguales en todos los polos (comprobado con los de
mayo de 2026): títulos distintos, hojas con sufijo («Lencería 2 MAYO»), el
desmanche partido en «sin remojo»/«con remojo», los inventarios en la hoja 1
y los códigos de prenda con dos significados (ver CODIGOS_*).
"""

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from agente_ca import (ALIAS_POLO, MAX_COLS, MAX_FILAS, POLOS, ErrorAgente, compacto,
                      contexto, coord, limpio, num, sin_tildes)

BLOQUES = ["montas_habitacion", "montas_cliente", "rechazo", "mal_uso",
           "inyecciones", "desmanche", "retiro", "faltantes"]

PRENDAS = ["sabana_personal", "sabana_king", "funda", "toalla_alfombra", "toalla_bano",
           "toalla_cara", "toalla_facial", "toalla_bidet", "toalla_piscina"]

# Clave de polo de Supabase (tabla_ca_meses, disponibilidad) por prefijo del CD.
CLAVE_POLO = {"LH": "habana", "VA": "varadero", "CA": "caibarien",
              "CC": "cayo_coco", "CCX": "cayo_cruz", "HOL": "holguin"}

# Los códigos de prenda no significan lo mismo en todos los CD. En el de
# Caibarién TC es la toalla de cuerpo (baño), TM la de mano (cara) y TB la
# bidet; en los de Varadero y La Habana TB es la de baño y TC la de cara.
# Se distingue por los códigos que trae el bloque: TM o ALF → modelo «cuerpo».
CODIGOS_CUERPO = {"SP": "sabana_personal", "SK": "sabana_king", "F": "funda",
                  "ALF": "toalla_alfombra", "TA": "toalla_alfombra", "TC": "toalla_bano",
                  "TM": "toalla_cara", "TF": "toalla_facial", "TB": "toalla_bidet",
                  "TP": "toalla_piscina"}
CODIGOS_BANO = {"SP": "sabana_personal", "SK": "sabana_king", "F": "funda",
                "TA": "toalla_alfombra", "ALF": "toalla_alfombra", "TB": "toalla_bano",
                "TC": "toalla_cara", "TF": "toalla_facial", "TP": "toalla_piscina"}


# ----------------------------------------------------------------- utilidades

def norm(s):
    """Texto normalizado: sin tildes, minúsculas, espacios simples."""
    return re.sub(r"\s+", " ", sin_tildes(str(s)).strip().lower())


def prenda_por_nombre(texto):
    """'Sabana Personal', 'Sábanas Personales', 'Toalla Baño'… → clave; None si no es prenda."""
    t = norm(texto)
    reglas = [(r"^sabanas? personal", "sabana_personal"), (r"^sabanas? king", "sabana_king"),
              (r"^fundas?\b", "funda"), (r"^(toallas? )?alfombra", "toalla_alfombra"),
              (r"^toallas? (de )?(bano|cuerpo)", "toalla_bano"),
              (r"^toallas? (de )?(cara|mano)", "toalla_cara"),
              (r"^toallas? facial", "toalla_facial"), (r"^toallas? bidet", "toalla_bidet"),
              (r"^toallas? (de )?piscina", "toalla_piscina")]
    for patron, clave in reglas:
        if re.search(patron, t):
            return clave
    return None


def tabla_codigos(codigos):
    """Elige el significado de los códigos según los que trae el bloque."""
    cods = {str(c).strip().upper() for c in codigos if c}
    return CODIGOS_CUERPO if cods & {"TM", "ALF"} else CODIGOS_BANO


class Rejilla:
    """Una hoja como lista de filas, con búsqueda de títulos."""

    def __init__(self, nombre, filas):
        self.nombre = nombre
        self.g = filas

    def v(self, r, c):
        if 0 <= r < len(self.g) and 0 <= c < len(self.g[r]):
            return self.g[r][c]
        return None

    def t(self, r, c):
        x = self.v(r, c)
        return norm(x) if isinstance(x, str) and x.strip() else ""

    def ref(self, r, c):
        return coord(r, c)

    def titulos(self, patron):
        """Celdas cuyo texto normalizado empieza por el patrón (regex)."""
        rx = re.compile(patron)
        return [(r, c) for r, fila in enumerate(self.g) for c, x in enumerate(fila)
                if isinstance(x, str) and rx.match(norm(x))]


def hojas_lenceria(wb):
    """Rejillas de las hojas de lencería del CD («5.Lenceria 1 (2)», «Lencería 2 MAYO»…)."""
    out = []
    for nombre in wb.sheetnames:
        if "lenceria" in compacto(nombre):
            filas = [list(f) for f in wb[nombre].iter_rows(max_row=MAX_FILAS, max_col=MAX_COLS,
                                                          values_only=True)]
            out.append(Rejilla(nombre, filas))
    return out


def resultado(estado, valores=None, totales=None, hoja=None, celda=None, avisos=None):
    return {"estado": estado, "valores": valores or {}, "totales": totales or {},
            "hoja": hoja, "celda": celda, "avisos": avisos or []}


def no_encontrado(que):
    return resultado("no_encontrado", avisos=[f"El CD no tiene el bloque «{que}»."])


# ------------------------------------------------------- bloques por hoteles

def fila_cabecera(h, r0, c0, palabras=("lenceria", "prenda")):
    """Fila de cabecera del bloque: la primera de las 3 siguientes con «Lencería»/«Prenda»."""
    for r in range(r0 + 1, r0 + 4):
        if any(h.t(r, c0).startswith(p) for p in palabras):
            return r
    return None


def col_cabecera(h, r, c0, *nombres, hasta=None):
    """Primera columna a la derecha de c0 cuya cabecera empieza por alguno de los nombres."""
    fin = min(len(h.g[r]) if r < len(h.g) else 0, hasta or MAX_COLS)
    for c in range(c0 + 1, fin):
        t = h.t(r, c)
        if t and any(t.startswith(n) for n in nombres):
            return c
    return None


def filas_prenda(h, r_cab, c0):
    """[(fila, prenda)] debajo de la cabecera, hasta TOTAL o la primera fila sin prenda."""
    out = []
    for r in range(r_cab + 1, r_cab + 16):
        t = h.t(r, c0)
        if t.startswith("total"):
            break
        p = prenda_por_nombre(t) if t else None
        if p:
            out.append((r, p))
        elif t and out:
            break
    return out


def leer_por_hoteles(h, r0, c0):
    """Bloque «RESUMEN … POR HOTELES»: prendas en filas, hoteles en columnas y SUMA.

    Devuelve (resultado, hay_prendas). Vacío si ninguna celda de hotel tiene número.
    """
    r_cab = fila_cabecera(h, r0, c0)
    if r_cab is None:
        return None, False
    c_suma = col_cabecera(h, r_cab, c0, "suma", "total")
    if c_suma is None:
        return resultado("no_encontrado", hoja=h.nombre, celda=h.ref(r0, c0),
                         avisos=[f"Sin columna SUMA en {h.nombre}!{h.ref(r_cab, c0)}."]), False
    filas = filas_prenda(h, r_cab, c0)
    if not filas:
        return None, False
    valores, escrito = {}, False
    for r, p in filas:
        if any(num(h.v(r, c)) is not None for c in range(c0 + 1, c_suma)):
            escrito = True
        valores[p] = {"cantidad": limpio(num(h.v(r, c_suma)))}
    if not escrito:
        return resultado("vacio", hoja=h.nombre, celda=h.ref(r0, c0),
                         avisos=["Bloque sin rellenar: las celdas de los hoteles están en blanco."]), True
    return resultado("ok", valores, hoja=h.nombre, celda=h.ref(r0, c0)), True


def bloque_por_hoteles(hojas, patron, nombre, preferir=None):
    """Busca el bloque por título y lo lee. Si hay varios con prendas, gana `preferir`
    y, si no, el primero rellenado (avisando)."""
    candidatos = []
    for h in hojas:
        for (r0, c0) in h.titulos(patron):
            res, con_prendas = leer_por_hoteles(h, r0, c0)
            if res is not None and con_prendas:
                candidatos.append((norm(h.v(r0, c0)), res))
    if not candidatos:
        return no_encontrado(nombre)
    if len(candidatos) > 1 and preferir:
        preferidos = [c for c in candidatos if re.search(preferir, c[0])]
        if preferidos:
            candidatos = preferidos + [c for c in candidatos if c not in preferidos]
    rellenados = [c for c in candidatos if c[1]["estado"] == "ok"]
    elegido = (rellenados or candidatos)[0][1]
    if len(rellenados) > 1:
        elegido["avisos"].append(f"Hay {len(rellenados)} bloques «{nombre}» rellenados; "
                                 f"se usa el de {elegido['hoja']}!{elegido['celda']}.")
    return elegido


# ------------------------------------------------------------------- montas

def bloque_montas(hojas, patron, nombre, con_td=False):
    for h in hojas:
        for (r0, c0) in h.titulos(patron):
            r_cab = fila_cabecera(h, r0, c0)
            if r_cab is None:
                continue
            c_tot = col_cabecera(h, r_cab, c0, "total", "promedio")
            if c_tot is None:  # sin cabecera: la última columna con texto de la fila
                textos = [c for c in range(c0 + 1, len(h.g[r_cab])) if h.t(r_cab, c)]
                c_tot = textos[-1] if textos else None
            if c_tot is None:
                continue
            valores, totales = {}, {}
            for r in range(r_cab + 1, r_cab + 14):
                t = h.t(r, c0)
                if t.startswith("td total"):
                    totales["td_totales"] = limpio(num(h.v(r, c_tot)))
                elif t.startswith("td promedio"):
                    totales["td_promedio"] = limpio(num(h.v(r, c_tot)))
                else:
                    p = prenda_por_nombre(t) if t else None
                    if p:
                        valores[p] = {"indice": limpio(num(h.v(r, c_tot)))}
                    elif t and valores:
                        break
            if not valores:
                continue
            avisos = [f"Columna del total: {h.ref(r_cab, c_tot)} «{h.v(r_cab, c_tot)}»."]
            if all(v["indice"] is None for v in valores.values()):
                # El total da #DIV/0! cuando algún hotel no tiene turistas-día (CD de
                # Caibarién, mayo de 2026). Si un solo hotel trae datos, su índice es el del polo.
                filas = {p: r for r in range(r_cab + 1, r_cab + 14)
                         for p in [prenda_por_nombre(h.t(r, c0)) if h.t(r, c0) else None] if p}
                con_datos = [c for c in range(c0 + 1, c_tot)
                             if any(num(h.v(r, c)) for r in filas.values())]
                if len(con_datos) == 1:
                    c1 = con_datos[0]
                    for p, r in filas.items():
                        valores[p] = {"indice": limpio(num(h.v(r, c1)))}
                    avisos.append(f"El total no trae valores; solo un hotel tiene datos "
                                  f"({h.v(r_cab, c1)}, {h.ref(r_cab, c1)}) y se usa el suyo.")
                    return resultado("ok", valores, totales if con_td else None, h.nombre,
                                     h.ref(r0, c0), avisos)
                return resultado("vacio", totales=totales if con_td else None, hoja=h.nombre,
                                 celda=h.ref(r0, c0), avisos=avisos + ["La columna del total no trae ningún valor."])
            return resultado("ok", valores, totales if con_td else None, h.nombre, h.ref(r0, c0), avisos)
    return no_encontrado(nombre)


# ---------------------------------------------------------------- desmanche

def leer_recuperaciones(h, r0, c0):
    """Bloque con códigos de prenda: Prenda | Producción | Recuperado | % Recuperación."""
    r_cab = fila_cabecera(h, r0, c0, palabras=("prenda",))
    if r_cab is None:
        return None
    c_prod = col_cabecera(h, r_cab, c0, "produccion")
    c_rec = col_cabecera(h, r_cab, c0, "recuperado")
    if c_prod is None or c_rec is None:
        return None
    filas = []
    for r in range(r_cab + 1, r_cab + 14):
        x = h.v(r, c0)
        if not isinstance(x, str) or not x.strip() or norm(x).startswith("total"):
            break
        filas.append((r, x.strip().upper()))
    codigos = tabla_codigos([c for _, c in filas])
    out = {}
    for r, cod in filas:
        p = codigos.get(cod)
        if p:
            out[p] = (num(h.v(r, c_prod)), num(h.v(r, c_rec)))
    return out


def bloque_desmanche(hojas):
    partes, origen, avisos = [], None, []
    for h in hojas:
        for (r0, c0) in h.titulos(r"^recuperaciones de lenceria"):
            titulo = norm(h.v(r0, c0))
            if "mal uso" in titulo:  # «… MAL USO HOTEL» no es desmanche
                continue
            datos = leer_recuperaciones(h, r0, c0)
            if datos is None:
                continue
            partes.append((titulo, datos))
            origen = origen or (h.nombre, h.ref(r0, c0))
    if not partes:
        return no_encontrado("Recuperaciones de lencería")
    if len(partes) > 1:
        avisos.append("Se suman " + " + ".join(f"«{t}»" for t, _ in partes) + ".")
    valores = {}
    for _, datos in partes:
        for p, (trat, rec) in datos.items():
            v = valores.setdefault(p, {"tratadas": None, "recuperadas": None})
            if trat is not None:
                v["tratadas"] = (v["tratadas"] or 0) + trat
            if rec is not None:
                v["recuperadas"] = (v["recuperadas"] or 0) + rec
    if not any((v["tratadas"] or 0) or (v["recuperadas"] or 0) for v in valores.values()):
        return resultado("vacio", hoja=origen[0], celda=origen[1],
                         avisos=avisos + ["Bloque sin rellenar: producción y recuperado vacíos o a cero."])
    for v in valores.values():
        v["pct"] = limpio(v["recuperadas"] / v["tratadas"]) if v["tratadas"] and v["recuperadas"] is not None else None
        v["tratadas"], v["recuperadas"] = limpio(v["tratadas"]), limpio(v["recuperadas"])
    return resultado("ok", valores, hoja=origen[0], celda=origen[1], avisos=avisos)


# ------------------------------------------------------------------- retiro

def bloque_retiro(hojas):
    # Modelo de Varadero y Caibarién: «DESCARTES - INVENTARIOS DE DESCARTE» con códigos.
    for h in hojas:
        for (r0, c0) in h.titulos(r"^descartes"):
            r_cab = fila_cabecera(h, r0, c0, palabras=("prenda",))
            if r_cab is None:
                continue
            cols = [c for c in range(c0 + 1, min(len(h.g[r_cab]), c0 + 12))
                    if h.t(r_cab, c).startswith("descarte generado")]
            if not cols:
                continue
            filas = []
            for r in range(r_cab + 1, r_cab + 14):
                x = h.v(r, c0)
                if not isinstance(x, str) or not x.strip() or norm(x).startswith("total"):
                    break
                filas.append((r, x.strip().upper()))
            codigos = tabla_codigos([c for _, c in filas])

            def columna(c):
                return {codigos[cod]: num(h.v(r, c)) for r, cod in filas if cod in codigos}

            avisos = []
            del_mes = next((c for c in cols if "mes" in h.t(r_cab, c)), cols[0])
            elegida = del_mes
            datos = columna(del_mes)
            # En el CD de Varadero de mayo de 2026 la columna «en el mes» trae negativos
            # (ajustes de inventario) y el mes va en la rotulada «… 2025».
            if any(v is not None and v < 0 for v in datos.values()):
                for c in cols:
                    otra = columna(c)
                    if c != del_mes and all(v is None or v >= 0 for v in otra.values()) \
                            and any(v for v in otra.values()):
                        avisos.append(f"«{h.v(r_cab, del_mes)}» trae negativos; se usa "
                                      f"«{h.v(r_cab, c)}» ({h.ref(r_cab, c)}).")
                        elegida, datos = c, otra
                        break
                else:
                    avisos.append(f"«{h.v(r_cab, del_mes)}» trae negativos.")
            if all(v is None for v in datos.values()):
                return resultado("vacio", hoja=h.nombre, celda=h.ref(r0, c0),
                                 avisos=avisos + ["Bloque sin rellenar."])
            valores = {p: {"cantidad": limpio(v)} for p, v in datos.items()}
            return resultado("ok", valores, hoja=h.nombre, celda=h.ref(r0, c0),
                             avisos=avisos + [f"Columna: {h.ref(r_cab, elegida)}."])
    # Modelo de La Habana: «RESUMEN DEL DESCARTE POR HOTELES», con SUMA.
    return bloque_por_hoteles(hojas, r"^resumen del descarte", "Descartes")


# ---------------------------------------------------------------- faltantes

def bloque_faltantes(hojas, mes_nombre=None):
    """Resultados de los inventarios: grupos Inicial / Real / Diferencia / Importe por hotel.

    Faltantes = suma de las diferencias negativas de los hoteles inventariados en el
    mes (P10); las positivas son sobrantes. Inventariado = su fila «Real» o su fila
    «Diferencia» tiene algún número distinto de cero.
    """
    for h in hojas:
        for r, fila in enumerate(h.g):
            for c, x in enumerate(fila):
                if not (isinstance(x, str) and norm(x) == "inicial"):
                    continue
                if h.t(r + 1, c) != "real" or h.t(r + 2, c) != "diferencia":
                    continue
                # Es el primer grupo del bloque: encima, la fila de códigos de prenda.
                r_cod = None
                for rr in range(r - 1, max(r - 4, -1), -1):
                    cods = [h.v(rr, cc) for cc in range(c + 1, c + 12)]
                    if sum(1 for k in cods if isinstance(k, str) and k.strip().upper() in CODIGOS_CUERPO) >= 5:
                        r_cod = rr
                        break
                if r_cod is None:
                    continue
                cols = [(cc, str(h.v(r_cod, cc)).strip().upper()) for cc in range(c + 1, c + 12)
                        if isinstance(h.v(r_cod, cc), str) and h.v(r_cod, cc).strip()]
                codigos = tabla_codigos([k for _, k in cols])
                cols = [(cc, codigos[k]) for cc, k in cols if k in codigos]
                return leer_inventarios(h, r, c, r_cod, cols, mes_nombre)
    return no_encontrado("Resultados de los inventarios realizados")


def leer_inventarios(h, r_ini, c, r_cod, cols, mes_nombre):
    faltan = {p: 0.0 for _, p in cols}
    sobran = {p: 0.0 for _, p in cols}
    hoteles, importe, avisos = [], 0.0, []
    etiqueta_mes = next((h.v(r_cod, cc) for cc in range(c - 1, c + 1)
                         if isinstance(h.v(r_cod, cc), str) and h.v(r_cod, cc).strip()), None)
    if mes_nombre and etiqueta_mes and norm(etiqueta_mes) != norm(mes_nombre):
        avisos.append(f"El bloque de inventarios dice «{etiqueta_mes}», no {mes_nombre}.")
    # Cada hotel ocupa cuatro filas: Inicial, Real, Diferencia e Importe. El nombre del
    # hotel va a la izquierda de «Inicial» y la etiqueta «Importe», a la izquierda de su valor.
    r = r_ini
    while h.t(r, c) == "inicial" and h.t(r + 1, c) == "real" and h.t(r + 2, c) == "diferencia":
        con_importe = h.t(r + 3, c - 1) == "importe"
        # Inventariado: trae «Real» o, como en el CD de Caibarién, solo la diferencia tecleada.
        if any(num(h.v(r + 1, cc)) or num(h.v(r + 2, cc)) for cc, _ in cols):
            nombre = h.v(r, c - 1)
            hoteles.append(nombre.strip() if isinstance(nombre, str) and nombre.strip() else h.ref(r, c))
            for cc, p in cols:
                d = num(h.v(r + 2, cc))
                if d is not None and d < 0:
                    faltan[p] += -d
                elif d:
                    sobran[p] += d
            if con_importe:
                importe += num(h.v(r + 3, c)) or 0
        r += 4 if con_importe else 3
    valores = {p: {"faltantes": limpio(faltan[p]), "sobrantes": limpio(sobran[p])} for _, p in cols}
    if not hoteles:
        avisos.append("Ningún hotel inventariado en el mes: faltantes a cero.")
    return resultado("ok", valores, {"hoteles": hoteles, "importe": limpio(importe)},
                     h.nombre, h.ref(r_ini, c), avisos)


# -------------------------------------------------------------------- leer

def leer_cd(ruta, mes_nombre=None):
    wb = load_workbook(ruta, data_only=True, read_only=True)
    hojas = hojas_lenceria(wb)
    if not hojas:
        raise ErrorAgente(f"El CD no tiene hojas de lencería. Hojas: {wb.sheetnames}")
    bloques = {
        "montas_habitacion": bloque_montas(hojas, r"^estado de las montas por habitacion",
                                           "Estado de las montas por habitación"),
        "montas_cliente": bloque_montas(hojas, r"^estado de las montas por cliente",
                                        "Estado de las montas por cliente", con_td=True),
        "rechazo": bloque_por_hoteles(hojas, r"^resumen de los rechazos", "Rechazos por prenda",
                                      preferir=r"tipo de prenda"),
        "mal_uso": bloque_por_hoteles(hojas, r"^resumen del mal uso", "Mal uso"),
        "inyecciones": bloque_por_hoteles(hojas, r"^resumen (de las )?inyecciones", "Inyecciones"),
        "desmanche": bloque_desmanche(hojas),
        "retiro": bloque_retiro(hojas),
        "faltantes": bloque_faltantes(hojas, mes_nombre),
    }
    return bloques


def principal(argv):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("leer")
    p.add_argument("cd")
    p.add_argument("--ruta", help="fichero del CD, si no está junto al script")
    p.add_argument("--polo", help="prefijo del polo (LH, VA, CA, CC, CCX, HOL) si el nombre no lo dice")
    p.add_argument("--anio", type=int)
    p.add_argument("--mes", type=int)
    a = ap.parse_args(argv)

    if a.polo or a.anio or a.mes:
        if not (a.polo and a.anio and a.mes):
            raise ErrorAgente("--polo, --anio y --mes van juntos")
        prefijo = ALIAS_POLO.get(a.polo.upper(), a.polo.upper())
        if prefijo not in POLOS:
            raise ErrorAgente(f"Polo desconocido: {a.polo}")
        anio, mes = a.anio, a.mes
        ruta = Path(a.ruta or a.cd)
    else:
        ctx = contexto(a.cd)
        prefijo, anio, mes = ctx["prefijo"], ctx["anio"], ctx["mes_num"]
        ruta = Path(a.ruta) if a.ruta else ctx["ruta_cd"]
    if not ruta.exists():
        raise ErrorAgente(f"No encuentro el fichero '{ruta.name}'.")
    from agente_ca import MESES
    bloques = leer_cd(ruta, MESES[mes - 1])
    return {"ok": True, "anio": anio, "mes": mes, "polo": CLAVE_POLO[prefijo],
            "fichero": Path(a.cd).name, "bloques": bloques}


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        salida = principal(sys.argv[1:])
    except ErrorAgente as e:
        salida = {"ok": False, "error": str(e)}
    except Exception as e:  # noqa: BLE001 — n8n tiene que leer el motivo
        salida = {"ok": False, "error": f"{type(e).__name__}: {e}"}
    sys.stdout.write(json.dumps(salida, ensure_ascii=False) + "\n")
