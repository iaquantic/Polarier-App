#!/usr/bin/env python3
"""
Herramientas del "Agente Modelo 1" (n8n + Claude). Sustituye a Modelo1.py y a
toda la rama CP de n8n (Sheets de staging en Drive, copia del GLOBAL en Drive).

El agente de n8n las llama por el pythonrunner:

    POST http://polarierauto_pythonrunner:8000/run
    {"script": "Leslie/agente_modelo1.py", "args": ["<comando>", "<json>"]}

Cada subida del chat es un "lote" (n8n usa el id de la ejecución). n8n deja el
fichero subido en Leslie/m1_<lote>__<nombre> (su nodo de escritura no crea
carpetas) y 'preparar' lo mueve a modelo1_lotes/<lote>/entrada/. Después:

    preparar     Desempaqueta (PDF suelto, ZIP o RAR), lee cada PDF por
                 coordenadas, deduce polo y mes DEL CONTENIDO, valida cada fila
                 con las identidades del propio PDF y compara con el GLOBAL.
                 Deja la extracción en el lote para que la escritura use
                 exactamente lo que el agente revisó.
    ver_filas    Filas extraídas de un PDF (todas o filtradas), con su estado.
    ver_pdf      Trozos de texto en bruto del PDF alrededor de un código, con
                 coordenadas, para comprobar a mano una fila dudosa.
    ver_global   Filas del GLOBAL para el polo y el mes de un PDF.
    escribir     Escribe en el GLOBAL los PDF aceptados, con las correcciones o
                 exclusiones que decida el agente. Copia de seguridad antes.
    verificar    Revisión final e independiente: vuelve a leer los PDF desde
                 cero y el GLOBAL desde disco y lo compara todo.
    cerrar       Devuelve el fichero a entregar en el chat (el GLOBAL escrito, o
                 un ZIP si el lote tocó GLOBAL de varios años) y borra el lote.

Reglas (ver vault/Generacion-Documentos.md, D1):
  - El polo y el mes salen del CONTENIDO del PDF ("Unidad: ... Lavanderia X" y
    "<Mes>/<AAAA>" de la cabecera visible). El nombre del fichero solo se usa
    para avisar si no coincide. Se ignora el texto fuera de la página.
  - Al GLOBAL va el bloque TOTAL (Plan y Real) de la fila del mes (M / MES).
    El cruce es por código ("Filas" del GLOBAL = código de 3 dígitos del PDF),
    nunca por nombre: hay nombres repetidos (Lavandería es 23, 33 y 72).
  - Año: GLOBAL-CP<aa>.xlsx, hoja <COD>CP<AAAA>. Nada fijo.
  - Celdas con fórmula: si la fórmula ya da el valor del PDF (±1 CUP de
    redondeo) se deja; si no cuadra o no se puede evaluar, se escribe el valor
    del PDF y se avisa.
  - Si el mes ya tenía valores, se sobrescriben y se devuelve "antes".

Toda salida es JSON por stdout. Los errores salen como {"ok": false, "error": ...}
con código 0, para que el agente pueda leerlos y razonar.
"""

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent            # /data/Polarier/Leslie
LOTES = BASE_DIR / "modelo1_lotes"
SALIDA = BASE_DIR / "modelo1_salida"                   # ZIP cuando un lote toca varios GLOBAL
RESPALDOS = BASE_DIR / "respaldos_global"
MAX_RESPALDOS = 20
TOL = 0.015            # tolerancia en valores con 2 decimales
TOL_PCT = 0.015        # tolerancia en porcentajes impresos con 2 decimales
TOL_FORMULA = 1.0      # una fórmula que suma decenas de partidas puede irse unos céntimos del total
                       # impreso por redondeo; hasta 1 CUP se considera que cuadra y se conserva

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
MES_NUM = {m: i + 1 for i, m in enumerate(MESES)}
MES_NUM["SETIEMBRE"] = 9

# código de hoja en el GLOBAL -> (nombre, marcas en "Unidad:", prefijos del nombre de fichero)
POLOS = {
    "LH": ("La Habana", ("ASTER", "HABANA"), ("HA", "LH")),
    "VAR": ("Varadero", ("VARADERO",), ("VA",)),
    "CA": ("Caibarién", ("CENTRO", "CAIBARIEN"), ("CA",)),
    "CC": ("Cayo Coco", ("COCO",), ("CC",)),
    "CCZ": ("Cayo Cruz", ("CRUZ",), ("CX", "CCZ", "CCX")),
    "HOL": ("Holguín", ("HOLGUIN",), ("HO", "HOL")),
}

RE_NUM = re.compile(r"^\(?-?\d{1,3}(,\d{3})*(\.\d+)?\)?$|^\(?-?\d+(\.\d+)?\)?$")
RE_MASK = re.compile(r"^[*,.]+$")
RE_MES = re.compile(r"\b(" + "|".join(MESES + ["SETIEMBRE"]) + r")\s*/\s*(\d{4})\b", re.I)
RE_CODIGO = re.compile(r"^\d{3}$")
MARCAS_FILA = {"M": "M", "A": "A", "MES": "M", "ACU": "A"}


# ----------------------------------------------------------------- utilidades

def sin_tildes(s):
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c))


def mayus(s):
    return re.sub(r"\s+", " ", sin_tildes(s).upper()).strip()


def limpio(v):
    if isinstance(v, float):
        if v != v:
            return None
        return int(v) if v.is_integer() else round(v, 6)
    if isinstance(v, datetime):
        return v.isoformat()
    return v


def a_num(t):
    """'1,234.56' -> 1234.56 ; '(5.00)' -> -5.0 ; máscara '**,***' -> 'MASK'."""
    t = t.strip()
    if RE_MASK.match(t) and "*" in t:
        return "MASK"
    if not RE_NUM.match(t):
        return None
    neg = t.startswith("(") and t.endswith(")")
    v = float(t.strip("()").replace(",", ""))
    return -v if neg else v


def iguales(a, b, tol=TOL):
    if a is None or b is None or a == "MASK" or b == "MASK":
        return None
    return abs(float(a) - float(b)) <= tol


def salir(obj):
    print(json.dumps(obj, ensure_ascii=False, default=str))
    sys.exit(0)


def error(msg, **extra):
    salir({"ok": False, "error": msg, **extra})


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(65536), b""):
            h.update(b)
    return h.hexdigest()


def nombre_seguro(n):
    n = sin_tildes(os.path.basename(str(n)))
    return re.sub(r"[^A-Za-z0-9_. -]", "_", n).strip() or "fichero"


def dir_lote(args):
    lote = re.sub(r"[^A-Za-z0-9_-]", "", str(args.get("lote", "")))
    if not lote:
        error("Falta 'lote'.")
    return LOTES / lote


# ------------------------------------------------------ desempaquetar el lote

def desempaquetar(lote_dir):
    """Deja todos los PDF del lote en lote/pdf/. Devuelve (pdfs, avisos)."""
    entrada = lote_dir / "entrada"
    destino = lote_dir / "pdf"
    shutil.rmtree(destino, ignore_errors=True)  # preparar se puede repetir sin duplicar PDF
    destino.mkdir(parents=True)
    avisos = []
    ficheros = sorted(p for p in entrada.iterdir() if p.is_file()) if entrada.exists() else []
    if not ficheros:
        raise FileNotFoundError(f"No hay ningún fichero en {entrada}")

    def guardar(nombre, datos):
        n = nombre_seguro(nombre)
        if not n.lower().endswith(".pdf"):
            return
        p = destino / n
        i = 2
        while p.exists():
            p = destino / f"{Path(n).stem}_{i}.pdf"
            i += 1
        p.write_bytes(datos)

    for f in ficheros:
        cab = f.read_bytes()[:8]
        ext = f.suffix.lower()
        if cab.startswith(b"%PDF") or ext == ".pdf":
            guardar(f.name, f.read_bytes())
        elif cab.startswith(b"PK") or ext == ".zip":
            with zipfile.ZipFile(f) as z:
                for info in z.infolist():
                    n = info.filename
                    if info.is_dir() or "__MACOSX" in n or os.path.basename(n).startswith("._"):
                        continue
                    if n.lower().endswith(".pdf"):
                        guardar(os.path.basename(n), z.read(info))
                    elif n.lower().endswith((".zip", ".rar")):
                        avisos.append(f"Se ignora el comprimido anidado {n}: sube los PDF sueltos o en un solo ZIP.")
                    else:
                        avisos.append(f"Se ignora {n}: no es un PDF.")
        elif cab.startswith(b"Rar!") or ext == ".rar":
            extraer_rar(f, destino, avisos)
        else:
            avisos.append(f"Se ignora {f.name}: no es PDF, ZIP ni RAR.")

    # lo que haya dejado un extractor externo con subcarpetas
    for p in list(destino.rglob("*")):
        if p.is_file() and p.parent != destino:
            if p.suffix.lower() == ".pdf" and "__MACOSX" not in str(p):
                guardar(p.name, p.read_bytes())
            p.unlink()
    for p in sorted(destino.rglob("*"), reverse=True):
        if p.is_dir():
            shutil.rmtree(p, ignore_errors=True)
    for p in destino.iterdir():
        if p.suffix.lower() != ".pdf":
            p.unlink()
    pdfs = sorted(destino.glob("*.pdf"))
    return pdfs, avisos


def extraer_rar(f, destino, avisos):
    try:
        import rarfile  # noqa: F401
        with rarfile.RarFile(str(f)) as r:
            for info in r.infolist():
                if not info.is_dir() and info.filename.lower().endswith(".pdf"):
                    (destino / nombre_seguro(info.filename)).write_bytes(r.read(info))
        return
    except ImportError:
        pass
    except Exception as e:  # rarfile presente pero sin backend
        avisos.append(f"rarfile no pudo abrir el RAR ({e}); se prueba con herramientas del sistema.")
    for cmd in (["bsdtar", "-xf", str(f), "-C", str(destino)],
                ["unar", "-q", "-f", "-o", str(destino), str(f)],
                ["7z", "x", "-y", f"-o{destino}", str(f)],
                ["unrar", "x", "-o+", str(f), str(destino) + os.sep]):
        if shutil.which(cmd[0]):
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
            if r.returncode == 0:
                return
            avisos.append(f"{cmd[0]} falló: {r.stderr.strip()[:200]}")
    raise RuntimeError("El servidor no puede abrir ficheros RAR (no hay rarfile, bsdtar, unar, 7z ni unrar). "
                       "Sube los PDF en un ZIP o sueltos.")


# ------------------------------------------------------------ lectura del PDF

def trozos_pagina(page):
    """Trozos de texto en orden de dibujo, con (x, y). Ignora vacíos."""
    out = []

    def visitor(text, cm, tm, fd, fs):
        if text and text.strip():
            x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
            y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
            out.append({"t": text.replace("\n", " ").strip(), "x": round(x, 1), "y": round(y, 1)})

    page.extract_text(visitor_text=visitor)
    return out


def tokens_de(trozos):
    """Parte cada trozo en palabras (heredan x/y del trozo) y une números
    partidos por el salto de línea ('132,665,142.1' + '7')."""
    toks = []
    for i, tr in enumerate(trozos):
        for w in tr["t"].split():
            toks.append({"t": w, "x": tr["x"], "y": tr["y"], "trozo": i})
    unidos = []
    for tk in toks:
        prev = unidos[-1] if unidos else None
        if (prev and re.fullmatch(r"\d", tk["t"]) and re.search(r"\d\.\d$", prev["t"])
                and prev["trozo"] != tk["trozo"]):
            prev["t"] += tk["t"]
            prev["unido"] = True
            continue
        unidos.append(dict(tk))
    return unidos


def maquetado(trozos):
    """L12: MLC/MN/TOTAL con Real 2024 · L9: MLC/MN/TOTAL sin Real 2024 · L3: solo TOTAL."""
    texto = " ".join(t["t"] for t in trozos[:40]).upper()
    if "M.L.C" in texto or "MONEDA" in texto:
        return "L12" if "2024" in texto or re.search(r"REAL\s*20\d\d", texto) else "L9"
    return "L3"


def leer_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        # pypdf es Python puro: si el runner no lo trae, vale el wheel suelto en Leslie/
        # (pypdf-*.whl, Python lo importa sin descomprimir) o el paquete en Leslie/_vendor/pypdf
        for p in sorted(BASE_DIR.glob("pypdf-*.whl"), reverse=True) + [BASE_DIR / "_vendor"]:
            sys.path.insert(0, str(p))
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("El pythonrunner no tiene pypdf. Instálalo en el runner o deja el wheel "
                               f"pypdf-<versión>-py3-none-any.whl en {BASE_DIR}.")
    reader = PdfReader(str(path))
    paginas = []
    for i, page in enumerate(reader.pages):
        W, H = float(page.mediabox.width), float(page.mediabox.height)
        trozos = trozos_pagina(page)
        dentro = [t for t in trozos if -120 <= t["x"] <= W + 5 and -120 <= t["y"] <= H + 5]
        fuera = [t for t in trozos if t not in dentro]
        paginas.append({"n": i + 1, "trozos": dentro, "fuera": fuera, "W": W, "H": H})
    return paginas, {k: str(v) for k, v in (reader.metadata or {}).items()}


def cabecera(paginas):
    """Unidad, polo y mes de la cabecera visible, con avisos."""
    avisos = []
    unidades, meses, ocultos = {}, {}, {}
    for pg in paginas:
        for t in pg["trozos"][:25]:
            if t["t"].upper().startswith("UNIDAD"):
                unidades[t["t"]] = unidades.get(t["t"], 0) + 1
        for t in pg["trozos"]:
            for m in RE_MES.finditer(t["t"]):
                k = (MES_NUM[m.group(1).upper()], int(m.group(2)))
                meses[k] = meses.get(k, 0) + 1
        for t in pg["fuera"]:
            if RE_MES.search(t["t"]):
                ocultos.setdefault(t["t"], []).append(pg["n"])
    for texto, pags in ocultos.items():
        avisos.append(f"Hay texto oculto fuera de la hoja con otra fecha ('{texto}', páginas {pags}). "
                      f"Se ignora: el PDF parece editado después de emitirse.")
    unidad = max(unidades, key=unidades.get) if unidades else None
    if len(unidades) > 1:
        avisos.append(f"El PDF mezcla varias unidades: {list(unidades)}")
    mes = anio = None
    if meses:
        (mes, anio), n = max(meses.items(), key=lambda kv: kv[1])
        if len(meses) > 1:
            avisos.append("La cabecera visible tiene más de un mes: "
                          + ", ".join(f"{MESES[m - 1].title()}/{a} ({c} veces)" for (m, a), c in meses.items()))
    polo = polo_por_unidad(unidad)
    return {"unidad": unidad, "polo": polo, "mes": mes, "anio": anio,
            "meses_distintos": len(meses)}, avisos


def polo_por_unidad(unidad):
    if not unidad:
        return None
    u = mayus(unidad)
    hits = [cod for cod, (_, marcas, _) in POLOS.items() if any(m in u for m in marcas)]
    return hits[0] if len(hits) == 1 else None


def polo_por_nombre(nombre):
    n = mayus(Path(nombre).stem)
    m = re.match(r"^(HOL|HO|VA|CA|HA|LH|CCX|CCZ|CX|CC)CP", n.replace(" ", ""))
    if m:
        pref = m.group(1)
        for cod, (_, _, prefijos) in POLOS.items():
            if pref in prefijos:
                return cod
    m = re.search(r"MODELO\s*1\s+(.+)$", n)
    if m:
        resto = m.group(1)
        for cod, (nombre_polo, marcas, _) in POLOS.items():
            if any(x in resto for x in marcas + (mayus(nombre_polo),)):
                return cod
    return None


def mes_por_nombre(nombre):
    n = mayus(Path(nombre).stem).replace(" ", "")
    m = re.search(r"CP(20\d{2})(0[1-9]|1[0-2])", n)
    if m:
        return int(m.group(2)), int(m.group(1))
    for nombre_mes, num in MES_NUM.items():
        if nombre_mes in n:
            a = re.search(r"(20\d{2})", n)
            return num, int(a.group(1)) if a else None
    return None


def filas_pdf(paginas):
    """Filas del PDF: código, indicador y los valores del mes (M) y acumulado (A)."""
    filas, orden = [], 0
    for pg in paginas:
        lay = maquetado(pg["trozos"])
        toks = tokens_de(pg["trozos"])
        # localizar el inicio de cada fila: un código de 3 dígitos al margen izquierdo
        inicios = [i for i, tk in enumerate(toks)
                   if RE_CODIGO.match(tk["t"]) and tk["x"] < 40
                   and (i == 0 or toks[i - 1]["trozo"] != tk["trozo"] or toks[i - 1]["t"] in MARCAS_FILA)]
        # cabecera: lo anterior a la primera fila no es de nadie
        for k, ini in enumerate(inicios):
            fin = inicios[k + 1] if k + 1 < len(inicios) else len(toks)
            bloque = toks[ini:fin]
            codigo = int(bloque[0]["t"])
            nombre, nums = [], []
            for tk in bloque[1:]:
                v = a_num(tk["t"])
                if v is None:
                    if tk["t"] in MARCAS_FILA:
                        continue
                    if not nums and tk["t"] not in ("Hoja",):
                        nombre.append(tk["t"])
                    continue
                nums.append({**tk, "v": v})
            filas.append(interpretar(codigo, " ".join(nombre), nums, lay, pg["n"], orden))
            orden += 1
    return filas


def interpretar(codigo, indicador, nums, lay, pagina, orden):
    fila = {"codigo": codigo, "indicador": indicador, "pagina": pagina, "maquetado": lay,
            "plan": None, "real": None, "pct": None, "plan_acu": None, "real_acu": None,
            "estado": "ok", "motivos": []}
    if not nums:
        fila["estado"] = "dudosa"
        fila["motivos"].append("sin números")
        return fila
    # líneas por y: la de arriba es el mes (M/MES), la siguiente el acumulado (A/ACU)
    ys = sorted({n["y"] for n in nums}, reverse=True)
    grupos = []
    for y in ys:
        if grupos and abs(grupos[-1][0] - y) <= 3:
            continue
        grupos.append((y, [n for n in nums if abs(n["y"] - y) <= 3]))
    # quitar duplicados (un número entra en un solo grupo)
    vistos, lineas = set(), []
    for y, g in grupos:
        g = [n for n in g if id(n) not in vistos]
        vistos.update(id(n) for n in g)
        if g:
            lineas.append(g)

    def valores(linea):
        if lay == "L12":
            return [n["v"] for n in sorted(linea, key=lambda n: n["x"])]
        return [n["v"] for n in linea]  # L9 y L3: orden de dibujo

    esperado = {"L12": 12, "L9": 9, "L3": 3}[lay]
    m = valores(lineas[0]) if lineas else []
    a = valores(lineas[1]) if len(lineas) > 1 else []
    fila["valores_mes"] = m
    fila["valores_acu"] = a
    if len(lineas) > 2:
        fila["motivos"].append(f"{len(lineas)} líneas de números (se esperaban 2)")

    def partes(v):
        if len(v) != esperado:
            return None
        if lay == "L12":
            return {"mlc_p": v[1], "mlc_r": v[2], "mn_p": v[5], "mn_r": v[6],
                    "p": v[9], "r": v[10], "pct": v[11], "mn_pct": v[7]}
        if lay == "L9":
            return {"mlc_p": v[0], "mlc_r": v[1], "mn_p": v[2], "mn_r": v[3],
                    "p": v[4], "r": v[5], "pct": v[8], "mn_pct": v[7]}
        return {"p": v[0], "r": v[1], "pct": v[2]}

    pm, pa = partes(m), partes(a)
    if pm is None:
        fila["estado"] = "dudosa"
        fila["motivos"].append(f"la línea del mes tiene {len(m)} números y se esperaban {esperado}")
        return fila
    fila["plan"], fila["real"], fila["pct"] = pm["p"], pm["r"], pm["pct"]
    if pa:
        fila["plan_acu"], fila["real_acu"] = pa["p"], pa["r"]
    elif a:
        fila["motivos"].append(f"acumulado con {len(a)} números (se esperaban {esperado}); no se usa")

    for etiqueta, p in (("mes", pm), ("acumulado", pa)):
        if not p:
            continue
        if "MASK" in (p["p"], p["r"]):
            if etiqueta == "mes":
                fila["estado"] = "dudosa"
                fila["motivos"].append("el PDF tapa el valor del mes con asteriscos")
            continue
        if lay != "L3":
            if iguales(p["p"], p["mlc_p"] + p["mn_p"] if "MASK" not in (p["mlc_p"], p["mn_p"]) else None) is False:
                fila["motivos"].append(f"{etiqueta}: Plan TOTAL ≠ MLC + MN")
                if etiqueta == "mes":
                    fila["estado"] = "dudosa"
            if iguales(p["r"], p["mlc_r"] + p["mn_r"] if "MASK" not in (p["mlc_r"], p["mn_r"]) else None) is False:
                fila["motivos"].append(f"{etiqueta}: Real TOTAL ≠ MLC + MN")
                if etiqueta == "mes":
                    fila["estado"] = "dudosa"
        if p["pct"] != "MASK" and p["p"] not in (0, "MASK") and p["r"] != "MASK":
            esperado_pct = p["r"] / p["p"] * 100
            if abs(esperado_pct - p["pct"]) > max(TOL_PCT, abs(esperado_pct) * 1e-4):
                fila["motivos"].append(f"{etiqueta}: % impreso {p['pct']} y Real/Plan da {round(esperado_pct, 2)}")
                if etiqueta == "mes":
                    fila["estado"] = "dudosa"
    return fila


def analizar_pdf(path):
    paginas, meta = leer_pdf(path)
    cab, avisos = cabecera(paginas)
    filas = filas_pdf(paginas)
    por_cod = {}
    for f in filas:
        por_cod.setdefault(f["codigo"], []).append(f)
    for cod, fs in por_cod.items():
        if len(fs) > 1:
            avisos.append(f"El código {cod:03d} aparece {len(fs)} veces (páginas {[f['pagina'] for f in fs]}).")
            for f in fs:
                f["estado"] = "dudosa"
                f["motivos"].append("código repetido en el PDF")
    nombre = path.name
    pn, mn = polo_por_nombre(nombre), mes_por_nombre(nombre)
    if cab["polo"] and pn and pn != cab["polo"]:
        avisos.append(f"El nombre del fichero apunta a {POLOS[pn][0]} pero el contenido es de "
                      f"{POLOS[cab['polo']][0]} ({cab['unidad']}). Manda el contenido.")
    if cab["mes"] and mn and (mn[0], mn[1] or cab["anio"]) != (cab["mes"], cab["anio"]):
        avisos.append(f"El nombre del fichero sugiere {MESES[mn[0] - 1].title()}/{mn[1]} pero el PDF es de "
                      f"{MESES[cab['mes'] - 1].title()}/{cab['anio']}. Manda el contenido.")
    if "Producer" in str(meta) and "PyPDF" in str(meta.get("/Producer", "")):
        avisos.append("El PDF fue reescrito con PyPDF2 después de emitirse (no es el original del sistema contable).")
    return {"archivo": nombre, "cabecera": cab, "polo_nombre_fichero": pn,
            "mes_nombre_fichero": mn, "maquetados": sorted({maquetado(p['trozos']) for p in paginas}),
            "paginas": len(paginas), "productor": meta.get("/Producer"), "filas": filas, "avisos": avisos}


# ---------------------------------------------------------------- el GLOBAL

def ruta_global(anio):
    aa = f"{anio % 100:02d}"
    for p in (BASE_DIR / f"GLOBAL-CP{aa}.xlsx", BASE_DIR / f"Leslie Global {anio}" / f"GLOBAL-CP{aa}.xlsx",
              BASE_DIR.parent / f"Leslie Global {anio}" / f"GLOBAL-CP{aa}.xlsx"):
        if p.exists():
            return p
    return None


def hoja_global(polo, anio):
    return f"{polo}CP{anio}"


def localizar_mes(ws, mes):
    """Columnas PLAN y REAL del mes (fila 4 el mes, fila 5 PLAN/REAL)."""
    nombre = MESES[mes - 1]
    for c in range(1, ws.max_column + 1):
        v = ws.cell(4, c).value
        if isinstance(v, str) and mayus(v) in (nombre, "SETIEMBRE" if mes == 9 else nombre):
            fin = c
            for c2 in range(c + 1, ws.max_column + 1):
                v2 = ws.cell(4, c2).value
                if v2 is not None and str(v2).strip():
                    break
                fin = c2
            cp = cr = None
            for c3 in range(c, fin + 1):
                t = mayus(ws.cell(5, c3).value or "")
                if "PLAN" in t and cp is None:
                    cp = c3
                if "REAL" in t and cr is None:
                    cr = c3
            if cp and cr:
                return cp, cr
    return None


def mapa_codigos(ws):
    col = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(4, c).value
        if isinstance(v, str) and mayus(v) == "FILAS":
            col = c
            break
    if col is None:
        raise KeyError("No se encontró la cabecera 'Filas' en la fila 4 del GLOBAL.")
    mapa = {}
    for r in range(6, ws.max_row + 1):
        v = ws.cell(r, col).value
        try:
            mapa.setdefault(int(v), r)
        except (TypeError, ValueError):
            continue
    return mapa, col


def evaluador(ws):
    sys.path.insert(0, str(BASE_DIR))
    from sync_global_ai import Evaluador
    return Evaluador(ws)


def evaluar(ev, r, c):
    try:
        v = ev.valor(r, c)
        return float(v) if isinstance(v, (int, float)) else None
    except Exception:
        return None


def es_formula(v):
    return isinstance(v, str) and v.startswith("=")


def coord(r, c):
    from openpyxl.utils import get_column_letter
    return f"{get_column_letter(c)}{r}"


# ---------------------------------------------------------- preparar y vistas

def resumen_fichero(an, lote_dir):
    """Cruce de un PDF con el GLOBAL: destino, cambios previstos y problemas."""
    import openpyxl
    cab = an["cabecera"]
    res = {"archivo": an["archivo"], "unidad": cab["unidad"],
           "polo": cab["polo"], "polo_nombre": POLOS[cab["polo"]][0] if cab["polo"] else None,
           "mes": cab["mes"], "anio": cab["anio"],
           "mes_texto": f"{MESES[cab['mes'] - 1].title()}/{cab['anio']}" if cab["mes"] else None,
           "maquetados": an["maquetados"], "paginas": an["paginas"],
           "filas_leidas": len(an["filas"]), "avisos": list(an["avisos"]), "bloqueos": []}
    if not cab["polo"]:
        res["bloqueos"].append(f"No se reconoce el polo en la cabecera ('{cab['unidad']}').")
    if not cab["mes"]:
        res["bloqueos"].append("No se encuentra el mes (<Mes>/<AAAA>) en la cabecera.")
    if cab["meses_distintos"] > 1:
        res["bloqueos"].append("La cabecera visible tiene varios meses.")
    dudosas = [f for f in an["filas"] if f["estado"] != "ok"]
    res["dudosas"] = [{"codigo": f["codigo"], "indicador": f["indicador"], "pagina": f["pagina"],
                       "plan": limpio(f["plan"]), "real": limpio(f["real"]), "motivos": f["motivos"]}
                      for f in dudosas]
    if res["bloqueos"]:
        return res
    gpath = ruta_global(cab["anio"])
    if not gpath:
        res["bloqueos"].append(f"No existe GLOBAL-CP{cab['anio'] % 100:02d}.xlsx en el servidor "
                               f"({BASE_DIR}). Hay que subir el maestro de {cab['anio']}.")
        return res
    hoja = hoja_global(cab["polo"], cab["anio"])
    res["global"] = gpath.name
    res["hoja"] = hoja
    wb = openpyxl.load_workbook(gpath)
    if hoja not in wb.sheetnames:
        res["bloqueos"].append(f"El GLOBAL {gpath.name} no tiene la hoja {hoja}.")
        return res
    ws = wb[hoja]
    cols = localizar_mes(ws, cab["mes"])
    if not cols:
        res["bloqueos"].append(f"La hoja {hoja} no tiene el mes {MESES[cab['mes'] - 1]} con PLAN y REAL.")
        return res
    cp, cr = cols
    mapa, _ = mapa_codigos(ws)
    ev = evaluador(ws)
    res["columnas"] = {"plan": coord(5, cp)[:-1], "real": coord(5, cr)[:-1]}
    nuevos, sobrescribe, iguales_n, formulas, fuera = 0, [], 0, [], []
    for f in an["filas"]:
        if f["codigo"] not in mapa:
            fuera.append(f["codigo"])
            continue
        r = mapa[f["codigo"]]
        for campo, c in (("plan", cp), ("real", cr)):
            nuevo = f[campo]
            if nuevo is None or nuevo == "MASK":
                continue
            actual = ws.cell(r, c).value
            if es_formula(actual):
                calc = evaluar(ev, r, c)
                formulas.append({"codigo": f["codigo"], "campo": campo, "celda": coord(r, c),
                                 "formula": actual, "pdf": limpio(nuevo), "calculado_hoy": limpio(calc)})
            elif actual is None or actual == "":
                nuevos += 1
            elif iguales(actual, nuevo):
                iguales_n += 1
            else:
                sobrescribe.append({"codigo": f["codigo"], "campo": campo, "celda": coord(r, c),
                                    "antes": limpio(actual), "pdf": limpio(nuevo)})
    res["cambios"] = {"celdas_vacias_a_rellenar": nuevos, "ya_iguales": iguales_n,
                      "sobrescribe": sobrescribe, "celdas_con_formula": len(formulas)}
    res["codigos_sin_fila_en_global"] = fuera
    vacias = [cod for cod in mapa if cod not in {f["codigo"] for f in an["filas"]}]
    res["codigos_global_sin_dato_en_pdf"] = len(vacias)
    return res


def recoger_subido(lote_dir, subido):
    """n8n no crea carpetas al escribir, así que deja el fichero plano en Leslie/
    como m1_<lote>__<nombre>; aquí se mueve a lote/entrada/."""
    if not subido:
        return
    subido = os.path.basename(str(subido))
    if not subido.startswith(f"m1_{lote_dir.name}__"):
        error(f"Nombre de fichero subido inesperado: {subido}")
    origen = BASE_DIR / subido
    if origen.exists():
        entrada = lote_dir / "entrada"
        entrada.mkdir(parents=True, exist_ok=True)
        shutil.move(str(origen), str(entrada / subido.split("__", 1)[1]))


def purgar_viejos(dias=3):
    """Lotes y subidas que quedaron a medias (un fallo antes de 'cerrar')."""
    limite = datetime.now().timestamp() - dias * 86400
    for p in list(LOTES.glob("*")) + list(BASE_DIR.glob("m1_*__*")):
        try:
            if p.stat().st_mtime < limite:
                shutil.rmtree(p) if p.is_dir() else p.unlink()
        except OSError:
            pass


def cmd_preparar(args):
    lote_dir = dir_lote(args)
    purgar_viejos()
    recoger_subido(lote_dir, args.get("subido"))
    try:
        pdfs, avisos = desempaquetar(lote_dir)
    except Exception as e:
        error(str(e))
    if not pdfs:
        error("El fichero subido no contiene ningún PDF.", avisos=avisos)
    analisis, resumen = {}, []
    for p in pdfs:
        try:
            an = analizar_pdf(p)
        except Exception as e:
            resumen.append({"archivo": p.name, "bloqueos": [f"No se pudo leer el PDF: {e}"]})
            continue
        analisis[p.name] = an
        try:
            resumen.append(resumen_fichero(an, lote_dir))
        except Exception as e:
            resumen.append({"archivo": p.name, "bloqueos": [f"No se pudo cruzar con el GLOBAL: {e}"]})
    # dos PDF con el mismo polo y mes en un lote
    vistos = {}
    for r in resumen:
        k = (r.get("polo"), r.get("mes"), r.get("anio"))
        if None not in k:
            vistos.setdefault(k, []).append(r["archivo"])
    for k, archivos in vistos.items():
        if len(archivos) > 1:
            avisos.append(f"Hay {len(archivos)} PDF del mismo polo y mes ({POLOS[k[0]][0]} "
                          f"{MESES[k[1] - 1].title()}/{k[2]}): {archivos}. Excluye todos menos uno.")
    (lote_dir / "extraccion.json").write_text(json.dumps(analisis, ensure_ascii=False, default=str), encoding="utf-8")
    salir({"ok": True, "script_sha256": sha256(__file__), "lote": lote_dir.name,
           "ficheros": resumen, "avisos": avisos})


def cargar_extraccion(lote_dir):
    p = lote_dir / "extraccion.json"
    if not p.exists():
        error("No hay extracción en este lote: hay que llamar antes a 'preparar'.")
    return json.loads(p.read_text(encoding="utf-8"))


def elegir(analisis, archivo):
    if archivo in analisis:
        return analisis[archivo]
    cands = [k for k in analisis if archivo and archivo.lower() in k.lower()]
    if len(cands) == 1:
        return analisis[cands[0]]
    error(f"No encuentro el PDF '{archivo}'. Disponibles: {list(analisis)}")


def cmd_ver_filas(args):
    an = elegir(cargar_extraccion(dir_lote(args)), args.get("archivo"))
    buscar = mayus(args.get("buscar") or "")
    desde, hasta = int(args.get("desde") or 0), int(args.get("hasta") or 999)
    filas = []
    for f in an["filas"]:
        if not (desde <= f["codigo"] <= hasta):
            continue
        if buscar and buscar not in mayus(f["indicador"]) and buscar != f"{f['codigo']:03d}":
            continue
        filas.append({"codigo": f["codigo"], "indicador": f["indicador"], "pagina": f["pagina"],
                      "plan": limpio(f["plan"]), "real": limpio(f["real"]), "pct": limpio(f["pct"]),
                      "plan_acu": limpio(f["plan_acu"]), "real_acu": limpio(f["real_acu"]),
                      "estado": f["estado"], "motivos": f["motivos"]})
    salir({"ok": True, "archivo": an["archivo"], "filas": filas})


def cmd_ver_pdf(args):
    lote_dir = dir_lote(args)
    an = elegir(cargar_extraccion(lote_dir), args.get("archivo"))
    paginas, _ = leer_pdf(lote_dir / "pdf" / an["archivo"])
    codigo = args.get("codigo")
    pagina = args.get("pagina")
    out = []
    for pg in paginas:
        if pagina and int(pagina) != pg["n"]:
            continue
        tr = pg["trozos"]
        if codigo:
            c = f"{int(codigo):03d}"
            idx = [i for i, t in enumerate(tr) if t["t"] == c or t["t"].startswith(c + " ")]
            for i in idx:
                out.append({"pagina": pg["n"], "trozos": tr[max(0, i - 4): i + 40]})
        else:
            out.append({"pagina": pg["n"], "trozos": tr[:400]})
    if not out:
        error(f"No encuentro el código {codigo} en el PDF.")
    salir({"ok": True, "archivo": an["archivo"],
           "nota": "Cada trozo es texto tal como está dibujado: x crece hacia la derecha, y hacia ARRIBA. "
                   "La línea del mes (M/MES) es la de mayor y; debajo va el acumulado (A/ACU).",
           "resultado": out})


def cmd_ver_global(args):
    import openpyxl
    an = elegir(cargar_extraccion(dir_lote(args)), args.get("archivo"))
    cab = an["cabecera"]
    gpath = ruta_global(cab["anio"]) if cab.get("anio") else None
    if not gpath or not cab.get("polo"):
        error("Este PDF no tiene GLOBAL o polo de destino.")
    wb = openpyxl.load_workbook(gpath)
    ws = wb[hoja_global(cab["polo"], cab["anio"])]
    cp, cr = localizar_mes(ws, cab["mes"])
    mapa, _ = mapa_codigos(ws)
    ev = evaluador(ws)
    buscar = mayus(args.get("buscar") or "")
    pdf = {f["codigo"]: f for f in an["filas"]}
    filas = []
    for cod, r in sorted(mapa.items()):
        titulo = str(ws.cell(r, 2).value or "").strip()
        if buscar and buscar not in mayus(titulo) and buscar != str(cod):
            continue
        fila = {"codigo": cod, "titulo": titulo, "fila_excel": r}
        for campo, c in (("plan", cp), ("real", cr)):
            v = ws.cell(r, c).value
            fila[campo] = {"celda": coord(r, c), "valor": limpio(v) if not es_formula(v) else None,
                           "formula": v if es_formula(v) else None,
                           "calculado": limpio(evaluar(ev, r, c)) if es_formula(v) else None}
        if cod in pdf:
            fila["pdf"] = {"plan": limpio(pdf[cod]["plan"]), "real": limpio(pdf[cod]["real"])}
        filas.append(fila)
    salir({"ok": True, "global": gpath.name, "hoja": ws.title, "mes": MESES[cab["mes"] - 1], "filas": filas})


# --------------------------------------------------------------- escritura

class Cerrojo:
    """Evita dos escrituras simultáneas del GLOBAL (Linux)."""

    def __enter__(self):
        BASE_DIR.joinpath(".modelo1.lock").touch()
        self.f = open(BASE_DIR / ".modelo1.lock", "r+")
        try:
            import fcntl
            fcntl.flock(self.f, fcntl.LOCK_EX)
        except ImportError:
            pass
        return self

    def __exit__(self, *a):
        try:
            import fcntl
            fcntl.flock(self.f, fcntl.LOCK_UN)
        except ImportError:
            pass
        self.f.close()


def respaldar(gpath):
    RESPALDOS.mkdir(exist_ok=True)
    destino = RESPALDOS / f"{gpath.stem}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
    shutil.copy2(gpath, destino)
    viejos = sorted(RESPALDOS.glob(f"{gpath.stem}_*.xlsx"))
    for v in viejos[:-MAX_RESPALDOS]:
        v.unlink()
    return destino


def decisiones(an, peticion):
    """Valores finales de un PDF: extracción + correcciones del agente - exclusiones."""
    excluir = {int(c) for c in (peticion.get("excluir") or [])}
    correcciones = {int(k): v for k, v in (peticion.get("correcciones") or {}).items()}
    valores, notas = {}, []
    for f in an["filas"]:
        cod = f["codigo"]
        if cod in excluir:
            notas.append({"codigo": cod, "nota": "excluida por el agente"})
            continue
        plan, real = f["plan"], f["real"]
        if cod in correcciones:
            corr = correcciones[cod]
            plan = corr.get("plan", plan)
            real = corr.get("real", real)
            notas.append({"codigo": cod, "nota": "corregida por el agente", "pdf_leido": [limpio(f["plan"]), limpio(f["real"])],
                          "decidido": [limpio(plan), limpio(real)]})
        elif f["estado"] != "ok":
            notas.append({"codigo": cod, "nota": "dudosa sin corregir ni excluir: no se escribe",
                          "motivos": f["motivos"]})
            continue
        valores[cod] = {campo: (float(v) if isinstance(v, (int, float)) else None)
                        for campo, v in (("plan", plan), ("real", real))}
    for cod in correcciones:
        if cod not in {f["codigo"] for f in an["filas"]}:
            corr = correcciones[cod]
            valores[cod] = {"plan": corr.get("plan"), "real": corr.get("real")}
            notas.append({"codigo": cod, "nota": "añadida por el agente (no estaba en la extracción)"})
    return valores, notas


def cmd_escribir(args):
    import openpyxl
    lote_dir = dir_lote(args)
    analisis = cargar_extraccion(lote_dir)
    peticiones = args.get("archivos")
    if isinstance(peticiones, str):
        peticiones = json.loads(peticiones)
    if not peticiones:
        error("Falta 'archivos': lista de {archivo, correcciones?, excluir?}.")

    planes, rechazados = {}, []
    for pet in peticiones:
        an = elegir(analisis, pet.get("archivo"))
        cab = an["cabecera"]
        polo = pet.get("polo") or cab["polo"]
        if polo not in POLOS:
            rechazados.append({"archivo": an["archivo"], "motivo": f"polo desconocido: {polo}"})
            continue
        if pet.get("polo") and cab["polo"] and pet["polo"] != cab["polo"]:
            rechazados.append({"archivo": an["archivo"], "motivo":
                               f"el PDF es de {POLOS[cab['polo']][0]}: no se puede escribir en otro polo"})
            continue
        if not cab["mes"] or cab["meses_distintos"] > 1:
            rechazados.append({"archivo": an["archivo"], "motivo": "mes no determinado sin ambigüedad"})
            continue
        gpath = ruta_global(cab["anio"])
        if not gpath:
            rechazados.append({"archivo": an["archivo"], "motivo": f"no existe GLOBAL-CP{cab['anio'] % 100:02d}.xlsx"})
            continue
        clave = (gpath, hoja_global(polo, cab["anio"]), cab["mes"])
        if any(k[1:] == clave[1:] and k[0] == clave[0] for k in [p["clave"] for g in planes.values() for p in g]):
            rechazados.append({"archivo": an["archivo"], "motivo": "otro PDF del lote ya escribe ese polo y mes"})
            continue
        valores, notas = decisiones(an, pet)
        planes.setdefault(gpath, []).append({"clave": clave, "archivo": an["archivo"], "valores": valores,
                                             "notas": notas, "polo": polo, "mes": cab["mes"], "anio": cab["anio"]})

    resultado = {"ok": True, "escritos": [], "rechazados": rechazados}
    registro = []
    with Cerrojo():
        for gpath, lista in planes.items():
            respaldo = respaldar(gpath)
            wb = openpyxl.load_workbook(gpath)
            for plan in lista:
                _, hoja, mes = plan["clave"]
                if hoja not in wb.sheetnames:
                    rechazados.append({"archivo": plan["archivo"], "motivo": f"no existe la hoja {hoja}"})
                    continue
                ws = wb[hoja]
                cols = localizar_mes(ws, mes)
                if not cols:
                    rechazados.append({"archivo": plan["archivo"], "motivo": f"la hoja {hoja} no tiene {MESES[mes - 1]}"})
                    continue
                cp, cr = cols
                mapa, _ = mapa_codigos(ws)
                cambios, formulas_ok, formulas_sustituidas, sin_fila = [], [], [], []
                pendientes_formula = []
                # 1) celdas sin fórmula
                for cod, vals in sorted(plan["valores"].items()):
                    if cod not in mapa:
                        sin_fila.append(cod)
                        continue
                    r = mapa[cod]
                    for campo, c in (("plan", cp), ("real", cr)):
                        nuevo = vals.get(campo)
                        if nuevo is None:
                            continue
                        actual = ws.cell(r, c).value
                        if es_formula(actual):
                            pendientes_formula.append((cod, campo, r, c, nuevo, actual))
                            continue
                        if not iguales(actual, nuevo):
                            cambios.append({"codigo": cod, "campo": campo, "celda": coord(r, c),
                                            "antes": limpio(actual), "ahora": limpio(nuevo)})
                        ws.cell(r, c).value = nuevo
                # 2) fórmulas: se quedan si ya dan el valor del PDF (con los detalles nuevos).
                #    Sustituir una puede descuadrar otra que dependa de ella (Ventas Netas
                #    usa Ventas Totales), así que se repite hasta que no cambie nada.
                while True:
                    ev = evaluador(ws)
                    nuevas = []
                    for cod, campo, r, c, nuevo, formula in pendientes_formula:
                        if not es_formula(ws.cell(r, c).value):
                            continue
                        calc = evaluar(ev, r, c)
                        if calc is None or not iguales(calc, nuevo, TOL_FORMULA):
                            nuevas.append({"codigo": cod, "campo": campo, "celda": coord(r, c),
                                           "formula": formula, "calculaba": limpio(calc), "pdf": limpio(nuevo)})
                            ws.cell(r, c).value = nuevo
                    formulas_sustituidas += nuevas
                    if not nuevas:
                        break
                for cod, campo, r, c, nuevo, formula in pendientes_formula:
                    if es_formula(ws.cell(r, c).value):
                        formulas_ok.append({"codigo": cod, "campo": campo, "celda": coord(r, c),
                                            "redondeo": round(nuevo - evaluar(ev, r, c), 2)})
                info = {"archivo": plan["archivo"], "global": gpath.name, "hoja": hoja,
                        "mes": MESES[mes - 1], "anio": plan["anio"],
                        "columnas": [coord(5, cp)[:-1], coord(5, cr)[:-1]],
                        "celdas_cambiadas": len(cambios), "cambios_con_valor_previo":
                            [x for x in cambios if x["antes"] not in (None, "")],
                        "formulas_conservadas": len(formulas_ok),
                        "formulas_con_redondeo": [x for x in formulas_ok if abs(x["redondeo"]) > TOL],
                        "formulas_sustituidas_por_valor_pdf": formulas_sustituidas,
                        "codigos_sin_fila_en_global": sin_fila, "notas": plan["notas"]}
                resultado["escritos"].append(info)
                registro.append({**info, "valores": {str(k): v for k, v in plan["valores"].items()},
                                 "formulas_conservadas_detalle": formulas_ok, "respaldo": str(respaldo),
                                 "celdas_escritas": [x["celda"] for x in cambios] + [x["celda"] for x in formulas_sustituidas]})
            tmp = gpath.with_suffix(".tmp.xlsx")
            wb.save(tmp)
            os.replace(tmp, gpath)
            resultado.setdefault("respaldos", []).append(respaldo.name)
    previo = []
    if (lote_dir / "escritura.json").exists():
        previo = json.loads((lote_dir / "escritura.json").read_text(encoding="utf-8"))
    (lote_dir / "escritura.json").write_text(json.dumps(previo + registro, ensure_ascii=False, default=str),
                                             encoding="utf-8")
    salir(resultado)


# --------------------------------------------------------------- verificación

def cmd_verificar(args):
    """Revisión final, independiente de la escritura: relee los PDF y el GLOBAL."""
    import openpyxl
    lote_dir = dir_lote(args)
    p = lote_dir / "escritura.json"
    if not p.exists():
        error("No se ha escrito nada en este lote.")
    historial = json.loads(p.read_text(encoding="utf-8"))
    registro = {}
    for r in historial:
        registro[r["archivo"]] = r  # la última escritura de cada PDF manda
    libros = {}

    def libro(nombre_global, anio):
        if nombre_global not in libros:
            libros[nombre_global] = openpyxl.load_workbook(ruta_global(anio))
        return libros[nombre_global]

    informe = {"ok": True, "ficheros": [], "libros": []}
    for archivo, reg in registro.items():
        an = analizar_pdf(lote_dir / "pdf" / archivo)      # segunda lectura, desde cero
        cab = an["cabecera"]
        res = {"archivo": archivo, "hoja": reg["hoja"], "mes": f"{reg['mes'].title()}/{reg['anio']}",
               "errores": [], "avisos": []}
        if cab["polo"] != reg["hoja"][:-6]:
            res["errores"].append(f"El PDF es de {cab['polo']} y se escribió en la hoja {reg['hoja']}.")
        if MESES[cab["mes"] - 1] != reg["mes"] or cab["anio"] != reg["anio"]:
            res["errores"].append(f"El PDF es de {MESES[cab['mes'] - 1]}/{cab['anio']} y se escribió en "
                                  f"{reg['mes']}/{reg['anio']}.")
        ws = libro(reg["global"], reg["anio"])[reg["hoja"]]
        cp, cr = localizar_mes(ws, MES_NUM[reg["mes"]])
        if [coord(5, cp)[:-1], coord(5, cr)[:-1]] != reg["columnas"]:
            res["errores"].append("Las columnas del mes no son las que se escribieron.")
        mapa, _ = mapa_codigos(ws)
        ev = evaluador(ws)
        decididos = {int(k): v for k, v in reg["valores"].items()}
        leidos = {f["codigo"]: f for f in an["filas"]}
        comprobadas = 0
        for cod, vals in decididos.items():
            if cod not in mapa:
                continue
            r = mapa[cod]
            for campo, c in (("plan", cp), ("real", cr)):
                quiero = vals.get(campo)
                if quiero is None:
                    continue
                v = ws.cell(r, c).value
                hay = evaluar(ev, r, c) if es_formula(v) else v
                comprobadas += 1
                if not iguales(hay, quiero, TOL_FORMULA if es_formula(v) else TOL):
                    res["errores"].append(f"{cod:03d} {campo}: el GLOBAL tiene {limpio(hay)} en {coord(r, c)} y "
                                          f"se decidió {limpio(quiero)}.")
                f = leidos.get(cod)
                if f and f["estado"] == "ok" and not iguales(f[campo], quiero):
                    res["avisos"].append(f"{cod:03d} {campo}: se escribió {limpio(quiero)} pero el PDF dice "
                                         f"{limpio(f[campo])} (corrección del agente).")
        for cod, f in leidos.items():
            if cod in mapa and cod not in decididos:
                res["avisos"].append(f"{cod:03d} {f['indicador']}: está en el PDF y no se escribió.")
        # acumulado: suma de los meses del GLOBAL contra el acumulado del PDF
        acum = []
        mes_n = MES_NUM[reg["mes"]]
        cols_meses = [localizar_mes(ws, m) for m in range(1, mes_n + 1)]
        for cod, f in leidos.items():
            if cod not in mapa or None in cols_meses:
                continue
            r = mapa[cod]
            for campo, idx, acu in (("plan", 0, f["plan_acu"]), ("real", 1, f["real_acu"])):
                if acu in (None, "MASK"):
                    continue
                suma = 0.0
                for cols in cols_meses:
                    c = cols[idx]
                    v = ws.cell(r, c).value
                    v = evaluar(ev, r, c) if es_formula(v) else v
                    suma += float(v) if isinstance(v, (int, float)) else 0.0
                if abs(suma - acu) > max(1.0, abs(acu) * 1e-6):
                    acum.append({"codigo": cod, "campo": campo, "suma_meses_global": round(suma, 2),
                                 "acumulado_pdf": limpio(acu), "diferencia": round(acu - suma, 2)})
        if acum:
            res["acumulado_no_cuadra"] = acum[:40]
            res["avisos"].append(f"{len(acum)} celdas: la suma de enero a {reg['mes'].lower()} en el GLOBAL "
                                 f"no da el acumulado del PDF (meses anteriores sin cargar o distintos).")
        res["celdas_comprobadas"] = comprobadas
        res["correcto"] = not res["errores"]
        informe["ficheros"].append(res)

    # nada fuera de lo previsto: cada GLOBAL contra su copia de antes del lote
    for nombre_global in libros:
        escrituras = [r for r in historial if r["global"] == nombre_global]
        original = escrituras[0]["respaldo"]
        permitidas = {(r["hoja"], c) for r in escrituras for c in r["celdas_escritas"]}
        res = {"global": nombre_global, "copia_previa": Path(original).name, "errores": []}
        try:
            wb, wb_prev = libros[nombre_global], openpyxl.load_workbook(original)
            revisadas = 0
            for hoja in wb.sheetnames:
                if hoja not in wb_prev.sheetnames:
                    res["errores"].append(f"Hoja nueva que no estaba antes: {hoja}")
                    continue
                a, b = wb_prev[hoja], wb[hoja]
                for row in b.iter_rows(min_row=1, max_row=max(a.max_row, b.max_row),
                                       max_col=max(a.max_column, b.max_column)):
                    for cell in row:
                        revisadas += 1
                        if a[cell.coordinate].value != cell.value and (hoja, cell.coordinate) not in permitidas:
                            res["errores"].append(f"Celda no prevista cambiada: {hoja}!{cell.coordinate} "
                                                  f"({limpio(a[cell.coordinate].value)} -> {limpio(cell.value)}).")
            res["celdas_revisadas"] = revisadas
        except Exception as e:
            res["errores"].append(f"No se pudo comparar con la copia previa: {e}")
        res["errores"] = res["errores"][:30]
        res["correcto"] = not res["errores"]
        informe["libros"].append(res)
    informe["todo_correcto"] = (all(f["correcto"] for f in informe["ficheros"])
                                and all(x["correcto"] for x in informe["libros"]))
    salir(informe)


def relativa(ruta):
    """Ruta respecto a la carpeta Polarier: el runner y n8n pueden montarla en sitios distintos."""
    return Path(ruta).resolve().relative_to(BASE_DIR.parent).as_posix()


def cmd_cerrar(args):
    lote_dir = dir_lote(args)
    globales = []
    p = lote_dir / "escritura.json"
    if p.exists():
        for r in json.loads(p.read_text(encoding="utf-8")):
            if r["global"] not in [g["nombre"] for g in globales]:
                ruta = ruta_global(r["anio"])
                globales.append({"nombre": r["global"], "ruta": str(ruta), "relativa": relativa(ruta)})
    entregar = None
    if len(globales) == 1:
        entregar = {"ruta": globales[0]["ruta"], "relativa": globales[0]["relativa"], "nombre": globales[0]["nombre"],
                    "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    elif globales:
        SALIDA.mkdir(exist_ok=True)
        for viejo in sorted(SALIDA.glob("GLOBALES_*.zip"))[:-5]:
            viejo.unlink()
        destino = SALIDA / f"GLOBALES_{lote_dir.name}.zip"
        with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
            for g in globales:
                z.write(g["ruta"], g["nombre"])
        entregar = {"ruta": str(destino), "relativa": relativa(destino), "nombre": destino.name,
                    "mime": "application/zip"}
    if not args.get("conservar"):
        shutil.rmtree(lote_dir, ignore_errors=True)
    salir({"ok": True, "globales": globales, "entregar": entregar})


def cmd_version(args):
    """Comprobación tras subir el script al VPS: huella y de dónde sale pypdf."""
    try:
        try:
            import pypdf
        except ImportError:
            for p in sorted(BASE_DIR.glob("pypdf-*.whl"), reverse=True) + [BASE_DIR / "_vendor"]:
                sys.path.insert(0, str(p))
            import pypdf
        info_pdf = {"version": pypdf.__version__, "origen": pypdf.__file__}
    except Exception as e:
        info_pdf = {"error": str(e)}
    import openpyxl
    ok = "error" not in info_pdf
    wheel = None
    if args.get("wheel"):
        p = BASE_DIR / os.path.basename(str(args["wheel"]))
        hay = sha256(p) if p.exists() else None
        wheel = {"fichero": p.name, "sha256": hay, "esperado": args.get("wheel_sha256"),
                 "coincide": bool(hay) and hay == args.get("wheel_sha256")}
        ok = ok and wheel["coincide"]
    salir({"ok": ok, "script_sha256": sha256(__file__), "python": sys.version.split()[0],
           "pypdf": info_pdf, "wheel": wheel, "openpyxl": openpyxl.__version__,
           "sync_global_ai": (BASE_DIR / "sync_global_ai.py").exists()})


COMANDOS = {
    "version": cmd_version,
    "preparar": cmd_preparar,
    "ver_filas": cmd_ver_filas,
    "ver_pdf": cmd_ver_pdf,
    "ver_global": cmd_ver_global,
    "escribir": cmd_escribir,
    "verificar": cmd_verificar,
    "cerrar": cmd_cerrar,
}


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in COMANDOS:
        error(f"Uso: agente_modelo1.py <{'|'.join(COMANDOS)}> '<json>'")
    try:
        args = json.loads(sys.argv[2]) if len(sys.argv) > 2 and sys.argv[2].strip() else {}
    except json.JSONDecodeError as e:
        error(f"JSON de argumentos inválido: {e}")
    try:
        COMANDOS[sys.argv[1]](args)
    except SystemExit:
        raise
    except Exception as e:
        error(f"{type(e).__name__}: {e}")


if __name__ == "__main__":
    main()
